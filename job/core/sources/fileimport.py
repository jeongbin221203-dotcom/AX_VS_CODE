"""CSV·엑셀로 공고 가져오기 — API가 없는 사이트(잡코리아·잡플래닛·인크루트 등)용.

사이트를 긁지 않고, 사용자가 옮겨 적거나 내려받은 표를 올린다. 열 이름은 한국어·영어 별칭을 받는다.
"""
from __future__ import annotations

import csv
import hashlib
import io

from core.postings import build

from . import SourceError

MAX_ROWS = 5000
COLUMNS = {
    "title": ("제목", "공고명", "공고 제목", "채용제목", "title", "position"),
    "company": ("회사", "회사명", "기업", "기업명", "company"),
    "url": ("링크", "url", "주소", "공고 링크", "지원 링크"),
    "location": ("지역", "근무지", "근무지역", "location", "region"),
    "salary_text": ("연봉", "급여", "salary", "임금"),
    "career": ("경력", "신입/경력", "career", "경력구분"),
    "education": ("학력", "education"),
    "employment_type": ("고용형태", "근무형태", "employment"),
    "deadline": ("마감", "마감일", "deadline", "접수마감"),
    "posted_at": ("등록일", "게시일", "posted"),
    "keywords": ("키워드", "기술", "스킬", "keywords", "skills"),
    "job_category": ("직무", "직종", "분야", "category"),
    "company_avg_salary": ("회사평균연봉", "평균연봉", "회사 평균연봉", "avg_salary"),
    "description": ("내용", "상세", "설명", "description"),
    "source": ("출처", "사이트", "source"),
}
TEMPLATE_HEADER = ["출처", "회사", "제목", "지역", "연봉", "경력", "학력", "고용형태", "마감일",
                   "링크", "키워드", "직무", "회사평균연봉"]
_SOURCE_ALIAS = {"잡코리아": "jobkorea", "jobkorea": "jobkorea", "사람인": "saramin", "saramin": "saramin",
                 "원티드": "wanted", "wanted": "wanted", "잡플래닛": "jobplanet", "jobplanet": "jobplanet",
                 "인크루트": "incruit", "incruit": "incruit", "고용24": "work24", "워크넷": "work24",
                 "링커리어": "linkareer", "linkareer": "linkareer", "자소설닷컴": "jasoseol", "자소설": "jasoseol",
                 "jasoseol": "jasoseol", "잡플레닛": "jobplanet", "리멤버": "remember", "remember": "remember"}


def read_table(filename: str, data: bytes) -> list[dict]:
    name = filename.lower()
    if name.endswith((".xlsx", ".xlsm")):
        return _read_xlsx(data)
    if name.endswith((".csv", ".txt")):
        return _read_csv(data)
    raise SourceError("CSV 또는 엑셀(.xlsx) 파일만 올릴 수 있습니다")


def _read_csv(data: bytes) -> list[dict]:
    for enc in ("utf-8-sig", "cp949"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise SourceError("파일 인코딩을 읽을 수 없습니다 (UTF-8 또는 CP949)")
    rows = list(csv.DictReader(io.StringIO(text)))
    if len(rows) > MAX_ROWS:
        raise SourceError(f"한 번에 {MAX_ROWS:,}행까지 올릴 수 있습니다")
    return rows


def _read_xlsx(data: bytes) -> list[dict]:
    try:
        from openpyxl import load_workbook
    except ImportError as e:
        raise SourceError("엑셀을 읽으려면 openpyxl 이 필요합니다 (pip install openpyxl)") from e
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as e:  # 손상된 파일, 압축 오류 등
        raise SourceError(f"엑셀 파일을 열 수 없습니다: {e.__class__.__name__}") from e
    ws = wb.worksheets[0]
    it = ws.iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else "" for h in next(it, [])]
    rows = []
    for values in it:
        if len(rows) >= MAX_ROWS:
            raise SourceError(f"한 번에 {MAX_ROWS:,}행까지 올릴 수 있습니다")
        if values and any(v not in (None, "") for v in values):
            rows.append({header[i]: v for i, v in enumerate(values) if i < len(header)})
    wb.close()
    return rows


def to_postings(rows: list[dict], default_source: str = "csv") -> tuple[list[dict], list[str]]:
    """(공고 목록, 건너뛴 행 설명)."""
    if not rows:
        raise SourceError("가져올 행이 없습니다")
    header = list(rows[0].keys())
    colmap = {}
    for field, aliases in COLUMNS.items():
        for h in header:
            if h and str(h).strip().lower().replace(" ", "") in {a.lower().replace(" ", "") for a in aliases}:
                colmap[field] = h
                break
    if "title" not in colmap or "company" not in colmap:
        raise SourceError("'회사'와 '제목' 열이 필요합니다. 양식 파일을 내려받아 확인하세요")

    out, skipped = [], []
    for i, row in enumerate(rows, start=2):
        v = {f: row.get(h) for f, h in colmap.items()}
        if not (v.get("title") and v.get("company")):
            skipped.append(f"{i}행: 회사·제목 없음")
            continue
        src_raw = str(v.pop("source", "") or "").strip()
        source = _SOURCE_ALIAS.get(src_raw.lower(), _SOURCE_ALIAS.get(src_raw, default_source))
        # 링크가 있으면 링크로, 없으면 회사+제목으로 같은 공고를 알아본다
        ident = str(v.get("url") or "") or f"{v['company']}|{v['title']}"
        sid = hashlib.sha1(ident.encode("utf-8")).hexdigest()[:16]
        out.append(build(source, sid, **{k: (str(x) if x is not None else None) for k, x in v.items()}))
    return out, skipped


def template_csv() -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(TEMPLATE_HEADER)
    w.writerow(["잡코리아", "(예) 한빛데이터", "백엔드 개발자 (Python)", "서울 강남구", "3,800~4,500만원", "경력 2년 이상",
                "대졸", "정규직", "2026-10-31", "https://www.jobkorea.co.kr/Recruit/GI_Read/...", "Python, Django, AWS",
                "소프트웨어 개발", "5200"])
    return buf.getvalue().encode("utf-8-sig")
