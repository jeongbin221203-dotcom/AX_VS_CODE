"""대기업 보완 — 거래처 중복·병합, 감사로그·백업 보관, 첨부, 회계연도, 개인정보 요청, 외화, 법인, 전자세금계산서."""
from __future__ import annotations

import io
import os
import re
import zipfile
from datetime import datetime, timedelta

import pytest
from conftest import biz, login, post, user
from openpyxl import load_workbook

from core import company
from core import database
from core import entities as ent_mod
from core import erp
from core import etax
from core import fiscal
from core import jobs
from core import privacy
from core import retention
from core import sales_db as db


# ── 1. 거래처 중복 · 병합 ─────────────────────────────────────────────────
def test_duplicate_customer_blocked_and_merge(app):
    rep = login(app, "김영업")
    number = biz("220811234")
    assert post(rep, "/customers/save", {"name": "(주)중복산업", "biz_no": number, "grade": "A",
                                         "industry": "제조"}).status_code == 302
    res = post(rep, "/customers/save", {"name": "다른이름", "biz_no": number.replace("", "")[:3] + "-" + number[3:5]
                                        + "-" + number[5:], "grade": "A", "industry": "제조"})
    assert res.status_code == 400 and "이미" in res.get_data(as_text=True)
    res = post(rep, "/customers/save", {"name": "중복산업 주식회사", "grade": "A", "industry": "제조"})
    assert res.status_code == 409 and "그래도 등록" in res.get_data(as_text=True)
    assert post(rep, "/customers/save", {"name": "중복산업 주식회사", "grade": "A", "industry": "제조",
                                         "confirm_similar": "1"}).status_code == 302
    src = db._one("SELECT id FROM customers WHERE name='중복산업 주식회사'")["id"]
    dst = db._one("SELECT id FROM customers WHERE name='(주)중복산업'")["id"]
    db.set_context("system", None)
    db.upsert_sale({"customer_id": src, "item": "병합전매출", "qty": 1, "unit_price": 10000,
                    "owner_id": user("김영업")["id"]})
    assert post(rep, "/customers/merge", {"source_id": src, "target_id": dst, "reason": "x"}).status_code == 403
    admin = login(app, "시스템관리자")
    assert "중복산업" in admin.get("/customers?tab=merge").get_data(as_text=True)
    post(admin, "/customers/merge", {"source_id": src, "target_id": dst, "reason": "같은 회사 이중 등록"})
    assert db._one("SELECT customer_id FROM sales WHERE item='병합전매출'")["customer_id"] == dst
    merged = db.get_customer(src)
    assert merged["merged_into"] == dst and merged["status"] == "종료"
    assert db._one("SELECT id FROM audit_log WHERE action='거래처병합'")


# ── 3. 감사로그 이관 · 백업 세대 ────────────────────────────────────────────
def test_audit_archive_keeps_chain_verifiable(app, isolated_db, monkeypatch, tmp_path):
    db.set_context("system", None)
    for i in range(5):
        db.audit("테스트", "시스템", i, {"n": i})
    total = int(db._scalar("SELECT COUNT(*) FROM audit_log"))
    result = retention.archive_audit(years=1, now=datetime.now() + timedelta(days=800))
    assert result["archived"] == total
    assert int(db._scalar("SELECT COUNT(*) FROM audit_log")) == 1        # 이관 기록 1줄만 남음
    chain = db.verify_audit_chain()
    assert chain["broken_id"] is None and chain["archived_until"] == result["last_id"]
    assert retention.verify_archive(1)["ok"]
    # 이관하지 않은 기록·이관 기록은 지울 수 없다
    with pytest.raises(Exception):
        with database.get_conn() as conn:
            conn.execute("DELETE FROM audit_log")
    with pytest.raises(Exception):
        with database.get_conn() as conn:
            conn.execute("DELETE FROM audit_archives")
    db.audit("이관뒤", "시스템", None, None)
    assert db.verify_audit_chain()["broken_id"] is None


def test_backup_generations(tmp_path, monkeypatch):
    folder = tmp_path / "bk"
    folder.mkdir()
    start = datetime(2024, 1, 1, 2, 0, 0)
    for day in range(0, 800, 3):                                       # 2년여 동안 3일마다
        ts = start + timedelta(days=day)
        (folder / f"sales_{ts:%Y%m%d_%H%M%S}_000000.db").write_bytes(b"x")
    values = {"backup_keep_daily": 5, "backup_keep_monthly": 6, "backup_keep_yearly": 3}
    monkeypatch.setattr(company, "get", lambda k: values[k])
    result = retention.prune_backups(folder, (".db",))
    names = sorted(os.listdir(folder))
    assert result["kept"] == len(names) and len(names) <= 5 + 6 + 3
    assert any(n.startswith("sales_2024") for n in names)               # 연말 백업은 오래 남는다


# ── 4. 첨부 ───────────────────────────────────────────────────────────────
def _docx(macro: bool = False) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("[Content_Types].xml", "<Types/>")
        z.writestr("word/document.xml", "<w:document/>")
        if macro:
            z.writestr("word/vbaProject.bin", b"\x00")
    return buf.getvalue()


def test_deal_and_activity_attachments(app):
    rep = login(app, "김영업")
    deal = db._one("SELECT id FROM deals WHERE owner_id=? ORDER BY id LIMIT 1", [user("김영업")["id"]])
    url = f"/attachments/deal/{deal['id']}"
    res = post(rep, url, {"kind": "제안서", "files": (io.BytesIO(_docx()), "제안서.docx")},
               content_type="multipart/form-data", follow_redirects=True)
    assert "첨부했습니다" in res.get_data(as_text=True)
    res = post(rep, url, {"kind": "계약서", "files": (io.BytesIO(_docx(True)), "계약.docm.docx")},
               content_type="multipart/form-data", follow_redirects=True)
    assert "매크로" in res.get_data(as_text=True)
    res = post(rep, url, {"files": (io.BytesIO(b"MZ\x90\x00"), "x.pdf")}, content_type="multipart/form-data",
               follow_redirects=True)
    assert "지원하지 않는" in res.get_data(as_text=True)
    att = db._one("SELECT * FROM attachments WHERE entity='deal' AND entity_id=?", [deal["id"]])
    assert rep.get(f"/attachments/{att['id']}/file").status_code == 200
    assert login(app, "박고객").get(f"/attachments/{att['id']}/file").status_code == 403
    assert "제안서.docx" in rep.get(f"/deals?tab=edit&id={deal['id']}").get_data(as_text=True)
    post(rep, f"/attachments/{att['id']}/void", {"reason": "판 교체"})
    assert db._one("SELECT voided_at FROM attachments WHERE id=?", [att["id"]])["voided_at"]
    cid = db._one("SELECT id FROM customers WHERE owner_id=? ORDER BY id LIMIT 1", [user("김영업")["id"]])["id"]
    post(rep, "/activities/add", {"customer_id": cid, "act_type": db.ACT_TYPES[0], "summary": "첨부 회의",
                                  "attach_kind": "회의록", "files": (io.BytesIO(b"%PDF-1.4 test"), "회의록.pdf")},
         content_type="multipart/form-data")
    act = db._one("SELECT id FROM activities WHERE summary='첨부 회의'")
    assert db._one("SELECT file_name FROM attachments WHERE entity='activity' AND entity_id=?", [act["id"]])


# ── 5. 회계연도 · 분기 ───────────────────────────────────────────────────
def test_fiscal_year_quarters(app, monkeypatch):
    monkeypatch.setattr(fiscal, "start_month", lambda: 4)
    assert fiscal.months(2026)[0] == "2026-04" and fiscal.months(2026)[-1] == "2027-03"
    assert fiscal.fy_of("2027-02") == (2026, 4) and fiscal.fy_of("2026-05") == (2026, 1)
    db.set_context("system", None)
    parts = fiscal.distribute(2026, 2, user("김영업")["id"], 10_000_001)
    assert [p[0] for p in parts] == ["2026-07", "2026-08", "2026-09"] and sum(p[1] for p in parts) == 10_000_001
    table = fiscal.summary(2026)
    row = table[table["담당자"] == "김영업"].iloc[0]
    assert int(row["Q2 목표"]) == 10_000_001
    admin = login(app, "시스템관리자")
    assert "회계연도" in admin.get("/targets?fy=2026").get_data(as_text=True)


# ── 6. 개인정보 열람·삭제 ────────────────────────────────────────────────
def test_privacy_request_export_and_erase(app):
    db.set_context("system", None)
    owner = user("김영업")["id"]
    cid = db.upsert_customer({"name": "정보주체상사", "owner_id": owner, "manager": "홍정보",
                              "phone": "010-7777-8888", "email": "hong@info.kr"})
    db.add_activity({"customer_id": cid, "act_type": db.ACT_TYPES[0], "owner_id": owner,
                     "summary": "홍정보 과장(010-7777-8888) 미팅"})
    found = privacy.search("01077778888")
    assert len(found["customers"]) == 1
    admin = login(app, "시스템관리자")
    res = post(admin, "/admin/privacy", {"term": "홍정보", "requester": "홍정보 본인", "action": "export"})
    wb = load_workbook(io.BytesIO(res.data))
    assert wb["거래처 고객담당자"].max_row >= 2
    aid = db._one("SELECT id FROM activities WHERE customer_id=?", [cid])["id"]
    post(admin, "/admin/privacy", {"term": "홍정보", "requester": "홍정보 본인", "action": "erase", "confirm": "1",
                                   "cid": str(cid), "aid": str(aid)})
    row = db.get_customer(cid)
    assert row["manager"] is None and row["phone"] is None and row["name"] == "정보주체상사"
    summary = db._one("SELECT summary FROM activities WHERE id=?", [aid])["summary"]
    assert "홍정보" not in summary and "010-7777-8888" not in summary and "[파기]" in summary
    assert db._df("SELECT kind FROM pii_requests")["kind"].tolist()[-2:] == ["열람", "삭제"]
    assert "홍정보" not in db._one("SELECT detail FROM audit_log WHERE action='개인정보삭제' ORDER BY id DESC")["detail"]


# ── 7·8. 외화 · 법인 ─────────────────────────────────────────────────────
def test_entities_and_foreign_currency(app):
    db.set_context("system", None)
    eid = ent_mod.upsert({"code": "VN01", "name": "베트남법인", "biz_no": biz("220811999"),
                          "erp_company_code": "3000", "sap_sales_org": "3100"})
    with pytest.raises(ValueError):
        ent_mod.upsert({"code": "VN01", "name": "중복"})
    ent_mod.set_rate("USD", "2026-09-01", 1380)
    rep = login(app, "김영업")
    cid = db.upsert_customer({"name": "수출상사", "owner_id": user("김영업")["id"], "erp_code": "X100"})
    res = post(rep, "/sales/add", {"customer_id": cid, "item": "수출장비", "item_code": "EX-01", "qty": "2", "unit_price": "0",
                                   "status": "입금대기", "tax_type": "영세", "currency": "USD",
                                   "foreign_unit_price": "1500", "entity_id": str(eid), "sale_date": "2026-09-20"})
    assert res.status_code == 302
    sale = db._one("SELECT * FROM sales WHERE item='수출장비'")
    assert (sale["currency"], float(sale["fx_rate"]), int(sale["amount"]), int(sale["vat_amount"])) == \
        ("USD", 1380.0, 2 * 1500 * 1380, 0)
    assert float(sale["foreign_amount"]) == 3000.0 and sale["entity_id"] == eid
    outbox = db._one("SELECT * FROM erp_outbox WHERE ref_id=?", [sale["id"]])
    doc = erp.build_document(outbox, erp.settings())
    assert doc["currency"] == "USD" and doc["foreign_unit_price"] == 1500.0 and doc["company_code"] == "3000"
    payload = erp.sap_sales_order_payload(doc, erp.settings())
    assert payload["SalesOrganization"] == "3100" and payload["TransactionCurrency"] == "USD"
    assert "VN01" in rep.get("/sales").get_data(as_text=True)
    # 환율이 없는 통화는 거부
    res = post(rep, "/sales/add", {"customer_id": cid, "item": "엔화", "qty": "1", "unit_price": "0",
                                   "currency": "JPY", "foreign_unit_price": "100", "status": "입금대기"})
    assert res.status_code == 400 and "환율" in res.get_data(as_text=True)
    with database.get_conn() as conn:
        conn.execute("UPDATE entities SET active=0 WHERE id=?", (eid,))


# ── 9. 전자세금계산서 발행 ───────────────────────────────────────────────
def test_etax_issue_mock_and_file_ack(app, monkeypatch, tmp_path):
    db.set_context("system", None)
    company.save({"company_name": "(주)발행테스트", "company_biz_no": biz("123456789")}, "test")
    owner = user("김영업")["id"]
    cid = db.upsert_customer({"name": "발행상사", "owner_id": owner, "biz_no": biz("314159265"), "erp_code": "E1"})
    sid = db.upsert_sale({"customer_id": cid, "item": "발행품목", "qty": 1, "unit_price": 1_000_000, "owner_id": owner,
                          "sale_date": datetime.now().strftime("%Y-%m-%d")})
    rep = login(app, "김영업")
    res = post(rep, f"/sales/{sid}/etax", {}, follow_redirects=True)
    assert "설정되지 않았습니다" in res.get_data(as_text=True)
    monkeypatch.setenv("SALES_ETAX_ADAPTER", "mock")
    assert post(rep, f"/sales/{sid}/etax", {}).status_code == 302
    jobs.run_pending()
    row = db._one("SELECT * FROM etax_invoices WHERE sale_id=?", [sid])
    assert row["status"] == "발행완료" and len(row["approval_no"]) == 24
    doc = db._one("SELECT * FROM sale_documents WHERE id=?", [row["document_id"]])
    assert doc["doc_type"] == "전자세금계산서" and doc["approval_no"] == row["approval_no"]
    assert int(doc["supply_amount"]) == 1_000_000 and int(doc["tax_amount"]) == 100_000
    res = post(rep, f"/sales/{sid}/etax", {}, follow_redirects=True)
    assert "이미 발행" in res.get_data(as_text=True)
    # 면세 매출은 세금계산서가 아니라 전자계산서(종류 0301, 세액 없음)
    sid2 = db.upsert_sale({"customer_id": cid, "item": "면세품목", "qty": 1, "unit_price": 50_000, "owner_id": owner,
                           "tax_type": "면세", "sale_date": datetime.now().strftime("%Y-%m-%d")})
    etax.request_issue(sid2, None, {"name": "t"})
    jobs.run_pending()
    row2 = db._one("SELECT * FROM etax_invoices WHERE sale_id=?", [sid2])
    doc2 = db._one("SELECT * FROM sale_documents WHERE id=?", [row2["document_id"]])
    assert row2["status"] == "발행완료" and doc2["doc_type"] == "전자계산서" and int(doc2["tax_amount"] or 0) == 0
    assert b"<TypeCode>0301</TypeCode>" in etax.build_xml(db.get_sale(sid2), {}, {}, "2026-10-05")
    # 파일 방식: XML 을 내려놓고, ASP 가 API 로 승인번호를 회신
    monkeypatch.setenv("SALES_ETAX_ADAPTER", "file")
    monkeypatch.setenv("SALES_ETAX_OUT_DIR", str(tmp_path / "etax"))
    sid3 = db.upsert_sale({"customer_id": cid, "item": "파일발행", "qty": 2, "unit_price": 300_000, "owner_id": owner,
                           "sale_date": datetime.now().strftime("%Y-%m-%d")})
    req = etax.request_issue(sid3, None, {"name": "t"})
    jobs.run_pending()
    assert list((tmp_path / "etax").glob("ETAX_*.xml"))
    admin = login(app, "시스템관리자")
    key_page = post(admin, "/admin/api/create", {"name": "ASP회신", "user_id": user("정임원")["id"],
                                                 "scopes": ["erp:write"]}).get_data(as_text=True)
    key = re.search(r"(sk_[0-9a-f]{8}_[A-Za-z0-9_\-]+)", key_page).group(1)
    number = datetime.now().strftime("%Y%m%d") + "41000000" + "12345678"
    res = app.test_client().post("/api/v1/etax/acks", json={"items": [{"request_id": req, "approval_no": number}]},
                                 headers={"Authorization": f"Bearer {key}"})
    assert res.get_json()["summary"] == {"발행완료": 1}
    assert db._one("SELECT status FROM etax_invoices WHERE id=?", [req])["status"] == "발행완료"
    company.save({"company_name": "", "company_biz_no": ""}, "test")
