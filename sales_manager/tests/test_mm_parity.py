"""자재관리와 맞춘 데이터 — 입금 내역·반제, 월 마감·월말 채권, 거래처 원장, ERP 거래처 마스터·거래정지, 직무 분리, 접속 IP."""
from __future__ import annotations

import io
import re
from datetime import date, timedelta

import pytest
from conftest import login, post, user
from openpyxl import load_workbook

from core import database
from core import enterprise as ent
from core import periods
from core import sales_db as db


def _sale(name: str, sale_date: str, unit: int = 1_000_000) -> tuple[int, int]:
    db.set_context("system", None)
    owner = user("김영업")["id"]
    cid = db.upsert_customer({"name": name, "owner_id": owner, "erp_code": "P" + name[-2:]})
    return cid, db.upsert_sale({"customer_id": cid, "item": "원장품목", "qty": 1, "unit_price": unit,
                                "owner_id": owner, "sale_date": sale_date})


def test_payment_history_and_reversal(app):
    _cid, sid = _sale("입금내역상사", date.today().isoformat())          # 합계 1,100,000
    rep = login(app, "김영업")
    post(rep, "/sales/payment", {"sale_id": sid, "amount": "600000", "pay_date": date.today().isoformat(),
                                 "method": "어음", "ref_no": "AB-123", "back": str(sid)})
    pays = ent.list_payments(sid)
    assert len(pays) == 1 and pays[0]["method"] == "어음" and pays[0]["ref_no"] == "AB-123"
    post(rep, f"/payments/{pays[0]['id']}/reverse", {"reason": "다른 거래처 입금을 잘못 넣음"})
    assert int(db.get_sale(sid)["paid_amount"]) == 600_000            # 영업사원은 바로 반제 못 함 (직무 분리)
    post(rep, f"/payments/{pays[0]['id']}/reverse-request", {"reason": "다른 거래처 입금을 잘못 넣음"})
    from core import credit
    req = db._one("SELECT id FROM fin_requests WHERE kind='입금반제' AND payment_id=? AND status='대기'", [pays[0]["id"]])
    credit.decide(int(req["id"]), True, "", user("한팀장"))           # 팀장 승인 → 반제
    sale = db.get_sale(sid)
    assert int(sale["paid_amount"]) == 0 and sale["status"] == "입금대기"
    rows = ent.list_payments(sid)
    assert [int(r["amount"]) for r in rows] == [600_000, -600_000] and rows[0]["reversed_by"]
    with pytest.raises(ValueError, match="이미 반제"):
        ent.reverse_payment(rows[0]["id"], "다시")
    page = rep.get(f"/sales?sid={sid}").get_data(as_text=True)
    assert "입금 내역" in page and "반제됨" in page


def test_payments_always_match_paid_amount(app):
    bad = db._df("SELECT s.id, COALESCE(s.paid_amount, 0) AS paid, COALESCE(SUM(p.amount), 0) AS rows_sum "
                 "FROM sales s LEFT JOIN payments p ON p.sale_id = s.id GROUP BY s.id, s.paid_amount "
                 "HAVING COALESCE(s.paid_amount, 0) <> COALESCE(SUM(p.amount), 0)")
    assert bad.empty, bad.head().to_dict("records")


def test_month_close_blocks_changes_and_snapshots(app, isolated_db, monkeypatch):
    db.set_context("system", None)
    ent.seed_org_demo()
    last_month = (date.today().replace(day=1) - timedelta(days=1))
    cid, sid = _sale("마감상사", last_month.replace(day=5).isoformat())
    ent.record_payment(sid, 300_000, pay_date=last_month.replace(day=20).isoformat())
    admin = {"name": "관리자", "id": 0, "role": "ADMIN"}
    with pytest.raises(ValueError, match="순서대로"):
        periods.close_month("2099-01", admin)
    # 지난달까지 차례로 마감 (ERP 연동이 켜져 있으면 미전송 건 때문에 막힌다)
    monkeypatch.setenv("SALES_ERP_ADAPTER", "none")
    while periods.next_closable() < date.today().strftime("%Y-%m"):
        periods.close_month(periods.next_closable(), admin)
    ym = last_month.strftime("%Y-%m")
    assert periods.closed_through() == ym
    snap = periods.snapshot_table(ym)
    row = snap[snap["거래처"] == "마감상사"].iloc[0]
    assert int(row["월말잔액"]) == 800_000 and int(row["입금누계"]) == 300_000
    with pytest.raises(ValueError, match="마감"):
        db.upsert_sale({**db.get_sale(sid), "id": sid, "memo": "마감 후 수정"})
    with pytest.raises(ValueError, match="마감"):
        db.cancel_sale(sid, "마감 후 취소")
    with pytest.raises(ValueError, match="마감"):
        ent.record_payment(sid, 1000, pay_date=last_month.isoformat())
    ent.record_payment(sid, 1000, pay_date=date.today().isoformat())           # 이번 달 입금은 된다
    with pytest.raises(ValueError, match="사유"):
        periods.reopen("", admin)
    periods.reopen("세금계산서 수정 반영", admin)
    assert periods.closed_through() < ym and periods.snapshot_table(ym).empty
    db.cancel_sale(sid, "해제 후 취소") if False else None


def test_customer_ledger(app):
    rep = login(app, "김영업")
    d1 = date.today().replace(day=1).isoformat()
    cid, s1 = _sale("원장상사", d1)
    db.set_context("system", None)
    s2 = db.upsert_sale({"customer_id": cid, "item": "두번째", "qty": 1, "unit_price": 500_000,
                         "owner_id": user("김영업")["id"], "sale_date": d1})
    ent.record_payment(s1, 1_100_000, pay_date=date.today().isoformat())
    opening, rows = periods.customer_ledger(cid, d1, date.today().isoformat())
    assert opening == 0 and list(rows["구분"]) == ["매출", "매출", "입금"]
    assert int(rows["잔액"].iloc[-1]) == 550_000
    res = rep.get(f"/sales?tab=ledger&customer_id={cid}&from={d1}&to={date.today().isoformat()}&export=ledger")
    assert res.status_code == 200 and load_workbook(io.BytesIO(res.data))["거래처원장"].max_row == 4
    assert login(app, "박고객").get(f"/sales?tab=ledger&customer_id={cid}").status_code == 403
    assert s2


def test_erp_customer_master_and_trade_block(app):
    admin = login(app, "시스템관리자")
    page = post(admin, "/admin/api/create", {"name": "ERP거래처", "user_id": user("정임원")["id"],
                                             "scopes": ["erp:write"]}).get_data(as_text=True)
    key = re.search(r"(sk_[0-9a-f]{8}_[A-Za-z0-9_\-]+)", page).group(1)
    client, h = app.test_client(), {"Authorization": f"Bearer {key}"}
    res = client.post("/api/v1/erp/customers", json={"items": [
        {"erp_code": "BP9001", "name": "ERP마스터상사", "payment_terms": 45, "credit_limit": 70_000_000,
         "owner_emp_no": user("김영업")["emp_no"]},
        {"erp_code": "", "name": "코드없음"}]}, headers=h).get_json()
    assert res["summary"] == {"등록": 1, "오류": 1}
    cust = db._one("SELECT * FROM customers WHERE erp_code='BP9001'")
    assert cust["owner_id"] == user("김영업")["id"] and cust["erp_synced_at"] and int(cust["payment_terms"]) == 45
    rep = login(app, "김영업")
    res = post(rep, "/customers/save", {"id": cust["id"], "name": "이름바꿈", "grade": "A", "industry": "제조",
                                        "erp_code": "BP9001", "payment_terms": "45", "credit_limit": "70000000"},
               follow_redirects=True)
    assert "ERP 에서 고쳐야" in res.get_data(as_text=True)
    assert post(rep, "/customers/save", {"id": cust["id"], "name": "ERP마스터상사", "grade": "VIP", "industry": "제조",
                                         "erp_code": "BP9001", "payment_terms": "45", "credit_limit": "70000000",
                                         "memo": "CRM 메모는 수정 가능"}).status_code == 302
    client.post("/api/v1/erp/customers", json={"items": [{"erp_code": "BP9001", "name": "ERP마스터상사",
                                                          "blocked": True}]}, headers=h)
    res = post(rep, "/sales/add", {"customer_id": cust["id"], "item": "정지중", "qty": "1", "unit_price": "1000",
                                   "status": "입금대기"})
    assert res.status_code == 400 and "거래정지" in res.get_data(as_text=True)


def test_cancel_segregation_and_audit_ip(app):
    _cid, sid = _sale("직무분리상사", date.today().isoformat())
    with database.get_conn() as conn:
        conn.execute("UPDATE sales SET erp_status='전송완료', erp_doc_no='X1', created_by_id=? WHERE id=?",
                     (user("한팀장")["id"], sid))
    rep = login(app, "김영업")
    res = post(rep, f"/sales/{sid}/cancel", {"reason": "취소"}, follow_redirects=True)
    assert "팀장 이상" in res.get_data(as_text=True)
    mgr = login(app, "한팀장")
    res = post(mgr, f"/sales/{sid}/cancel", {"reason": "취소"}, follow_redirects=True)
    assert "본인이 등록한" in res.get_data(as_text=True)
    exec_ = login(app, "정임원")
    post(exec_, f"/sales/{sid}/cancel", {"reason": "계약 해지"})
    assert db.get_sale(sid)["status"] == db.SALE_CANCELLED
    detail = db._one("SELECT detail FROM audit_log WHERE action='취소' AND entity_id=? ORDER BY id DESC", [sid])["detail"]
    assert "접속IP" in detail and "127.0.0.1" in detail
