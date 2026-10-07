"""올린 회사 엑셀 → 같은 모양으로 내려받기: 제목·머리글(3행)·합계 줄이 있는 회사 파일을 올려 모양을 기억하고, 나중에 그 모양으로 받는다."""
from __future__ import annotations

import io
from datetime import date

import pytest
from conftest import login, post, user
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from core import excel_forms as xf
from core import sales_db as db

TODAY = date.today().isoformat()


def _company_file(customer: str) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "매출집계"
    ws["A1"] = "㈜한빛 월 매출집계표"
    ws["A1"].font = Font(bold=True, size=16)
    ws["A2"] = "결재: 담당 / 팀장 / 임원"
    for i, h in enumerate(["일자", "상호", "품명", "수량", "판매단가", "공급가액", "영업담당"], start=1):
        ws.cell(row=3, column=i, value=h).font = Font(bold=True)
    sales_rep = user("김영업")["name"]
    rows = [(TODAY, customer, "양식시험A", 2, 50_000, 100_000, sales_rep), (TODAY, customer, "양식시험B", 1, 30_000, 30_000, sales_rep)]
    for r, row in enumerate(rows, start=4):
        for c, v in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=v)
    ws["A6"] = "합계"
    ws["F6"] = "=SUM(F4:F5)"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_remember_company_layout_and_export_same_shape(app):
    db.set_context("system", None)
    cid = db.upsert_customer({"name": "양식기억상사", "owner_id": user("김영업")["id"]})
    data = _company_file("양식기억상사")
    r = xf.remember_upload("매출", data, "월매출.xlsx", "한빛 매출집계")
    assert r["template"] and r["import_id"] and r["export_id"]
    exp = xf.get_form(r["export_id"])
    assert exp["header_row"] == 3 and {m["field"] for m in exp["column_map"]} >= {"매출일", "거래처", "공급가액"}

    # 서식 파일에는 업로드한 데이터 행이 남아 있지 않고, 제목·머리글·합계 자리표시자만 있다
    ws = load_workbook(io.BytesIO(__import__("core.storage", fromlist=["get_storage"]).get_storage().get(exp["template_key"]))).active
    texts = [str(c.value) for row in ws.iter_rows() for c in row if c.value not in (None, "")]
    assert "㈜한빛 월 매출집계표" in texts and "양식시험A" not in texts and "{{합계:공급가액}}" in texts

    # 같은 양식으로 내려받기: 시스템 데이터가 회사 열 이름·자리에 들어가고 합계가 채워진다
    import pandas as pd
    frame = pd.DataFrame([{"매출일": TODAY, "거래처": "양식기억상사", "품목": "내보내기1", "수량": 3, "단가": 1000,
                           "공급가액": 3000, "담당자": "김영업"},
                          {"매출일": TODAY, "거래처": "양식기억상사", "품목": "내보내기2", "수량": 1, "단가": 7000,
                           "공급가액": 7000, "담당자": "김영업"}])
    out = load_workbook(io.BytesIO(xf.render_export(exp, frame, xf.export_meta(user("시스템관리자"), "시험")))).active
    assert out["A1"].value == "㈜한빛 월 매출집계표" and out["B3"].value == "상호"
    assert [out.cell(row=4, column=c).value for c in (2, 3, 6)] == ["양식기억상사", "내보내기1", 3000]
    assert out["A6"].value == "합계" and out["F6"].value == 10000          # 합계 줄이 데이터 아래로, SUM 대신 합계 값


def test_remember_route_needs_support_role(app, tmp_path):
    db.set_context("system", None)
    db.upsert_customer({"name": "양식권한상사", "owner_id": user("김영업")["id"]})
    data = _company_file("양식권한상사")
    rep, admin = login(app, "김영업"), login(app, "시스템관리자")
    for client in (rep, admin):
        post(client, "/data/import/check", {"entity": "매출", "file": (io.BytesIO(data), "월매출.xlsx")},
             content_type="multipart/form-data")
    assert post(rep, "/data/import/remember", {"name": "권한시험"}).status_code == 403          # 영업사원
    res = post(admin, "/data/import/remember", {"name": "권한시험"})
    assert res.status_code == 302 and res.headers["Location"].endswith("/data/import/recheck")   # 기억한 뒤 같은 파일을 이어서 검증
    assert any(f["name"] == "권한시험 (내려받기)" for f in xf.list_forms("export"))
    page = admin.get("/data/import/recheck")
    text = page.get_data(as_text=True)
    assert page.status_code == 200 and "등록 가능" in text and "권한시험 (업로드)" in text   # 방금 만든 업로드 양식으로 변환해 읽음
