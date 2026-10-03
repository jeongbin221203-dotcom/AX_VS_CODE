"""회사 엑셀 양식 — 각 회사가 이미 쓰는 엑셀 파일을 그대로 올리고, 그대로 받는다.

  업로드(import) 양식
    회사 파일의 시트·머리글 행·열 이름을 시스템 항목에 연결한다(열 매핑). 머리글이 3행에 있어도,
    '상호'·'품명'·'공급가액'처럼 이름이 달라도, 과세구분이 '01' 같은 코드여도 받는다.
    병합 셀(빈 칸)은 위 값을 이어 쓰고(fill_down), '합계'·'소계' 행은 건너뛴다(stop_words).
    변환된 결과는 표준 업로드와 같은 검증(dataio.import_rows)을 거친다 — 양식이 달라도 통제는 같다.

  내려받기(export) 양식
    회사 서식 파일(.xlsx: 로고·제목·결재란·글꼴·테두리)을 올려 두면 데이터만 채워서 돌려준다.
    데이터 시작 행의 서식을 아래 행에 복사하고, 그 아래 내용(합계 줄 등)은 데이터 건수만큼 내린다.
    셀 안의 {{추출일시}} {{추출자}} {{기간}} {{건수}} {{합계:공급가액}} 같은 자리표시자를 값으로 바꾼다.
    서식 파일이 없으면 매핑한 열 이름·순서로 새 파일을 만든다.
"""
from __future__ import annotations

import io
import json
import re
import secrets
from copy import copy
from datetime import datetime
from typing import Any, Optional

import pandas as pd

from . import dataio
from . import sales_db as db
from .storage import get_storage

DIRECTIONS = {"import": "업로드(일괄 등록)", "export": "내려받기(추출)"}
MAX_TEMPLATE_BYTES = 5 * 1024 * 1024
DEFAULT_STOP_WORDS = ["합계", "총계", "소계", "계"]

# 회사 양식에서 자주 쓰는 머리글 → 시스템 항목 (자동 제안용)
SYNONYMS: dict[str, list[str]] = {
    "거래처명": ["상호", "거래처", "고객사", "고객명", "업체명", "회사명", "매출처", "공급받는자상호", "customer"],
    "담당자": ["영업담당", "영업담당자", "담당", "담당자명", "영업사원", "사원명", "사원"],
    "사업자번호": ["사업자등록번호", "등록번호", "공급받는자등록번호", "bizno"],
    "고객담당자": ["고객담당", "구매담당", "담당자고객", "거래처담당자"],
    "연락처": ["전화", "전화번호", "휴대폰", "핸드폰", "tel"],
    "이메일": ["email", "메일", "e-mail"],
    "주소": ["소재지", "사업장주소"],
    "여신한도": ["신용한도", "여신", "한도"],
    "결제조건일": ["결제조건", "지급조건", "결제일수"],
    "ERP코드": ["거래처코드", "고객코드", "sap코드", "erp거래처코드", "bp코드"],
    "매출일": ["일자", "전표일자", "매출일자", "거래일", "거래일자", "작성일자", "출고일"],
    "품목": ["품명", "품목명", "제품명", "상품명", "item"],
    "품목코드": ["품번", "자재코드", "자재번호", "제품코드", "상품코드"],
    "수량": ["qty", "개수", "판매수량"],
    "단가": ["판매단가", "unitprice", "공급단가"],
    "금액": ["공급가액", "공급가", "매출액", "금액원", "판매금액"],
    "과세구분": ["과세유형", "세구분", "부가세구분", "과세여부", "과세코드"],
    "수금상태": ["입금상태", "수금여부", "입금여부"],
    "결제기일": ["수금예정일", "입금예정일", "만기일", "결제예정일"],
    "메모": ["비고", "적요", "remark", "특이사항"],
    "기회명": ["건명", "프로젝트명", "안건명", "영업건명", "딜명"],
    "단계": ["진행단계", "영업단계", "stage"],
    "정가": ["견적가", "리스트가", "정상가"],
    "할인율": ["할인", "dc율", "dc"],
    "예상금액": ["제안금액", "제안가", "예상매출"],
    "예상마감일": ["마감예정일", "수주예정일", "계약예정일"],
    "예측구분": ["forecast", "포캐스트"],
    "유입경로": ["리드출처", "유입"],
    "경쟁사": ["경쟁업체"],
    "활동일": ["방문일", "상담일", "활동일자"],
    "활동내용": ["내용", "상담내용", "방문내용"],
    "유형": ["활동유형", "구분"],
    "관련기회": ["관련건명", "관련안건"],
    "다음액션": ["후속조치", "다음조치"],
    "다음일정": ["후속일정", "다음방문일"],
    "월": ["기준월", "년월", "연월"],
    "목표금액": ["목표", "매출목표", "목표매출"],
}


def _norm(text: Any) -> str:
    return re.sub(r"[\s()\[\]{}·_./\-:*※]", "", str(text or "")).lower()


# ---------------------------------------------------------------------------
# 조회 · 저장
# ---------------------------------------------------------------------------
def _decode(row: Optional[dict]) -> Optional[dict]:
    if not row:
        return None
    row = dict(row)
    row["column_map"] = json.loads(row.get("column_map") or "[]")
    row["value_map"] = json.loads(row.get("value_map") or "{}")
    row["fill_down"] = json.loads(row.get("fill_down") or "[]")
    row["stop_words"] = json.loads(row["stop_words"]) if row.get("stop_words") else list(DEFAULT_STOP_WORDS)
    return row


def list_forms(direction: str | None = None, entity: str | None = None, active_only: bool = True) -> list[dict]:
    sql, params = "SELECT * FROM excel_forms WHERE 1=1", []
    if direction:
        sql += " AND direction=?"
        params.append(direction)
    if entity:
        sql += " AND entity=?"
        params.append(entity)
    if active_only:
        sql += " AND active=1"
    return [_decode(r) for r in db._df(sql + " ORDER BY entity, name", params).to_dict("records")]


def table() -> pd.DataFrame:
    rows = [{"id": f["id"], "양식명": f["name"], "구분": DIRECTIONS.get(f["direction"], f["direction"]),
             "항목": f["entity"], "시트": f["sheet_name"] or "(첫 시트)", "머리글 행": f["header_row"],
             "열 매핑": len(f["column_map"]), "서식 파일": f["template_name"] or "-", "수정일": f["updated_at"]}
            for f in list_forms()]
    return pd.DataFrame(rows)


def get_form(form_id: int) -> Optional[dict]:
    return _decode(db._one("SELECT * FROM excel_forms WHERE id=?", [int(form_id)]))


def system_fields(direction: str, entity: str) -> list[str]:
    """매핑할 수 있는 시스템 항목."""
    if direction == "import":
        spec = dataio.IMPORT_SPECS[entity]
        return spec["required"] + spec["optional"]
    frame = dataio.collect([entity]).get(entity, pd.DataFrame())
    return [c for c in frame.columns]


def _validate_template(data: bytes, filename: str) -> None:
    if not filename.lower().endswith(".xlsx"):
        raise ValueError("서식 파일은 .xlsx 만 받습니다 (매크로가 든 .xlsm·옛 .xls 는 받지 않습니다).")
    if len(data) > MAX_TEMPLATE_BYTES:
        raise ValueError("서식 파일은 5MB 이하여야 합니다.")
    from openpyxl import load_workbook
    try:
        load_workbook(io.BytesIO(data))
    except Exception as exc:   # noqa: BLE001
        raise ValueError(f"엑셀 파일을 열 수 없습니다: {exc}") from exc


def save_form(data: dict, template: bytes | None = None, template_name: str | None = None) -> int:
    name = str(data.get("name") or "").strip()
    direction, entity = data.get("direction"), data.get("entity")
    if not name:
        raise ValueError("양식 이름을 입력하세요 (예: 본사 매출집계표).")
    if direction not in DIRECTIONS:
        raise ValueError("양식 구분을 고르세요.")
    valid_entities = list(dataio.IMPORT_SPECS) if direction == "import" else dataio.EXPORT_SOURCES
    if entity not in valid_entities:
        raise ValueError(f"항목이 올바르지 않습니다: {entity}")
    header_row = int(data.get("header_row") or 1)
    if not 1 <= header_row <= 100:
        raise ValueError("머리글 행은 1~100 사이여야 합니다.")
    start = int(data.get("data_start_row") or header_row + 1)
    if start <= header_row:
        raise ValueError("데이터 시작 행은 머리글 행보다 아래여야 합니다.")
    mapping = [m for m in (data.get("column_map") or []) if m.get("column") and m.get("field")]
    fields = system_fields(direction, entity)
    bad = [m["field"] for m in mapping if m["field"] not in fields]
    if bad:
        raise ValueError(f"시스템에 없는 항목입니다: {', '.join(bad)}")
    if not mapping:
        raise ValueError("열 매핑을 한 개 이상 지정하세요.")
    dup = {m["column"] for m in mapping if sum(1 for x in mapping if x["column"] == m["column"]) > 1}
    if dup:
        raise ValueError(f"같은 회사 열이 두 번 매핑되었습니다: {', '.join(dup)}")
    if direction == "import":
        mapped = {m["field"] for m in mapping}
        missing = [f for f in dataio.IMPORT_SPECS[entity]["required"] if f not in mapped]
        if missing:
            raise ValueError(f"필수 항목이 매핑되지 않았습니다: {', '.join(missing)}")
    value_map = data.get("value_map") or {}
    if isinstance(value_map, str):
        value_map = parse_value_map(value_map)
    fill_down = [f for f in (data.get("fill_down") or []) if f in fields]
    stop_words = data.get("stop_words")
    if isinstance(stop_words, str):
        stop_words = [w.strip() for w in stop_words.split(",") if w.strip()]

    prev = get_form(int(data["id"])) if data.get("id") else None
    key = prev["template_key"] if prev else None
    if template:
        _validate_template(template, template_name or "")
        key = f"forms/{secrets.token_hex(8)}.xlsx"
        get_storage().put(key, template, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    dup_name = db._one("SELECT id FROM excel_forms WHERE name=? AND id<>?", [name, int(data.get("id") or 0)])
    if dup_name:
        raise ValueError(f"'{name}' 양식이 이미 있습니다.")
    record = {"name": name, "direction": direction, "entity": entity,
              "sheet_name": (data.get("sheet_name") or "").strip() or None, "header_row": header_row,
              "data_start_row": start, "column_map": json.dumps(mapping, ensure_ascii=False),
              "value_map": json.dumps(value_map, ensure_ascii=False),
              "fill_down": json.dumps(fill_down, ensure_ascii=False),
              "stop_words": json.dumps(stop_words if stop_words is not None else DEFAULT_STOP_WORDS,
                                       ensure_ascii=False),
              "template_key": key, "template_name": template_name if template else (prev or {}).get("template_name"),
              "memo": data.get("memo")}
    cols = list(record)
    with db.get_conn() as conn:
        if prev:
            conn.execute(f"UPDATE excel_forms SET {', '.join(f'{c}=?' for c in cols)}, updated_at=? WHERE id=?",
                         (*record.values(), db._now(), prev["id"]))
            fid = int(prev["id"])
        else:
            cur = conn.execute(f"INSERT INTO excel_forms ({', '.join(cols)}, created_by, created_at, updated_at) "
                               f"VALUES ({', '.join('?' * len(cols))}, ?, ?, ?)",
                               (*record.values(), db.current_actor(), db._now(), db._now()))
            fid = int(cur.lastrowid)
    if template and prev and prev.get("template_key") and prev["template_key"] != key:
        get_storage().delete(prev["template_key"])
    db.audit("수정" if prev else "등록", "시스템", fid,
             {"엑셀양식": name, "구분": direction, "항목": entity, "열": len(mapping), "서식파일": bool(key)})
    return fid


def deactivate(form_id: int) -> None:
    with db.get_conn() as conn:
        conn.execute("UPDATE excel_forms SET active=0, updated_at=? WHERE id=?", (db._now(), form_id))
    db.audit("삭제", "시스템", form_id, {"엑셀양식": "사용 중지"})


def parse_value_map(text: str) -> dict:
    """'과세구분: 01=과세, 02=영세, 03=면세' 줄 단위 입력을 {항목: {원래값: 바꿀값}} 으로."""
    out: dict[str, dict[str, str]] = {}
    for line in (text or "").splitlines():
        if not line.strip():
            continue
        if ":" not in line:
            raise ValueError(f"값 변환 형식: '항목: 원래값=바꿀값, …' ({line.strip()})")
        field, pairs = line.split(":", 1)
        for pair in pairs.split(","):
            if pair.strip():
                if "=" not in pair:
                    raise ValueError(f"값 변환 형식: '원래값=바꿀값' ({pair.strip()})")
                src, dst = pair.split("=", 1)
                out.setdefault(field.strip(), {})[src.strip()] = dst.strip()
    return out


def value_map_text(value_map: dict) -> str:
    return "\n".join(f"{f}: " + ", ".join(f"{s}={d}" for s, d in m.items()) for f, m in (value_map or {}).items())


# ---------------------------------------------------------------------------
# 양식 등록 도우미: 샘플 파일에서 머리글을 읽고 매핑을 제안한다
# ---------------------------------------------------------------------------
def inspect_file(data: bytes, filename: str, sheet: str | None = None, header_row: int | None = None) -> dict:
    """시트 목록, (지정 없으면 추정한) 머리글 행, 머리글 목록, 앞부분 미리보기."""
    name = filename.lower()
    if name.endswith(".csv"):
        frame = dataio.read_upload(io.BytesIO(data), filename)
        raw = pd.concat([pd.DataFrame([list(frame.columns)]), pd.DataFrame(frame.values)], ignore_index=True)
        sheets, sheet = ["(CSV)"], None
    else:
        book = pd.ExcelFile(io.BytesIO(data))
        sheets = book.sheet_names
        sheet = sheet if sheet in sheets else sheets[0]
        raw = pd.read_excel(book, sheet_name=sheet, header=None, dtype=object, nrows=60)
    if not header_row:
        header_row = _guess_header_row(raw)
    headers = [str(v).strip() for v in raw.iloc[header_row - 1].tolist()] if len(raw) >= header_row else []
    headers = [h for h in headers if h and h.lower() != "nan"]
    from openpyxl.utils import get_column_letter
    preview = raw.head(max(header_row + 5, 8)).fillna("").copy()
    preview.columns = [get_column_letter(i + 1) for i in range(len(preview.columns))]
    preview.insert(0, "행", range(1, len(preview) + 1))
    return {"sheets": sheets, "sheet": sheet, "header_row": header_row, "headers": headers, "preview": preview}


def _guess_header_row(raw: pd.DataFrame) -> int:
    """글자 칸이 가장 많고, 시스템 항목 이름과 가장 많이 맞는 행을 머리글로 본다(제목·결재란 건너뛰기)."""
    known = {_norm(k) for k in SYNONYMS} | {_norm(s) for v in SYNONYMS.values() for s in v}
    best, best_score = 1, -1.0
    for i in range(min(len(raw), 30)):
        cells = [c for c in raw.iloc[i].tolist() if isinstance(c, str) and c.strip()]
        score = len(cells) + 3 * sum(1 for c in cells if _norm(c) in known)
        if score > best_score:
            best, best_score = i + 1, score
    return best


def suggest_mapping(headers: list[str], fields: list[str]) -> list[dict]:
    lookup: dict[str, str] = {}
    for field in fields:
        lookup.setdefault(_norm(field), field)
        for syn in SYNONYMS.get(field, []):
            lookup.setdefault(_norm(syn), field)
    used: set[str] = set()
    out = []
    for h in headers:
        key = _norm(h)
        field = lookup.get(key) or next((f for k, f in lookup.items() if k and (k in key or key in k)
                                         and len(key) >= 2), "")
        if field in used:
            field = ""
        used.add(field)
        out.append({"column": h, "field": field})
    return out


# ---------------------------------------------------------------------------
# 업로드: 회사 파일 → 표준 컬럼
# ---------------------------------------------------------------------------
def to_standard(form: dict, data: bytes, filename: str) -> tuple[pd.DataFrame, int]:
    """회사 양식 파일을 표준 업로드 컬럼으로 바꾼다. (DataFrame, 엑셀 행 번호 보정값)."""
    header = int(form["header_row"])
    if filename.lower().endswith(".csv"):
        frame = dataio.read_upload(io.BytesIO(data), filename)
        header = 1
    else:
        book = pd.ExcelFile(io.BytesIO(data))
        sheet = form["sheet_name"] if form.get("sheet_name") in book.sheet_names else book.sheet_names[0]
        if form.get("sheet_name") and form["sheet_name"] not in book.sheet_names:
            raise ValueError(f"'{form['sheet_name']}' 시트가 없습니다 (파일의 시트: {', '.join(book.sheet_names)}).")
        frame = pd.read_excel(book, sheet_name=sheet, header=header - 1, dtype=object)
    frame.columns = [str(c).strip() for c in frame.columns]
    by_norm = {_norm(c): c for c in frame.columns}
    mapping = {}
    missing = []
    for m in form["column_map"]:
        col = m["column"] if m["column"] in frame.columns else by_norm.get(_norm(m["column"]))
        if col is None:
            missing.append(m["column"])
        else:
            mapping[col] = m["field"]
    if missing:
        raise ValueError(f"'{form['name']}' 양식의 열을 찾지 못했습니다: {', '.join(missing)} "
                         f"(머리글 행 {header}행 · 시트 이름을 확인하세요)")
    out = frame[list(mapping)].rename(columns=mapping)
    out = out.dropna(how="all")
    stop = {_norm(w) for w in form.get("stop_words") or []}
    if stop and len(out.columns):
        first_text = out.apply(lambda r: next((_norm(v) for v in r.tolist()
                                               if isinstance(v, str) and v.strip()), ""), axis=1)
        out = out[~first_text.isin(stop)]
    for field in form.get("fill_down") or []:
        if field in out.columns:
            out[field] = out[field].ffill()
    for field, table_ in (form.get("value_map") or {}).items():
        if field in out.columns:
            out[field] = out[field].map(lambda v, t=table_: t.get(str(v).strip(), v) if pd.notna(v) else v)
    return out, header + 1


def blank_template(form: dict) -> bytes:
    """업로드용 빈 양식: 회사 서식 파일이 있으면 그대로, 없으면 머리글 행 위치에 회사 열 이름을 둔다."""
    if form.get("template_key"):
        return get_storage().get(form["template_key"])
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = (form.get("sheet_name") or form["entity"])[:31]
    if form["header_row"] > 1:
        ws.cell(row=1, column=1, value=form["name"]).font = Font(bold=True, size=14)
    for i, m in enumerate(form["column_map"], start=1):
        cell = ws.cell(row=form["header_row"], column=i, value=m["column"])
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDEBF7")
        ws.column_dimensions[cell.column_letter].width = max(12, len(m["column"]) * 2 + 2)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# 내려받기: 표준 데이터 → 회사 서식
# ---------------------------------------------------------------------------
_PLACEHOLDER = re.compile(r"\{\{\s*([^}]+?)\s*\}\}")


def render_export(form: dict, frame: pd.DataFrame, meta: dict) -> bytes:
    columns = [m for m in form["column_map"] if m["field"] in frame.columns]
    body = dataio.neutralize_formulas(frame[[m["field"] for m in columns]]) if columns else pd.DataFrame()
    body.columns = [m["column"] for m in columns]
    if not form.get("template_key"):
        return dataio.to_excel({form.get("sheet_name") or form["entity"]: body}, meta)

    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(get_storage().get(form["template_key"])))
    ws = wb[form["sheet_name"]] if form.get("sheet_name") in wb.sheetnames else wb.worksheets[0]
    header_row, start = int(form["header_row"]), int(form.get("data_start_row") or form["header_row"] + 1)

    # 회사 머리글 위치 찾기 (없으면 오른쪽 빈 열에 머리글을 새로 쓴다)
    positions: dict[str, int] = {}
    for cell in ws[header_row]:
        if cell.value is not None and str(cell.value).strip():
            positions.setdefault(_norm(cell.value), cell.column)
    next_col = max([c.column for c in ws[header_row] if c.value is not None] or [0]) + 1
    targets = []
    for m in columns:
        col = positions.get(_norm(m["column"]))
        if col is None:
            col = next_col
            ws.cell(row=header_row, column=col, value=m["column"])
            next_col += 1
        targets.append(col)

    # 데이터 시작 행 아래 내용(합계 줄 등)은 데이터 건수만큼 내린다 — 수식의 행 번호도 함께 옮긴다
    n = len(body)
    if n > 1 and ws.max_row > start:
        ws.move_range(f"A{start + 1}:{ws.cell(row=ws.max_row, column=ws.max_column).coordinate}",
                      rows=n - 1, translate=True)
    styles = {col: ws.cell(row=start, column=col) for col in targets}
    style_attrs = ("font", "border", "fill", "number_format", "alignment", "protection")
    for r, row in enumerate(body.itertuples(index=False), start=start):
        for col, value in zip(targets, row):
            cell = ws.cell(row=r, column=col)
            cell.value = None if (not isinstance(value, str) and pd.isna(value)) else (
                value.item() if hasattr(value, "item") else value)
            if r != start:
                src = styles[col]
                for attr in style_attrs:
                    setattr(cell, attr, copy(getattr(src, attr)))
            if cell.number_format == "General" and isinstance(cell.value, (int, float)) and \
                    any(h in str(body.columns[targets.index(col)]) for h in dataio.MONEY_HINTS):
                cell.number_format = "#,##0"

    values = {**{k: str(v) for k, v in meta.items()}, "건수": f"{n:,}"}
    for m in columns:
        if pd.api.types.is_numeric_dtype(frame[m["field"]]):
            total = frame[m["field"]].sum()
            total = total.item() if hasattr(total, "item") else total      # numpy → 파이썬 숫자 (엑셀 숫자 셀)
            values[f"합계:{m['field']}"] = values[f"합계:{m['column']}"] = total
    for sheet in wb.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and "{{" in cell.value:
                    whole = _PLACEHOLDER.fullmatch(cell.value.strip())
                    if whole and isinstance(values.get(whole.group(1)), (int, float)):
                        cell.value = values[whole.group(1)]
                        cell.number_format = "#,##0"
                    else:
                        cell.value = _PLACEHOLDER.sub(lambda mt: str(values.get(mt.group(1), mt.group(0))), cell.value)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_meta(user: dict, period: str) -> dict:
    return {"추출일시": datetime.now().strftime("%Y-%m-%d %H:%M"), "추출자": user.get("name", ""),
            "기간": period, "작성일": datetime.now().strftime("%Y-%m-%d")}
