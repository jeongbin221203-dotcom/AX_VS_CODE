"""영업관리 시스템 - 데이터 일괄 등록 / 추출

대량 데이터를 다루는 실무 원칙을 그대로 구현한다.
  * 등록 전 반드시 검증만 먼저 돌려 볼 수 있다 (dry-run)
  * 오류는 '행 번호 + 사유' 로 돌려준다 (엑셀에서 바로 찾아 고칠 수 있도록)
  * 본인 접근범위 밖의 담당자로는 등록할 수 없다 (권한 우회 차단)
  * 담당자는 사번 또는 이름으로 받되 등록된 사용자로 확정한다 (동명이인은 사번 필수)
  * 모든 일괄 작업은 감사로그에 건수와 함께 남는다
  * 추출 파일은 엑셀 수식 주입을 막고, 연락처·이메일은 기본으로 가린다

UI 의존성이 없으므로 배치 스크립트에서도 그대로 쓸 수 있다.
"""
from __future__ import annotations

import io
import re
from datetime import date, datetime, timedelta
from typing import Any, Callable, Optional

import pandas as pd

from . import enterprise as ent
from . import sales_db as db

# ============================================================================
# 공통 파서 / 검증 유틸
# ============================================================================
def _clean(value: Any) -> str:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return ""
    return str(value).strip()


def parse_date(value: Any, field: str, required: bool = False) -> Optional[str]:
    """'2026-01-05', '2026/1/5', '20260105', 엑셀 날짜셀을 모두 받아들인다."""
    raw = _clean(value)
    if not raw:
        if required:
            raise ValueError(f"{field}: 필수 항목입니다")
        return None
    if isinstance(value, (datetime, pd.Timestamp)):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    text = raw.replace(".", "-").replace("/", "-").split(" ")[0].strip("-")
    if re.fullmatch(r"\d{5}(\.0)?", raw):                      # 엑셀 날짜 일련번호 (46299 = 2026-10-04)
        return (date(1899, 12, 30) + timedelta(days=int(float(raw)))).strftime("%Y-%m-%d")
    if re.fullmatch(r"\d{8}", text):
        text = f"{text[:4]}-{text[4:6]}-{text[6:]}"
    try:
        parts = [int(p) for p in text.split("-")[:3]]
        if len(parts) < 3:
            raise ValueError
        if parts[0] < 100:                                       # '26.10.04' → 2026-10-04
            parts[0] += 2000
        if not 2000 <= parts[0] <= 2100:
            raise ValueError
        return date(parts[0], parts[1], parts[2]).strftime("%Y-%m-%d")
    except (ValueError, IndexError) as exc:
        raise ValueError(f"{field}: 날짜 형식이 아닙니다 ('{raw}') → 예: 2026-01-05") from exc


# 실무 엑셀에 흔한 한글 금액 표기 ('3억', '5천만', '1억2천만')를 숫자로 바꾼다
KOREAN_UNITS = [("억", 100_000_000), ("천만", 10_000_000), ("백만", 1_000_000),
                ("십만", 100_000), ("만", 10_000), ("천", 1_000)]


def _korean_number(text: str) -> Optional[int]:
    """'1억2천만' → 120000000. 해석할 수 없으면 None."""
    remain, total, matched = text, 0, False
    while remain:
        match = re.match(r"\s*([\d.]+)\s*(억|천만|백만|십만|만|천)?", remain)
        if not match or not match.group(1):
            return None
        number = float(match.group(1))
        unit = match.group(2)
        total += number * dict(KOREAN_UNITS)[unit] if unit else number
        matched = True
        remain = remain[match.end():]
    return int(round(total)) if matched else None


def parse_int(value: Any, field: str, default: int = 0, minimum: int | None = None) -> int:
    raw = _clean(value).replace(",", "").replace("원", "").replace(" ", "")
    if not raw:
        return default
    try:
        number = int(round(float(raw)))
    except OverflowError:
        raise ValueError(f"{field}: 숫자가 너무 큽니다 ('{raw[:30]}')") from None
    except ValueError:
        number = _korean_number(raw)
        if number is None:
            raise ValueError(f"{field}: 숫자가 아닙니다 ('{raw[:30]}')")
    if abs(number) > 10 ** 15:
        raise ValueError(f"{field}: 숫자가 너무 큽니다 ('{raw[:30]}')")
    if minimum is not None and number < minimum:
        raise ValueError(f"{field}: {minimum} 이상이어야 합니다 (입력값 {number})")
    return number


def parse_float(value: Any, field: str, default: float = 0.0) -> float:
    raw = _clean(value).replace(",", "").replace("%", "")
    if not raw:
        return default
    try:
        number = float(raw)
    except ValueError as exc:
        raise ValueError(f"{field}: 숫자가 아닙니다 ('{raw[:30]}')") from exc
    if number != number or abs(number) > 10 ** 15:
        raise ValueError(f"{field}: 숫자가 올바르지 않습니다 ('{raw[:30]}')")
    return number


def parse_choice(value: Any, field: str, choices: list[str], default: str | None = None,
                 required: bool = False) -> Optional[str]:
    raw = _clean(value)
    if not raw:
        if required:
            raise ValueError(f"{field}: 필수 항목입니다")
        return default
    for choice in choices:                      # 공백/대소문자 차이는 허용
        if raw.replace(" ", "").lower() == choice.replace(" ", "").lower():
            return choice
    raise ValueError(f"{field}: 허용되지 않는 값입니다 ('{raw}') → {', '.join(choices)}")


def parse_ym(value: Any, field: str) -> str:
    raw = _clean(value).replace(".", "-").replace("/", "-")
    if re.fullmatch(r"\d{6}", raw):
        raw = f"{raw[:4]}-{raw[4:]}"
    if not re.fullmatch(r"\d{4}-\d{2}", raw):
        raise ValueError(f"{field}: 'YYYY-MM' 형식이어야 합니다 ('{_clean(value)}')")
    year, month = int(raw[:4]), int(raw[5:])
    if not 1 <= month <= 12:
        raise ValueError(f"{field}: 월은 01~12 여야 합니다 ('{_clean(value)}')")
    if not 2000 <= year <= 2100:
        raise ValueError(f"{field}: 연도가 올바르지 않습니다 ('{_clean(value)}')")
    return raw


# ============================================================================
# 업로드 양식 정의
# ============================================================================
IMPORT_SPECS: dict[str, dict] = {
    "거래처": {
        "required": ["거래처명", "담당자"],
        "optional": ["등급", "업종", "사업자번호", "고객담당자", "연락처", "이메일",
                     "주소", "여신한도", "결제조건일", "ERP코드", "메모"],
        "sample": {"거래처명": "예시상사", "담당자": "김영업", "등급": "A", "업종": "제조",
                   "사업자번호": "123-45-67890", "고객담당자": "홍길동 과장",
                   "연락처": "010-1234-5678", "이메일": "hong@example.co.kr",
                   "주소": "서울시 강남구", "여신한도": 300000000, "결제조건일": 30,
                   "ERP코드": "C100234", "메모": "신규 발굴"},
        "key": "거래처명",
    },
    "영업기회": {
        "required": ["거래처명", "기회명", "담당자"],
        "optional": ["단계", "정가", "할인율", "예상금액", "예상마감일", "예측구분",
                     "유입경로", "경쟁사", "메모"],
        "sample": {"거래처명": "예시상사", "기회명": "ERP 라이선스 갱신", "담당자": "김영업",
                   "단계": db.OPEN_STAGES[2], "정가": 50000000, "할인율": 5, "예상금액": "",
                   "예상마감일": "2026-03-31", "예측구분": "Best Case", "유입경로": "기존고객",
                   "경쟁사": "A社", "메모": ""},
        "key": "기회명",
    },
    "영업활동": {
        "required": ["거래처명", "활동일", "담당자", "활동내용"],
        "optional": ["유형", "관련기회", "다음액션", "다음일정"],
        "sample": {"거래처명": "예시상사", "활동일": "2026-02-10", "담당자": "김영업",
                   "활동내용": "구매팀 미팅, 도입 일정 협의", "유형": "방문",
                   "관련기회": "ERP 라이선스 갱신", "다음액션": "수정 견적 발송",
                   "다음일정": "2026-02-17"},
        "key": None,
    },
    "매출": {
        "required": ["거래처명", "매출일", "품목", "담당자"],
        "optional": ["품목코드", "수량", "단가", "금액", "과세구분", "수금상태", "결제기일", "메모"],
        "sample": {"거래처명": "예시상사", "매출일": "2026-02-28", "품목": "ERP 라이선스",
                   "품목코드": "SW-ERP-01", "담당자": "김영업", "수량": 10, "단가": 3000000, "금액": "",
                   "과세구분": "과세", "수금상태": "입금대기", "결제기일": "2026-03-31", "메모": ""},
        "key": None,
    },
    "목표": {
        "required": ["월", "담당자", "목표금액"],
        "optional": [],
        "sample": {"월": "2026-03", "담당자": "김영업", "목표금액": 80000000},
        "key": "월+담당자",
    },
}


def template_df(entity: str) -> pd.DataFrame:
    """업로드 양식(헤더 + 예시 1행)을 만든다."""
    spec = IMPORT_SPECS[entity]
    columns = spec["required"] + spec["optional"]
    return pd.DataFrame([{c: spec["sample"].get(c, "") for c in columns}], columns=columns)


MAX_UPLOAD_ROWS = 20_000              # 한 번에 올리는 행 수 상한 (넘으면 나눠 올리게)
MAX_XLSX_UNPACKED = 200 * 1024 * 1024  # 엑셀(zip) 풀었을 때 크기 상한 — 작은 파일이 수 GB 로 풀리는 압축 폭탄 차단


def check_excel_safe(data: bytes) -> None:
    """xlsx 는 zip 이다: 풀린 크기·압축률을 먼저 본다 (pandas 가 통째로 풀기 전에)."""
    import zipfile
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            unpacked = sum(i.file_size for i in zf.infolist())
    except zipfile.BadZipFile:
        return                                         # .xls(옛 형식) 등은 그대로 — pandas 가 판단
    if unpacked > MAX_XLSX_UNPACKED or (len(data) and unpacked / len(data) > 200):
        raise ValueError("엑셀 파일이 너무 큽니다(풀면 200MB 초과). 행을 나눠 올려 주세요.")


def read_upload(file_obj, filename: str = "") -> pd.DataFrame:
    """CSV / Excel 업로드 파일을 DataFrame 으로 읽는다(한글 인코딩 자동 처리)."""
    name = (filename or getattr(file_obj, "name", "")).lower()
    if name.endswith((".xlsx", ".xlsm", ".xls")):
        data = file_obj.read() if hasattr(file_obj, "read") else open(file_obj, "rb").read()
        check_excel_safe(data)
        df = pd.read_excel(io.BytesIO(data), dtype=object, nrows=MAX_UPLOAD_ROWS + 1)
        if len(df) > MAX_UPLOAD_ROWS:
            raise ValueError(f"한 번에 {MAX_UPLOAD_ROWS:,}행까지 올릴 수 있습니다. 나눠 올려 주세요.")
        return df
    for encoding in ("utf-8-sig", "cp949", "euc-kr", "utf-8"):
        try:
            if hasattr(file_obj, "seek"):
                file_obj.seek(0)
            return pd.read_csv(file_obj, dtype=object, encoding=encoding)
        except (UnicodeDecodeError, LookupError):
            continue
    raise ValueError("CSV 인코딩을 인식하지 못했습니다. UTF-8 또는 CP949(한글 Windows)로 저장해 주세요.")


# ============================================================================
# 행 → 레코드 변환기 (엔터티별)
# ============================================================================
def _customer_id_by_name(name: str, cache: dict, db_path: str | None) -> int:
    """거래처명 → id. 정식 이름 · 다른 이름 · '(주)'·띄어쓰기를 뺀 이름 순으로 찾는다 (core/customer_names.py)."""
    from . import customer_names as cn
    key = name.strip()
    if key in cache:
        return cache[key]
    if "_name_index" not in cache:
        cache["_name_index"] = cn.Index(db_path)
    idx = cache["_name_index"]
    cid, _how = idx.resolve(key)
    if cid is None:
        cn.record_unknown(key, "엑셀 업로드", db_path)
        hint = idx.suggest(key)
        raise ValueError(f"거래처명: 등록되지 않은 거래처입니다 ('{key}')"
                         + (f" — 비슷한 거래처: {', '.join(n for _i, n in hint)}" if hint else "")
                         + " → 거래처 > 🏷️ 이름 정리에서 기존 거래처에 연결하거나 먼저 등록하세요")
    cache[key] = cid
    return cache[key]


def _owner(row: pd.Series, ctx: dict) -> dict:
    """담당자 칸(사번 또는 이름) → 등록된 사용자. 권한 범위 밖이면 여기서 오류가 난다."""
    raw = _clean(row.get("담당자"))
    if not raw:
        raise ValueError("담당자: 필수 항목입니다")
    if raw not in ctx["owner_cache"]:
        try:
            ctx["owner_cache"][raw] = db.resolve_owner(raw, ctx["db_path"])
        except ValueError as exc:
            ctx["owner_cache"][raw] = exc
    found = ctx["owner_cache"][raw]
    if isinstance(found, Exception):
        raise ValueError(f"담당자: {found}")
    return {"owner_id": found[0], "owner": found[1]}


CUSTOMER_COLUMNS = {"grade": "등급", "industry": "업종", "biz_no": "사업자번호", "manager": "고객담당자",
                    "phone": "연락처", "email": "이메일", "address": "주소", "credit_limit": "여신한도",
                    "payment_terms": "결제조건일", "erp_code": "ERP코드", "memo": "메모"}


def _row_customer(row: pd.Series, ctx: dict) -> tuple[str, dict]:
    name = _clean(row.get("거래처명"))
    if not name:
        raise ValueError("거래처명: 필수 항목입니다")
    blank = [f for f, col in CUSTOMER_COLUMNS.items() if not _clean(row.get(col))]
    data = {
        "name": name,
        "grade": parse_choice(row.get("등급"), "등급", db.GRADES, "B"),
        "industry": parse_choice(row.get("업종"), "업종", db.INDUSTRIES, "기타"),
        "biz_no": _clean(row.get("사업자번호")) or None,
        "manager": _clean(row.get("고객담당자")) or None,
        "phone": _clean(row.get("연락처")) or None,
        "email": _clean(row.get("이메일")) or None,
        "address": _clean(row.get("주소")) or None,
        "credit_limit": parse_int(row.get("여신한도"), "여신한도", 0, 0),
        "payment_terms": parse_int(row.get("결제조건일"), "결제조건일", 30, 0),
        "erp_code": _clean(row.get("ERP코드")) or None,
        "memo": _clean(row.get("메모")) or None,
        "status": "활성",
        **_owner(row, ctx),
        "_blank": blank,
    }
    return data["owner"], data


def _row_deal(row: pd.Series, ctx: dict) -> tuple[str, dict]:
    cid = _customer_id_by_name(_clean(row.get("거래처명")), ctx["cust_cache"], ctx["db_path"])
    title = _clean(row.get("기회명"))
    if not title:
        raise ValueError("기회명: 필수 항목입니다")
    stage = parse_choice(row.get("단계"), "단계", db.STAGES, db.OPEN_STAGES[0])
    list_amount = parse_int(row.get("정가"), "정가", 0, 0)
    discount = parse_float(row.get("할인율"), "할인율", 0.0)
    amount = parse_int(row.get("예상금액"), "예상금액", 0, 0)
    if not amount:
        amount = int(list_amount * (100 - discount) / 100)
    if not list_amount:
        list_amount = amount
    if stage == db.STAGE_LOST and not _clean(row.get("메모")):
        pass                                   # 실주사유는 선택 항목으로 둔다(이관 데이터 고려)
    data = {
        "customer_id": cid, "title": title, **_owner(row, ctx),
        "stage": stage, "amount": amount, "list_amount": list_amount,
        "discount_rate": discount, "probability": db.STAGE_PROB[stage],
        "expected_close": parse_date(row.get("예상마감일"), "예상마감일"),
        "forecast_category": parse_choice(row.get("예측구분"), "예측구분", db.FORECAST_CATS, "Pipeline"),
        "source": parse_choice(row.get("유입경로"), "유입경로", db.LEAD_SOURCES, "기타"),
        "competitor": _clean(row.get("경쟁사")) or None,
        "memo": _clean(row.get("메모")) or None,
    }
    return data["owner"], data


def _row_activity(row: pd.Series, ctx: dict) -> tuple[str, dict]:
    cid = _customer_id_by_name(_clean(row.get("거래처명")), ctx["cust_cache"], ctx["db_path"])
    summary = _clean(row.get("활동내용"))
    if not summary:
        raise ValueError("활동내용: 필수 항목입니다")
    deal_id = None
    deal_title = _clean(row.get("관련기회"))
    if deal_title:
        found = db._one("SELECT id FROM deals WHERE customer_id=? AND title=? ORDER BY id DESC",
                        [cid, deal_title], ctx["db_path"])
        if not found:
            raise ValueError(f"관련기회: 해당 거래처에 '{deal_title}' 기회가 없습니다")
        deal_id = int(found["id"])
    data = {
        "customer_id": cid, "deal_id": deal_id,
        "act_date": parse_date(row.get("활동일"), "활동일", required=True),
        "act_type": parse_choice(row.get("유형"), "유형", db.ACT_TYPES, "기타"),
        **_owner(row, ctx), "summary": summary,
        "next_action": _clean(row.get("다음액션")) or None,
        "next_date": parse_date(row.get("다음일정"), "다음일정"),
    }
    return data["owner"], data


def _row_sale(row: pd.Series, ctx: dict) -> tuple[str, dict]:
    cid = _customer_id_by_name(_clean(row.get("거래처명")), ctx["cust_cache"], ctx["db_path"])
    item = _clean(row.get("품목"))
    if not item:
        raise ValueError("품목: 필수 항목입니다")
    qty = parse_int(row.get("수량"), "수량", 1, 0)
    unit = parse_int(row.get("단가"), "단가", 0, 0)
    amount = parse_int(row.get("금액"), "금액", 0, 0) or qty * unit
    if amount <= 0:
        raise ValueError("금액: 금액 또는 (수량 × 단가)가 0보다 커야 합니다")
    data = {
        "customer_id": cid, "sale_date": parse_date(row.get("매출일"), "매출일", required=True),
        "item": item, "item_code": _clean(row.get("품목코드")) or None,
        "qty": qty, "unit_price": unit or (amount // qty if qty else amount),
        "amount": amount, **_owner(row, ctx),
        "status": parse_choice(row.get("수금상태"), "수금상태", db.SALE_STATUS, "입금대기"),
        "tax_type": parse_choice(row.get("과세구분"), "과세구분", db.TAX_TYPES, "과세"),
        "due_date": parse_date(row.get("결제기일"), "결제기일"),
        "memo": _clean(row.get("메모")) or None,
    }
    return data["owner"], data


def _row_target(row: pd.Series, ctx: dict) -> tuple[str, dict]:
    data = {"yyyymm": parse_ym(row.get("월"), "월"), **_owner(row, ctx),
            "amount": parse_int(row.get("목표금액"), "목표금액", 0, 0)}
    return data["owner"], data


ROW_PARSERS: dict[str, Callable] = {
    "거래처": _row_customer, "영업기회": _row_deal, "영업활동": _row_activity,
    "매출": _row_sale, "목표": _row_target,
}


# ============================================================================
# 일괄 등록
# ============================================================================
def import_rows(entity: str, df: pd.DataFrame, user: dict, dry_run: bool = True,
                on_duplicate: str = "건너뛰기", db_path: str | None = None, row_offset: int = 2) -> dict:
    """업로드 데이터를 검증하고(dry_run=True) 또는 실제로 등록한다.

    영업기회는 관리자만 이관 데이터로 보아 Stage Gate 를 우회할 수 있다(감사로그에 사유 기록).
    그 외 사용자의 영업기회 업로드는 화면 입력과 똑같이 단계 조건을 검사한다.
    반환: {"total", "ok", "skipped", "errors"[(엑셀행번호, 사유)], "preview" DataFrame}
    """
    if len(df) > MAX_UPLOAD_ROWS:
        raise ValueError(f"한 번에 {MAX_UPLOAD_ROWS:,}행까지 올릴 수 있습니다 ({len(df):,}행). 나눠 올려 주세요.")
    if entity not in IMPORT_SPECS:
        raise ValueError(f"지원하지 않는 항목입니다: {entity}")
    spec = IMPORT_SPECS[entity]
    df = df.dropna(how="all")
    missing = [c for c in spec["required"] if c not in df.columns]
    if missing:
        raise ValueError(f"필수 컬럼이 없습니다: {', '.join(missing)} "
                         f"(양식을 내려받아 헤더를 맞춰 주세요)")

    ctx = {"cust_cache": {}, "owner_cache": {}, "db_path": db_path}
    parser = ROW_PARSERS[entity]
    errors: list[tuple[int, str]] = []
    records: list[tuple[int, dict]] = []
    skipped = 0
    migrate = entity == "영업기회" and ent.has_role(user, "ADMIN")

    for idx, row in df.iterrows():
        excel_row = int(idx) + row_offset      # 머리글 행 + 0-base 보정 (회사 양식은 머리글이 아래에 있다)
        try:
            _owner_name, data = parser(row, ctx)
            if entity == "영업기회" and not migrate:
                gate = db.validate_stage(data, data["stage"], db_path) if data["stage"] != db.OPEN_STAGES[0] else []
                if gate:
                    raise ValueError(f"단계: '{data['stage']}' 조건 미충족 → {', '.join(gate)}")
            records.append((excel_row, data))
        except ValueError as exc:
            errors.append((excel_row, str(exc)))

    preview = pd.DataFrame([r[1] for r in records[:20]]) if records else pd.DataFrame()
    if not preview.empty:
        preview = preview.drop(columns=[c for c in ("owner_id", "customer_id") if c in preview.columns])

    if dry_run:
        return {"total": len(df), "ok": len(records), "skipped": 0,
                "errors": errors, "preview": preview}

    inserted = 0
    for excel_row, data in records:
        try:
            if entity == "거래처":
                blank = data.pop("_blank", [])
                exists = _existing_customer(data, ctx, db_path)
                if exists:
                    if on_duplicate == "건너뛰기" or not db.in_scope(exists["owner_id"]):
                        skipped += 1
                        continue
                    # 덮어쓰기: 파일에서 비어 있거나 없는 열은 기존 값 그대로 (두 열만 올려도 나머지가 지워지지 않게)
                    for field in blank:
                        data[field] = exists.get(field)
                    data["name"] = exists["name"]           # 다른 이름으로 맞춘 경우 정식 이름 유지
                    data["id"] = int(exists["id"])
                db.upsert_customer(data, db_path)
            elif entity == "영업기회":
                db.upsert_deal(data, db_path, force=migrate,
                               force_reason="일괄 등록(이관 데이터)" if migrate else "")
            elif entity == "영업활동":
                db.add_activity(data, db_path)
            elif entity == "매출":
                if _sale_exists(data, db_path):
                    skipped += 1                        # 같은 파일을 두 번 올려도 매출이 두 번 생기지 않게
                    errors.append((excel_row, "건너뜀: 같은 거래처·일자·품목·수량·금액의 매출이 이미 있습니다"))
                    continue
                db.upsert_sale(data, db_path)
            else:
                db.upsert_target(data["yyyymm"], data["owner_id"], data["amount"], db_path)
            inserted += 1
        except (ValueError, PermissionError) as exc:   # 한 행 실패가 전체를 막지 않도록
            errors.append((excel_row, f"등록 실패: {exc}"))

    db.audit("일괄등록", entity, None,
             {"요청": len(df), "등록": inserted, "건너뜀": skipped, "오류": len(errors),
              "단계검증우회": migrate}, db_path)
    return {"total": len(df), "ok": inserted, "skipped": skipped,
            "errors": errors, "preview": preview}


def _existing_customer(data: dict, ctx: dict, db_path: str | None) -> Optional[dict]:
    from . import customer_names as cn
    if "_name_index" not in ctx:
        ctx["_name_index"] = cn.Index(db_path)
    cid, _how = ctx["_name_index"].resolve(data["name"])
    if cid is None and db.biz_digits(data.get("biz_no")):
        row = db._one("SELECT id FROM customers WHERE biz_no_norm=? AND merged_into IS NULL",
                      [db.biz_digits(data.get("biz_no"))], db_path)
        cid = row["id"] if row else None
    return db.get_customer(int(cid), db_path) if cid else None


def _sale_exists(data: dict, db_path: str | None) -> bool:
    qty = int(data.get("qty") or 0)
    amount = int(data.get("amount") or qty * int(data.get("unit_price") or 0))
    return bool(db._one("SELECT id FROM sales WHERE customer_id=? AND sale_date=? AND item=? AND qty=? AND amount=? "
                        "AND status <> ? LIMIT 1",
                        [int(data["customer_id"]), data.get("sale_date"), str(data.get("item") or "").strip(), qty,
                         amount, db.SALE_CANCELLED], db_path))


def errors_to_df(errors: list[tuple[int, str]]) -> pd.DataFrame:
    return pd.DataFrame(errors, columns=["엑셀 행", "오류 내용"]) if errors else pd.DataFrame()


# ============================================================================
# 데이터 추출
# ============================================================================
EXPORT_SOURCES = ["거래처", "영업기회", "영업활동", "매출", "목표",
                  "채권(미수)", "결재이력", "담당자별예측", "감사로그"]

DROP_COLUMNS = {"id", "customer_id", "deal_id", "org_id", "parent_id", "역할코드", "owner_id",
                "requested_by_id"}


PII_COLUMNS = {"연락처": "phone", "이메일": "email", "고객담당자": "name"}


def mask_phone(value: Any) -> Any:
    text = _clean(value)
    digits = re.sub(r"\D", "", text)
    if len(digits) < 7:
        return "***" if text else value
    return f"{digits[:3]}-****-{digits[-4:]}"


def mask_email(value: Any) -> Any:
    text = _clean(value)
    if "@" not in text:
        return "***" if text else value
    local, domain = text.split("@", 1)
    return f"{local[:2]}***@{domain}"


def mask_name(value: Any) -> Any:
    text = _clean(value)
    return (text[0] + "*" * (len(text) - 1)) if len(text) > 1 else text


def mask_pii(frame: pd.DataFrame) -> pd.DataFrame:
    """개인정보(고객 측 연락처·이메일·담당자명)를 가린 사본."""
    if frame.empty:
        return frame
    out = frame.copy()
    funcs = {"phone": mask_phone, "email": mask_email, "name": mask_name}
    for col, kind in PII_COLUMNS.items():
        if col in out.columns:
            out[col] = out[col].map(funcs[kind])
    return out


def collect(sources: list[str], date_from: str = "", date_to: str = "", owner_id: int | None = None,
            stage: str = "", ym: str = "", include_pii: bool = False,
            db_path: str | None = None) -> dict[str, pd.DataFrame]:
    """선택한 항목을 조건에 맞춰 조회한다. 로그인 사용자의 접근범위가 그대로 적용된다.

    include_pii=False 이면 고객 연락처·이메일·담당자명을 가린다.
    """
    ym_from = date_from[:7] if date_from else ""
    ym_to = date_to[:7] if date_to else ""
    days = 3650
    if date_from:
        try:
            days = max((date.today() - datetime.strptime(date_from, "%Y-%m-%d").date()).days, 1)
        except ValueError:
            days = 3650

    out: dict[str, pd.DataFrame] = {}
    for source in sources:
        if source == "거래처":
            frame = db.list_customers(owner_id=owner_id, db_path=db_path)
        elif source == "영업기회":
            frame = db.list_deals(owner_id=owner_id, stage=stage, db_path=db_path)
            if not frame.empty and date_from:
                frame = frame[frame["예상마감일"].fillna("") >= date_from]
            if not frame.empty and date_to:
                frame = frame[frame["예상마감일"].fillna("9999-12-31") <= date_to]
        elif source == "영업활동":
            frame = db.list_activities(days=days, owner_id=owner_id, db_path=db_path)
            if not frame.empty and date_to:
                frame = frame[frame["활동일"] <= date_to]
        elif source == "매출":
            frame = db.list_sales(ym_from=ym_from, ym_to=ym_to, owner_id=owner_id, db_path=db_path)
        elif source == "목표":
            frame = db.list_targets(ym, db_path=db_path)
        elif source == "채권(미수)":
            frame = ent.ar_aging(db_path=db_path)
        elif source == "결재이력":
            frame = ent.list_approvals(db_path=db_path)
        elif source == "담당자별예측":
            frame = ent.forecast_by_owner(ym or date.today().strftime("%Y-%m"), db_path=db_path)
        elif source == "감사로그":
            frame = db.list_audit(5000, db_path=db_path)
            if not frame.empty and date_from:
                frame = frame[frame["시각"] >= date_from]
            if not frame.empty and date_to:
                frame = frame[frame["시각"] <= date_to + " 23:59:59"]
        else:
            continue
        frame = frame.drop(columns=[c for c in DROP_COLUMNS if c in frame.columns])
        out[source] = frame if include_pii else mask_pii(frame)
    return out


MONEY_HINTS = ("금액", "매출", "목표", "단가", "여신", "미수", "정가", "제안가", "입금", "공급가액", "부가세", "합계",
               "파이프라인", "Commit", "Best Case", "Pipeline", "확정", "갭")


FORMULA_PREFIX = ("=", "+", "-", "@", "\t", "\r")


def neutralize_formulas(frame: pd.DataFrame) -> pd.DataFrame:
    """엑셀 수식 주입(CSV/Formula Injection) 차단.

    '=HYPERLINK(...)' 처럼 수식으로 시작하는 '문자열' 앞에 작은따옴표를 붙여 글자로 취급되게 한다.
    숫자(음수 포함)는 건드리지 않는다.
    """
    if frame.empty:
        return frame
    out = frame.copy()
    for col in out.columns:
        if out[col].dtype == object or pd.api.types.is_string_dtype(out[col]):
            out[col] = out[col].map(
                lambda v: "'" + v if isinstance(v, str) and v.startswith(FORMULA_PREFIX) else v)
    return out


def to_excel(sheets: dict[str, pd.DataFrame], meta: dict | None = None) -> bytes:
    """다중 시트 엑셀로 내보낸다. 헤더 고정·열너비·금액 서식을 적용한다."""
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        if meta:
            neutralize_formulas(pd.DataFrame(list(meta.items()), columns=["항목", "값"])).to_excel(
                writer, sheet_name="추출정보", index=False)
        for name, frame in sheets.items():
            sheet = re.sub(r"[\\/*?:\[\]]", "_", name)[:31]
            body = neutralize_formulas(frame) if not frame.empty else pd.DataFrame({"결과": ["데이터 없음"]})
            body.to_excel(writer, sheet_name=sheet, index=False)
            ws = writer.sheets[sheet]
            header_fill = PatternFill("solid", fgColor="1F4E78")
            for col_idx, column in enumerate(body.columns, start=1):
                cell = ws.cell(row=1, column=col_idx)
                cell.font = Font(color="FFFFFF", bold=True)
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center", vertical="center")
                width = max(len(str(column)) * 2,
                            *(len(str(v)) for v in body[column].head(200).tolist() or [""]))
                ws.column_dimensions[get_column_letter(col_idx)].width = min(max(width + 2, 10), 40)
                if any(hint in str(column) for hint in MONEY_HINTS):
                    for row_idx in range(2, len(body) + 2):
                        ws.cell(row=row_idx, column=col_idx).number_format = "#,##0"
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
    return buf.getvalue()


def to_csv(frame: pd.DataFrame) -> bytes:
    """한글 Windows 엑셀에서 바로 열리도록 UTF-8 BOM 으로 인코딩한다(수식 주입 차단 포함)."""
    return neutralize_formulas(frame).to_csv(index=False).encode("utf-8-sig")


# 자주 쓰는 추출 프리셋 (보고서 목적별 묶음)
EXPORT_PRESETS = {
    "월간 실적 보고": ["매출", "목표", "담당자별예측"],
    "파이프라인 리뷰": ["영업기회", "영업활동", "담당자별예측"],
    "채권 현황": ["채권(미수)", "매출"],
    "내부통제 점검": ["결재이력", "감사로그"],
    "전체 백업": ["거래처", "영업기회", "영업활동", "매출", "목표", "결재이력"],
}