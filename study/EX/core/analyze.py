"""엑셀/CSV 파일 → 표 → 요약 대시보드(지표·분류별 집계·월별 추이·피벗) + 같은 결과를 내는 엑셀 수식."""
import csv
import datetime as dt
import io
import json
import math
import re
import secrets
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font, PatternFill

from . import formula as fx
from . import xlsx

MAX_ROWS = 20000
MAX_COLS = 60
AGGS = {'sum': '합계', 'avg': '평균', 'count': '개수', 'max': '최대', 'min': '최소'}
AGG_FUNCS = {'sum': 'SUMIFS', 'avg': 'AVERAGEIFS', 'count': 'COUNTIFS', 'max': 'MAXIFS', 'min': 'MINIFS'}
AVG_WORDS = re.compile(r'(단가|가격|율|비율|평균|점수|price|rate|%)', re.I)
ID_WORDS = re.compile(r'(번호|코드|id|no\.?|순번|사번|학번)$', re.I)


# ------------------------------------------------------------ 읽기 ----------
def read_sheets(data, filename):
    """→ [(시트 이름, 행 목록)]. 행은 셀 값 목록(날짜는 date)."""
    name = filename.lower()
    if name.endswith('.csv'):
        for enc in ('utf-8-sig', 'cp949', 'utf-16'):
            try:
                text = data.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        else:
            raise xlsx.BadFile('CSV 글자 인코딩을 알 수 없습니다(UTF-8 또는 CP949 로 저장해 주세요).')
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=',;\t')
        except csv.Error:
            dialect = csv.excel
        try:
            rows = [[_csv_value(v) for v in row[:MAX_COLS + 1]]
                    for _, row in zip(range(MAX_ROWS + 50), csv.reader(io.StringIO(text), dialect))]
        except csv.Error as e:
            raise xlsx.BadFile(f'CSV 를 읽을 수 없습니다: {e}')
        return [('CSV', rows)]
    if not name.endswith(('.xlsx', '.xlsm')):
        raise xlsx.BadFile('.xlsx · .xlsm · .csv 파일만 올릴 수 있습니다.')
    xlsx.check_zip(data)
    import openpyxl
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001
        raise xlsx.BadFile(f'파일을 열 수 없습니다: {type(e).__name__}')
    out = []
    try:
        for ws in wb.worksheets:
            rows = []
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i >= MAX_ROWS + 50:
                    break
                rows.append([_csv_value(v) if isinstance(v, str) else v for v in row[:MAX_COLS + 1]])   # 글자로 저장된 숫자·날짜도
            out.append((ws.title, rows))
    except Exception as e:  # noqa: BLE001 — 깨진 시트 XML 등
        raise xlsx.BadFile(f'시트를 읽을 수 없습니다: {type(e).__name__}')
    finally:
        wb.close()
    return out


ILLEGAL = re.compile(r'[\x00-\x08\x0b\x0c\x0e-\x1f]')


def _csv_value(s):
    s = ILLEGAL.sub('', s).strip()
    if s == '':
        return None
    n = s.replace(',', '')
    if re.fullmatch(r'\(\s*[\d.]+\s*\)', n):                  # (3) = 회계식 음수
        n = '-' + n.strip('() ')
    elif re.fullmatch(r'-?\d{1,3}( \d{3})+(\.\d+)?', n):         # 2 500 = 2500
        n = n.replace(' ', '')
    n = re.sub(r'^(-?[\d.]+)\s*(개|명|건|회|점|대|권|장|개월|일|kg|㎏)$', r'\1', n)    # 10개
    n = re.sub(r'^([-+]?)[₩$￦]\s*', r'\1', n)          # ₩120,000 · 120,000원 같은 금액
    if re.fullmatch(r'-?[\d.]+\s*원', n):
        n = re.sub(r'\s*원$', '', n)
    if re.fullmatch(r'-?\d+(?:\.\d+)?%', n):
        return float(n[:-1]) / 100
    if re.fullmatch(r'-?\d+', n):
        return int(n)
    if re.fullmatch(r'-?\d*\.\d+(?:[eE][-+]?\d+)?|-?\d+[eE][-+]?\d+', n):
        return float(n)
    d = fx.parse_date_text(s)
    if d is not None:
        return fx.serial_date(d)
    return s


def _is_blank(v):
    return v is None or (isinstance(v, str) and not v.strip())


MEMO_RE = re.compile(r'(※|\*|출처\s*[:：]|주\d*\)|주\s*[:：]|비고\s*[:：]|참고\s*[:：]|단위\s*[:：]|\(단위|자료\s*[:：]|note\s*[:：])', re.I)
UNIT_RE = re.compile(r'\(.{1,12}\)')
TOTAL_RE = re.compile(r'(총\s*)?(합\s*계|총\s*계|소\s*계|누\s*계|총합계|계)(\s*\(.*\))?')


def to_table(sheet_name, rows):
    """머리글 행 찾기 → 열 이름·종류·값 정리."""
    head_i = None
    for i, row in enumerate(rows[:30]):
        filled = [v for v in row if not _is_blank(v)]
        texts = [v for v in filled if isinstance(v, str)]
        if len(filled) >= 2 and len(texts) >= max(2, len(filled) * 0.6):
            head_i = i
            break
    if head_i is None:
        raise xlsx.BadFile(f"{'파일' if sheet_name == 'CSV' else repr(sheet_name) + ' 시트'}에서 머리글(열 이름) 행을 찾지 못했습니다. 첫 행에 열 이름을 넣어 주세요.")
    wide = any(len(r) > MAX_COLS for r in rows)
    rows = [r[:MAX_COLS] for r in rows]
    head = rows[head_i]
    used = [j for j, v in enumerate(head) if not _is_blank(v)]
    first_c, last_c = used[0], used[-1]
    # 두 줄 머리글(병합한 '매출' 아래 '1분기·2분기'): 다음 행이 숫자 없는 글자뿐이고 위 머리글에 빈칸이 있으면 합친다
    sub = rows[head_i + 1] if head_i + 1 < len(rows) else []
    sub_vals = [sub[j] if j < len(sub) else None for j in range(first_c, max(last_c, len(sub) - 1) + 1)]
    sub_strs = [v for v in sub_vals if isinstance(v, str) and v.strip()]
    units_row = bool(sub_strs) and len(sub_strs) == sum(1 for v in sub_vals if not _is_blank(v)) and \
        all(UNIT_RE.fullmatch(v.strip()) for v in sub_strs)              # (개)·(원) 같은 단위만 있는 둘째 줄
    two_line = units_row or (sum(1 for v in sub_vals if isinstance(v, str) and v.strip()) >= 2
                and not any(fx.is_num(v) or isinstance(v, (dt.date, dt.datetime)) for v in sub_vals if not _is_blank(v))
                and any(_is_blank(head[j]) if j < len(head) else True for j in range(first_c, len(sub_vals) + first_c)))
    if two_line:
        last_c = max(last_c, first_c + max(k for k, v in enumerate(sub_vals) if not _is_blank(v)))
    names, seen, top = [], Counter(), ''
    for j in range(first_c, last_c + 1):
        h = str(head[j]).strip() if j < len(head) and not _is_blank(head[j]) else ''
        if two_line:
            top = h or top
            s_ = str(sub[j]).strip() if j < len(sub) and not _is_blank(sub[j]) else ''
            n = (f'{top} {s_}' if top and s_ and top != s_ else (s_ or top)) or f'열{j + 1}'
        else:
            n = h or f'열{j + 1}'
        seen[n] += 1
        names.append(n if seen[n] == 1 else f'{n}_{seen[n]}')
    body, skipped = [], 0
    for row in rows[head_i + (2 if two_line else 1):]:
        vals = [row[j] if j < len(row) else None for j in range(first_c, last_c + 1)]
        if all(_is_blank(v) for v in vals):
            continue
        texts = [str(v).strip() for v in vals if isinstance(v, str) and v.strip()]
        if any(TOTAL_RE.fullmatch(t) for t in texts) or (len(texts) == 1 and MEMO_RE.match(texts[0]) and
                                                          sum(1 for v in vals if not _is_blank(v)) == 1):
            skipped += 1                       # 합계·소계 행과 ※ 메모 행은 자료가 아님(넣으면 합계가 두 배)
            continue
        body.append(vals)
        if len(body) >= MAX_ROWS:
            break
    cols = []
    for k, n in enumerate(names):
        vals = [r[k] for r in body if not _is_blank(r[k])]
        nums = sum(1 for v in vals if fx.is_num(v))
        dates = sum(1 for v in vals if isinstance(v, (dt.date, dt.datetime)))
        kind = 'text'
        if vals and dates >= 0.8 * len(vals):
            kind = 'date'
        elif vals and nums >= 0.8 * len(vals):
            kind = 'number'
        cols.append({'name': n, 'type': kind, 'letter': fx.col_name(first_c + 1 + k)})
    clean = []
    for r in body:
        out = []
        for k, v in enumerate(r):
            t = cols[k]['type']
            if _is_blank(v):
                out.append(None)
            elif t == 'date':
                out.append(v.date().isoformat() if isinstance(v, dt.datetime) else (v.isoformat() if isinstance(v, dt.date) else None))
            elif t == 'number':
                out.append(v if fx.is_num(v) and not isinstance(v, bool) else None)
            else:
                if isinstance(v, (dt.date, dt.datetime)):
                    out.append(v.strftime('%Y-%m-%d'))
                elif isinstance(v, float) and v == int(v):
                    out.append(str(int(v)))
                else:
                    out.append(ILLEGAL.sub('', str(v)).strip())
        clean.append(out)
    return {'sheet': sheet_name, 'columns': cols, 'rows': clean, 'head_row': head_i + 1,
            'truncated': len(body) >= MAX_ROWS, 'skipped': skipped, 'wide': wide}


# ------------------------------------------------------------ 저장 ----------
def save_upload(folder, data, filename):
    sheets = read_sheets(data, filename)
    tables, errors = {}, {}
    for name, rows in sheets:
        try:
            tables[name] = to_table(name, rows)
        except xlsx.BadFile as e:
            errors[name] = str(e)
    if not tables:
        raise xlsx.BadFile(next(iter(errors.values()), '읽을 수 있는 표가 없습니다.'))
    uid = secrets.token_hex(8)
    folder = Path(folder)
    folder.mkdir(parents=True, exist_ok=True)
    doc = {'name': filename, 'tables': tables, 'skipped': errors}
    (folder / f'{uid}.json').write_text(json.dumps(doc, ensure_ascii=False), encoding='utf-8')
    return uid, default_table(doc)


def default_table(doc):
    return max(doc['tables'].values(), key=lambda t: len(t['rows']))


def load_upload(folder, uid):
    if not re.fullmatch(r'[0-9a-f]{16}', uid or ''):
        return None
    p = Path(folder) / f'{uid}.json'
    if not p.exists():
        return None
    return json.loads(p.read_text(encoding='utf-8'))


def delete_upload(folder, uid):
    if re.fullmatch(r'[0-9a-f]{16}', uid or ''):
        (Path(folder) / f'{uid}.json').unlink(missing_ok=True)


# ------------------------------------------------------------ 분석 ----------
def profile(table):
    out = []
    n = len(table['rows'])
    for k, col in enumerate(table['columns']):
        vals = [r[k] for r in table['rows'] if r[k] is not None]
        p = {'name': col['name'], 'type': col['type'], 'letter': col['letter'], 'filled': len(vals),
             'blank': n - len(vals), 'unique': len(set(vals))}
        if col['type'] == 'number' and vals:
            p.update(sum=fx.fix(math.fsum(vals)), avg=fx.fix(math.fsum(vals) / len(vals)), min=min(vals), max=max(vals))
        elif col['type'] == 'date' and vals:
            p.update(min=min(vals), max=max(vals))
        elif vals:
            p['top'] = Counter(vals).most_common(3)
        out.append(p)
    return out


def _is_id_like(col, prof):
    return ID_WORDS.search(col['name'].replace(' ', '')) is not None or \
        (prof['filled'] > 10 and prof['unique'] == prof['filled'] and col['type'] == 'text')


def suggest(table, prof):
    """기본으로 보여 줄 분류 열·값 열·날짜 열·피벗 열."""
    cats = [(k, p) for k, (c, p) in enumerate(zip(table['columns'], prof))
            if c['type'] == 'text' and 2 <= p['unique'] <= 40 and not _is_id_like(c, p)]
    cats.sort(key=lambda kp: (kp[1]['unique'] >= kp[1]['filled'], abs(kp[1]['unique'] - 6), kp[0]))
    rest = sorted(cats[1:], key=lambda kp: (kp[1]['unique'] >= kp[1]['filled'], kp[1]['unique'], kp[0]))
    nums = [(k, p) for k, (c, p) in enumerate(zip(table['columns'], prof))
            if c['type'] == 'number' and not _is_id_like(c, p)]
    nums.sort(key=lambda kp: (bool(AVG_WORDS.search(table['columns'][kp[0]]['name'])), -abs(kp[1].get('sum', 0))))
    dates = [k for k, c in enumerate(table['columns']) if c['type'] == 'date']
    return {'group': cats[0][0] if cats else None,
            'pivot': rest[0][0] if rest else None,
            'value': nums[0][0] if nums else None,
            'date': dates[0] if dates else None,
            'cats': [k for k, _ in cats], 'nums': [k for k, _ in nums]}


def _agg(vals, how, count):
    if how == 'count':
        return count
    if not vals:
        return 0 if how == 'sum' else None      # 금액이 빈 항목: 합계 0(엑셀 SUM 과 같음), 평균·최대·최소는 없음
    if how == 'sum':
        return fx.fix(math.fsum(vals))
    if how == 'avg':
        return fx.fix(math.fsum(vals) / len(vals))
    if how == 'max':
        return max(vals)
    if how == 'min':
        return min(vals)
    raise ValueError(how)


def filtered(table, fcol=None, fval=None):
    if fcol is None or fval in (None, ''):
        return table['rows']
    return [r for r in table['rows'] if (r[fcol] if r[fcol] is not None else '(빈칸)') == fval]


def group_by(rows, gcol, vcol, how, limit=15):
    buckets = defaultdict(list)
    counts = Counter()
    for r in rows:
        key = r[gcol] if r[gcol] is not None else '(빈칸)'
        counts[key] += 1
        if vcol is not None and fx.is_num(r[vcol]):
            buckets[key].append(r[vcol])
    items = [(k, _agg(buckets[k], how, counts[k]), counts[k]) for k in counts]
    items.sort(key=lambda t: (t[1] is None, -(t[1] or 0)))
    if len(items) > limit and how in ('sum', 'count'):
        head, tail = items[:limit - 1], items[limit - 1:]
        other = fx.fix(sum(t[1] or 0 for t in tail))
        items = head + [(f'기타 {len(tail)}개', other, sum(t[2] for t in tail))]
    return items


def monthly(rows, dcol, vcol, how):
    buckets, counts = defaultdict(list), Counter()
    for r in rows:
        d = r[dcol]
        if not d:
            continue
        key = d[:7]
        counts[key] += 1
        if vcol is not None and fx.is_num(r[vcol]):
            buckets[key].append(r[vcol])
    if not counts:
        return []
    keys = sorted(counts)
    y, m = map(int, keys[0].split('-'))
    y2, m2 = map(int, keys[-1].split('-'))
    if (y2 - y) * 12 + (m2 - m) > 36:                  # 3년 넘게 이어지면 월별 대신 연별
        yb, yc = defaultdict(list), Counter()
        for k in keys:
            yc[k[:4]] += counts[k]
            yb[k[:4]] += buckets[k]
        return [(str(yy), _agg(yb[str(yy)], how, yc[str(yy)]) if yc[str(yy)] else 0, yc[str(yy)]) for yy in range(y, y2 + 1)]
    out = []
    while (y, m) <= (y2, m2) and len(out) < 60:
        k = f'{y:04d}-{m:02d}'
        out.append((k, _agg(buckets[k], how, counts[k]) if counts[k] else 0, counts[k]))
        m += 1
        if m > 12:
            y, m = y + 1, 1
    return out


def pivot(rows, rcol, ccol, vcol, how, max_cols=10, max_rows=25):
    col_keys = [k for k, _ in Counter((r[ccol] if r[ccol] is not None else '(빈칸)') for r in rows).most_common(max_cols)]
    row_keys = [k for k, _ in Counter((r[rcol] if r[rcol] is not None else '(빈칸)') for r in rows).most_common(max_rows)]
    cells = defaultdict(list)
    counts = Counter()
    for r in rows:
        rk = r[rcol] if r[rcol] is not None else '(빈칸)'
        ck = r[ccol] if r[ccol] is not None else '(빈칸)'
        for key in ((rk, ck), (rk, None), (None, ck), (None, None)):
            counts[key] += 1
            if vcol is not None and fx.is_num(r[vcol]):
                cells[key].append(r[vcol])
    val = lambda key: _agg(cells[key], how, counts[key]) if counts[key] else None  # noqa: E731
    matrix = [{'key': rk, 'vals': [val((rk, ck)) for ck in col_keys], 'total': val((rk, None))} for rk in row_keys]
    matrix.sort(key=lambda row: -(row['total'] or 0))
    return {'cols': col_keys, 'rows': matrix, 'totals': [val((None, ck)) for ck in col_keys], 'grand': val((None, None)),
            'more_cols': len({r[ccol] for r in rows}) > max_cols, 'more_rows': len({r[rcol] for r in rows}) > max_rows}


def excel_formula(table, gcol, vcol, how, key, sheet_ref=True):
    """'이 값을 엑셀 수식으로 구하면' — 원래 시트 열 위치 기준."""
    n = table['head_row'] + len(table['rows'])
    r1 = table['head_row'] + 1
    prefix = ''
    if sheet_ref:
        s = table['sheet']
        prefix = (s if re.fullmatch(r'[A-Za-z_가-힣][\w가-힣]*', s) else "'" + s.replace("'", "''") + "'") + '!'
    rng = lambda k: f"{prefix}${table['columns'][k]['letter']}${r1}:${table['columns'][k]['letter']}${n}"  # noqa: E731
    if key == '(빈칸)':
        lit = '""'                                       # 빈 칸 조건은 "" (글자 '(빈칸)' 이 아님)
    else:
        lit = '"' + str(key).replace('"', '""') + '"' if not fx.is_num(key) else str(key)
    if how == 'count':
        return f'=COUNTIFS({rng(gcol)},{lit})'
    return f'={AGG_FUNCS[how]}({rng(vcol)},{rng(gcol)},{lit})'


def dashboard(table, params):
    prof = profile(table)
    sug = suggest(table, prof)
    cols = table['columns']

    def pick(name, default, allowed):
        try:
            v = int(params.get(name))
        except (TypeError, ValueError):
            return default
        return v if v in allowed else default

    all_idx = range(len(cols))
    text_idx = [k for k in all_idx if cols[k]['type'] in ('text', 'date')]
    num_idx = [k for k in all_idx if cols[k]['type'] == 'number']
    g = pick('g', sug['group'], text_idx)
    v = pick('v', sug['value'], num_idx)
    c = pick('c', sug['pivot'], text_idx)
    how = params.get('a') if params.get('a') in AGGS else ('sum' if v is not None else 'count')
    if v is None:
        how = 'count'
    fcol = pick('f', None, text_idx)
    fval = params.get('fv') if fcol is not None else None
    rows = filtered(table, fcol, fval)

    kpis = [{'label': '행 수', 'value': len(rows), 'fmt': 'int'}]
    for k in sug['nums'][:3]:
        vals = [r[k] for r in rows if fx.is_num(r[k])]
        if vals:
            avg = math.fsum(vals) / len(vals)
            if AVG_WORDS.search(cols[k]['name']):  # 단가·비율은 더해도 의미가 없다
                kpis.append({'label': f"{cols[k]['name']} 평균", 'value': fx.fix(round(avg, 2)), 'fmt': 'num',
                             'sub': f'최소 {min(vals):,} · 최대 {max(vals):,}'})
            else:
                kpis.append({'label': f"{cols[k]['name']} 합계", 'value': fx.fix(math.fsum(vals)), 'fmt': 'num',
                             'sub': f'평균 {avg:,.2f}'.rstrip('0').rstrip('.')})
    out = {'profile': prof, 'suggest': sug, 'kpis': kpis, 'g': g, 'v': v, 'c': c, 'a': how, 'f': fcol, 'fv': fval,
           'rows_n': len(rows), 'text_idx': text_idx, 'num_idx': num_idx}
    if g is not None:
        grp = group_by(rows, g, v, how)
        out['group'] = grp
        out['group_formula'] = excel_formula(table, g, v, how, grp[0][0]) if grp and not str(grp[0][0]).startswith('기타') else None
    if sug['date'] is not None:
        out['trend'] = monthly(rows, sug['date'], v, how)
        out['date_name'] = cols[sug['date']]['name']
    if g is not None and c is not None and c != g:
        out['pivot'] = pivot(rows, g, c, v, how)
    cats = []
    for k in sug['cats'][:4]:
        if k == g:
            continue
        cats.append({'name': cols[k]['name'], 'idx': k, 'items': group_by(rows, k, None, 'count', limit=8)})
    out['cat_counts'] = cats[:3]
    if fcol is not None:
        out['filter_values'] = [x for x, _ in Counter((r[fcol] if r[fcol] is not None else '(빈칸)')
                                                      for r in table['rows']).most_common(50)]
    elif text_idx:
        out['filter_values'] = []
    out['preview'] = rows[:15]
    return out


def export_xlsx(table, dash):
    """요약 결과를 엑셀로 — 집계표 + 막대 차트 + 피벗 표."""
    wb = Workbook()
    ws = wb.active
    ws.title = '요약'
    cols = table['columns']
    head_fill = PatternFill('solid', fgColor='1F6E43')
    white = Font(bold=True, color='FFFFFF')
    agg_name = AGGS[dash['a']]
    vname = cols[dash['v']]['name'] if dash.get('v') is not None and dash['a'] != 'count' else '건수'
    if dash.get('group'):
        ws.append([cols[dash['g']]['name'], f'{vname} {agg_name}' if dash['a'] != 'count' else '건수', '건수'])
        for c in ws[1]:
            c.fill, c.font = head_fill, white
        for k, val, n in dash['group']:
            ws.append([k, val, n])
        for r in range(2, ws.max_row + 1):
            ws.cell(r, 2).number_format = '#,##0.##'
        ch = BarChart()
        ch.title = f"{cols[dash['g']]['name']}별 {vname}"
        ch.add_data(Reference(ws, min_col=2, min_row=1, max_row=ws.max_row), titles_from_data=True)
        ch.set_categories(Reference(ws, min_col=1, min_row=2, max_row=ws.max_row))
        ch.width, ch.height = 16, 8
        ws.add_chart(ch, 'E2')
        ws.column_dimensions['A'].width = 18
        ws.column_dimensions['B'].width = 16
    if dash.get('pivot'):
        pv = dash['pivot']
        ps = wb.create_sheet('피벗')
        ps.append([f"{cols[dash['g']]['name']} ↓ · {cols[dash['c']]['name']} →"] + [str(x) for x in pv['cols']] + ['합계'])
        for c in ps[1]:
            c.fill, c.font = head_fill, white
        for row in pv['rows']:
            ps.append([row['key']] + row['vals'] + [row['total']])
        ps.append(['합계'] + pv['totals'] + [pv['grand']])
        for c in ps[ps.max_row]:
            c.font = Font(bold=True)
        ps.column_dimensions['A'].width = 18
    if dash.get('trend'):
        ts = wb.create_sheet('월별')
        ts.append(['월', f'{vname} {agg_name}' if dash['a'] != 'count' else '건수', '건수'])
        for c in ts[1]:
            c.fill, c.font = head_fill, white
        for k, val, n in dash['trend']:
            ts.append([k, val, n])
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


def sample_file():
    """샘플: 실습용 매출 데이터를 엑셀 파일로."""
    from .build import sales_data
    head, rows = sales_data()
    wb = Workbook()
    ws = wb.active
    ws.title = '매출'
    ws.append(head)
    for r in rows:
        ws.append(r)
    for r in range(2, ws.max_row + 1):
        ws.cell(r, 2).number_format = 'yyyy-mm-dd'
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()
