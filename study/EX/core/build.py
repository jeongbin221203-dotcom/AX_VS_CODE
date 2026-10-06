"""대시보드 만들기 실습: 연습 파일 만들기, 완성 예시, 올린 파일 채점."""
import datetime as dt
import io
import random

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.formatting.rule import DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from . import formula as fx
from . import pivots, xlsx

DATA = '데이터'
DASH = '대시보드'
BASE_DATE = dt.date(2026, 9, 30)


# ------------------------------------------------------------ 데이터 --------
def sales_data(n=120, seed=7):
    rnd = random.Random(seed)
    regions = {'서울': ['김민준', '박지호'], '부산': ['이서연', '강도윤'], '대구': ['최유나'],
               '광주': ['정하은'], '대전': ['윤시우']}
    products = [('노트북', '전자', 1250000), ('모니터', '전자', 330000), ('키보드', '주변기기', 48000),
                ('마우스', '주변기기', 27000), ('프린터', '사무기기', 290000)]
    qty_range = {'노트북': (1, 4), '모니터': (1, 8), '키보드': (2, 20), '마우스': (5, 40), '프린터': (1, 3)}
    start = dt.date(2026, 1, 2)
    rows = []
    for i in range(n):
        d = start + dt.timedelta(days=rnd.randint(0, 179))
        reg = rnd.choices(list(regions), weights=[34, 24, 16, 13, 13])[0]
        name, cat, price = rnd.choices(products, weights=[18, 22, 24, 24, 12])[0]
        q = rnd.randint(*qty_range[name])
        rows.append([d, reg, rnd.choice(regions[reg]), name, cat, q, price])
    rows.sort(key=lambda r: r[0])
    out = []
    for i, (d, reg, rep, name, cat, q, price) in enumerate(rows, 1):
        out.append([f'S{2026}{i:04d}', d, reg, rep, name, cat, q, price, q * price])
    return ['주문번호', '날짜', '지역', '담당자', '제품', '분류', '수량', '단가', '금액'], out


def hr_data(n=60, seed=11):
    rnd = random.Random(seed)
    family = '김이박최정강조윤장임한오서신권황안송류홍'
    given = ['민준', '서연', '도윤', '하은', '지호', '수아', '예준', '지유', '시우', '채원', '주원', '다은',
             '건우', '서윤', '현우', '지민', '우진', '하린', '선우', '유나']
    depts = ['영업', '인사', '생산', '연구', '재무']
    ranks = [('사원', 3400, 0, 4), ('대리', 4300, 3, 8), ('과장', 5400, 7, 13), ('차장', 6500, 11, 18),
             ('부장', 7800, 15, 25)]
    rows = []
    for i in range(1, n + 1):
        rank, pay, ymin, ymax = rnd.choices(ranks, weights=[30, 25, 20, 13, 12])[0]
        years = rnd.randint(ymin, ymax)
        hired = BASE_DATE - dt.timedelta(days=years * 365 + rnd.randint(0, 300))
        salary = pay + rnd.randint(-4, 8) * 100
        rows.append([f'E{1000 + i}', rnd.choice(family) + rnd.choice(given), rnd.choice(depts), rank, hired,
                     rnd.choice(['남', '여']), salary, rnd.choices(['S', 'A', 'B', 'C'], weights=[10, 30, 45, 15])[0]])
    return ['사번', '이름', '부서', '직급', '입사일', '성별', '연봉(만원)', '평가'], rows


def inventory_data(n=40, seed=5):
    rnd = random.Random(seed)
    items = {'사무': [('A4 용지', 4500), ('볼펜', 800), ('포스트잇', 1500), ('클립', 2000), ('파일철', 1200),
                    ('스테이플러', 6500), ('수정테이프', 1800), ('형광펜', 900), ('바인더', 3500), ('메모지', 1000)],
             '전자': [('USB 메모리', 12000), ('무선 마우스', 25000), ('충전 케이블', 9000), ('보조배터리', 35000),
                    ('이어폰', 19000), ('멀티탭', 15000), ('웹캠', 48000), ('HDMI 케이블', 8000), ('키보드', 32000),
                    ('외장 SSD', 89000)],
             '청소': [('물티슈', 3200), ('세제', 7800), ('종이컵', 30), ('고무장갑', 2500), ('휴지', 18000),
                    ('쓰레기봉투', 5500), ('걸레', 3000), ('손세정제', 4200), ('락스', 2800), ('행주', 1500)],
             '식음료': [('생수', 500), ('커피 믹스', 13000), ('녹차', 6000), ('컵라면', 1100), ('과자', 2500),
                     ('설탕', 3000), ('우유', 2600), ('주스', 1500), ('보리차', 4500), ('사탕', 4000)]}
    codes = {'사무': 'OF', '전자': 'EL', '청소': 'CL', '식음료': 'FD'}
    rows = []
    for cat, lst in items.items():
        for k, (name, price) in enumerate(lst, 1):
            safe = rnd.choice([10, 15, 20, 30, 50])
            stock = max(0, int(safe * rnd.uniform(0.2, 3.5)))
            last = BASE_DATE - dt.timedelta(days=rnd.randint(3, 200))
            rows.append([f'{codes[cat]}-{100 + k}', name, cat, rnd.choice(['본사', '물류센터']), stock, safe, price, last])
    rows = rows[:n]
    return ['품목코드', '품목명', '분류', '창고', '재고량', '안전재고', '단가', '최근입고일'], rows


# ------------------------------------------------------------ 과제 정의 -----
def _mission_sales():
    head, rows = sales_data()
    n = len(rows) + 1
    R = lambda col: f'{DATA}!${col}$2:${col}${n}'  # noqa: E731
    return {
        'key': 'sales', 'title': '상반기 매출 대시보드', 'level': 1, 'tracks': ['work', 'c2'],
        'desc': '주문 120건으로 핵심 지표·지역별·월별·제품별 매출을 요약하고 차트를 붙입니다.',
        'head': head, 'rows': rows, 'formats': {'B': 'yyyy-mm-dd', 'H': '#,##0', 'I': '#,##0'},
        'labels': {'A1': '상반기 매출 대시보드', 'A3': '총매출', 'A4': '주문 건수', 'A5': '평균 주문 금액',
                   'A6': '최대 주문 금액', 'A8': '지역', 'B8': '매출', 'C8': '비중',
                   **{f'A{9 + i}': r for i, r in enumerate(['서울', '부산', '대구', '광주', '대전'])},
                   'E8': '월', 'F8': '매출', **{f'E{9 + i}': i + 1 for i in range(6)},
                   'A16': '제품', 'B16': '판매 수량',
                   **{f'A{17 + i}': p for i, p in enumerate(['노트북', '모니터', '키보드', '마우스', '프린터'])},
                   'A23': '판매 수량 1위 제품'},
        'dash_formats': {'B3': '#,##0', 'B4': '#,##0', 'B5': '#,##0', 'B6': '#,##0', 'B9:B13': '#,##0',
                         'C9:C13': '0.0%', 'E9:E14': '0"월"', 'F9:F14': '#,##0', 'B17:B21': '#,##0'},
        'checks': [
            {'id': 'total', 'kind': 'cell', 'at': 'B3', 'label': '총매출', 'answer': f'=SUM({R("I")})',
             'hint': '금액 열 전체를 SUM 으로 더합니다.'},
            {'id': 'count', 'kind': 'cell', 'at': 'B4', 'label': '주문 건수', 'answer': f'=COUNTA({R("A")})',
             'hint': '주문번호가 몇 개인지 COUNTA 로 셉니다(COUNT 는 숫자만 셉니다).'},
            {'id': 'avg', 'kind': 'cell', 'at': 'B5', 'label': '평균 주문 금액', 'answer': f'=AVERAGE({R("I")})',
             'hint': 'AVERAGE 또는 =B3/B4.'},
            {'id': 'max', 'kind': 'cell', 'at': 'B6', 'label': '최대 주문 금액', 'answer': f'=MAX({R("I")})',
             'hint': 'MAX 함수.'},
            {'id': 'region', 'kind': 'range', 'at': 'B9:B13', 'label': '지역별 매출',
             'answer': f'=SUMIFS({R("I")},{R("C")},A9)',
             'hint': 'B9 에 =SUMIFS(금액, 지역, A9) 를 만들고 범위는 $로 고정한 뒤 아래로 채웁니다.'},
            {'id': 'share', 'kind': 'range', 'at': 'C9:C13', 'label': '지역별 비중', 'answer': '=B9/$B$3',
             'hint': '지역 매출 ÷ 총매출. 총매출 셀은 $B$3 으로 고정합니다.'},
            {'id': 'month', 'kind': 'range', 'at': 'F9:F14', 'label': '월별 매출',
             'answer': f'=SUMIFS({R("I")},{R("B")},">="&DATE(2026,E9,1),{R("B")},"<"&DATE(2026,E9+1,1))',
             'hint': '날짜가 그 달 1일 이상, 다음 달 1일 미만인 금액을 SUMIFS 로 더합니다. '
                     'SUMPRODUCT((MONTH(날짜)=E9)*금액) 도 됩니다.'},
            {'id': 'product', 'kind': 'range', 'at': 'B17:B21', 'label': '제품별 판매 수량',
             'answer': f'=SUMIFS({R("G")},{R("E")},A17)', 'hint': '수량 열을 제품 기준으로 SUMIFS.'},
            {'id': 'top', 'kind': 'cell', 'at': 'B23', 'label': '판매 수량 1위 제품',
             'answer': '=INDEX(A17:A21,MATCH(MAX(B17:B21),B17:B21,0))',
             'hint': 'MAX 로 가장 큰 수량을 찾고 MATCH 로 위치, INDEX 로 제품 이름을 꺼냅니다.'},
            {'id': 'chart', 'kind': 'chart', 'label': '차트 1개 이상', 'min': 1,
             'hint': '지역별 매출(A8:B13)을 선택하고 [삽입] > [세로 막대형 차트].'},
            {'id': 'cf', 'kind': 'cf', 'label': '조건부 서식 1개 이상',
             'hint': 'B9:B13 을 선택하고 [홈] > [조건부 서식] > [데이터 막대].'},
            {'id': 'pivot', 'kind': 'pivot', 'label': '(선택) 피벗 테이블', 'optional': True,
             'pivot': {'rows': ['지역'], 'cols': ['분류'], 'values': [['금액', 'sum']]},
             'hint': '데이터 시트에서 [삽입] > [피벗 테이블] — 행: 지역, 열: 분류, 값: 금액 합계.'},
        ],
        'answer_charts': [('bar', 'A8:B13', 'H3', '지역별 매출'), ('line', 'E8:F14', 'H20', '월별 매출')],
        'answer_cf': 'B9:B13',
    }


def _mission_hr():
    head, rows = hr_data()
    n = len(rows) + 1
    R = lambda col: f'{DATA}!${col}$2:${col}${n}'  # noqa: E731
    depts = ['영업', '인사', '생산', '연구', '재무']
    ranks = ['사원', '대리', '과장', '차장', '부장']
    return {
        'key': 'hr', 'title': '인사 현황 대시보드', 'level': 2, 'tracks': ['work', 'c2', 'c1'],
        'desc': '사원 60명의 부서·직급·연봉·평가를 요약합니다. COUNTIF(S)·AVERAGEIF·INDEX/MATCH·날짜 조건 연습.',
        'head': head, 'rows': rows, 'formats': {'E': 'yyyy-mm-dd', 'G': '#,##0'},
        'labels': {'A1': '인사 현황 대시보드', 'A2': '기준일', 'B2': BASE_DATE,
                   'A4': '전체 인원', 'A5': '평균 연봉(만원)', 'A6': '여성 비율',
                   'A8': '부서', 'B8': '인원', 'C8': '평균 연봉', 'D8': 'S·A 평가 인원',
                   **{f'A{9 + i}': d for i, d in enumerate(depts)},
                   'F8': '직급', 'G8': '인원', **{f'F{9 + i}': r for i, r in enumerate(ranks)},
                   'A16': '근속 10년 이상 인원', 'A17': '최고 연봉자'},
        'dash_formats': {'B2': 'yyyy-mm-dd', 'B5': '#,##0', 'B6': '0.0%', 'C9:C13': '#,##0'},
        'checks': [
            {'id': 'count', 'kind': 'cell', 'at': 'B4', 'label': '전체 인원', 'answer': f'=COUNTA({R("A")})',
             'hint': '사번 개수를 COUNTA 로.'},
            {'id': 'avg', 'kind': 'cell', 'at': 'B5', 'label': '평균 연봉(반올림)',
             'answer': f'=ROUND(AVERAGE({R("G")}),0)', 'hint': 'AVERAGE 결과를 ROUND(…,0) 으로 정수로.'},
            {'id': 'female', 'kind': 'cell', 'at': 'B6', 'label': '여성 비율',
             'answer': f'=COUNTIF({R("F")},"여")/B4', 'hint': 'COUNTIF(성별,"여") ÷ 전체 인원.'},
            {'id': 'dept_n', 'kind': 'range', 'at': 'B9:B13', 'label': '부서별 인원',
             'answer': f'=COUNTIF({R("C")},A9)', 'hint': 'COUNTIF(부서 범위, A9) — 범위는 $ 고정.'},
            {'id': 'dept_pay', 'kind': 'range', 'at': 'C9:C13', 'label': '부서별 평균 연봉(반올림)',
             'answer': f'=ROUND(AVERAGEIF({R("C")},A9,{R("G")}),0)', 'hint': 'ROUND(AVERAGEIF(부서, A9, 연봉), 0).'},
            {'id': 'dept_sa', 'kind': 'range', 'at': 'D9:D13', 'label': '부서별 S·A 평가 인원',
             'answer': f'=COUNTIFS({R("C")},A9,{R("H")},"S")+COUNTIFS({R("C")},A9,{R("H")},"A")',
             'hint': 'COUNTIFS 두 개를 더하거나 SUM(COUNTIFS(…,{"S","A"})).'},
            {'id': 'rank_n', 'kind': 'range', 'at': 'G9:G13', 'label': '직급별 인원',
             'answer': f'=COUNTIF({R("D")},F9)', 'hint': 'COUNTIF(직급 범위, F9).'},
            {'id': 'senior', 'kind': 'cell', 'at': 'B16', 'label': '근속 10년 이상',
             'answer': f'=COUNTIF({R("E")},"<="&EDATE(B2,-120))',
             'hint': '입사일이 기준일 120개월 전(EDATE(B2,-120)) 이하인 사람 수.'},
            {'id': 'top', 'kind': 'cell', 'at': 'B17', 'label': '최고 연봉자',
             'answer': f'=INDEX({R("B")},MATCH(MAX({R("G")}),{R("G")},0))',
             'hint': 'INDEX(이름, MATCH(MAX(연봉), 연봉, 0)) 또는 XLOOKUP.'},
            {'id': 'chart', 'kind': 'chart', 'label': '차트 1개 이상', 'min': 1,
             'hint': '부서별 인원이나 직급별 인원으로 막대·원형 차트.'},
            {'id': 'cf', 'kind': 'cf', 'label': '조건부 서식 1개 이상',
             'hint': '평균 연봉 C9:C13 에 색조 또는 데이터 막대.'},
        ],
        'answer_charts': [('bar', 'A8:B13', 'I3', '부서별 인원'), ('pie', 'F8:G13', 'I20', '직급 구성')],
        'answer_cf': 'C9:C13',
    }


def _mission_inventory():
    head, rows = inventory_data()
    n = len(rows) + 1
    R = lambda col: f'{DATA}!${col}$2:${col}${n}'  # noqa: E731
    cats = ['사무', '전자', '청소', '식음료']
    return {
        'key': 'inventory', 'title': '재고 점검 대시보드', 'level': 3, 'tracks': ['work', 'c1'],
        'desc': '품목 40개의 재고 금액·안전재고 미달·창고별 현황을 SUMPRODUCT 와 드롭다운(유효성 검사)으로 만듭니다.',
        'head': head, 'rows': rows, 'formats': {'G': '#,##0', 'H': 'yyyy-mm-dd'},
        'labels': {'A1': '재고 점검 대시보드', 'A2': '기준일', 'B2': BASE_DATE,
                   'A4': '총 재고 금액', 'A5': '품목 수', 'A6': '안전재고 미달 품목 수',
                   'A8': '분류', 'B8': '품목 수', 'C8': '재고 금액', **{f'A{9 + i}': c for i, c in enumerate(cats)},
                   'E8': '창고', 'F8': '재고 금액', 'E9': '본사', 'E10': '물류센터',
                   'A15': '분류 선택', 'B15': '사무', 'A16': '선택 분류의 미달 품목 수',
                   'A17': '90일 넘게 입고 없는 품목 수'},
        'dash_formats': {'B2': 'yyyy-mm-dd', 'B4': '#,##0', 'C9:C12': '#,##0', 'F9:F10': '#,##0'},
        'checks': [
            {'id': 'value', 'kind': 'cell', 'at': 'B4', 'label': '총 재고 금액',
             'answer': f'=SUMPRODUCT({R("E")},{R("G")})', 'hint': 'SUMPRODUCT(재고량, 단가) — 곱한 뒤 모두 더합니다.'},
            {'id': 'items', 'kind': 'cell', 'at': 'B5', 'label': '품목 수', 'answer': f'=COUNTA({R("A")})',
             'hint': 'COUNTA(품목코드).'},
            {'id': 'short', 'kind': 'cell', 'at': 'B6', 'label': '안전재고 미달 품목 수',
             'answer': f'=SUMPRODUCT(--({R("E")}<{R("F")}))',
             'hint': '두 열을 비교하는 조건은 COUNTIF 로 안 됩니다. SUMPRODUCT(--(재고량<안전재고)).'},
            {'id': 'cat_n', 'kind': 'range', 'at': 'B9:B12', 'label': '분류별 품목 수',
             'answer': f'=COUNTIF({R("C")},A9)', 'hint': 'COUNTIF(분류, A9).'},
            {'id': 'cat_v', 'kind': 'range', 'at': 'C9:C12', 'label': '분류별 재고 금액',
             'answer': f'=SUMPRODUCT(({R("C")}=A9)*{R("E")}*{R("G")})',
             'hint': 'SUMPRODUCT((분류=A9)*재고량*단가).'},
            {'id': 'wh_v', 'kind': 'range', 'at': 'F9:F10', 'label': '창고별 재고 금액',
             'answer': f'=SUMPRODUCT(({R("D")}=E9)*{R("E")}*{R("G")})', 'hint': '분류별과 같은 방식으로 창고 기준.'},
            {'id': 'dv', 'kind': 'dv', 'at': 'B15', 'label': 'B15 드롭다운(목록 유효성 검사)',
             'hint': 'B15 선택 → [데이터] > [데이터 유효성 검사] > 제한 대상 "목록", 원본 =$A$9:$A$12.'},
            {'id': 'pick', 'kind': 'cell', 'at': 'B16', 'label': '선택 분류의 미달 품목 수',
             'answer': f'=SUMPRODUCT(({R("C")}=B15)*({R("E")}<{R("F")}))',
             'hint': '분류 조건과 미달 조건을 곱합니다. B15 를 바꾸면 결과도 바뀌어야 합니다.'},
            {'id': 'old', 'kind': 'cell', 'at': 'B17', 'label': '90일 넘게 입고 없는 품목 수',
             'answer': f'=COUNTIF({R("H")},"<"&B2-90)', 'hint': 'COUNTIF(최근입고일, "<"&B2-90).'},
            {'id': 'chart', 'kind': 'chart', 'label': '차트 1개 이상', 'min': 1, 'hint': '분류별 재고 금액 막대 차트.'},
            {'id': 'cf', 'kind': 'cf', 'label': '조건부 서식 1개 이상',
             'hint': '데이터 시트에서 =$E2<$F2 수식 규칙으로 미달 품목 행 강조.'},
        ],
        'answer_charts': [('bar', 'A8:C12', 'H3', '분류별 재고')],
        'answer_cf': 'C9:C12',
    }


MISSIONS = {m['key']: m for m in (_mission_sales(), _mission_hr(), _mission_inventory())}


def mission(key):
    return MISSIONS.get(key)


# ------------------------------------------------------------ 파일 만들기 ---
HEAD_FILL = PatternFill('solid', fgColor='1F6E43')
INPUT_FILL = PatternFill('solid', fgColor='FFF4C2')
LABEL_FILL = PatternFill('solid', fgColor='E8F1EC')
THIN = Side(style='thin', color='B7C4BC')


def _cells_of(ref):
    r1, c1, r2, c2 = fx.parse_range(ref)
    return [(r, c) for r in range(r1, r2 + 1) for c in range(c1, c2 + 1)]


def _apply_formats(ws, fmts):
    for ref, fmt in fmts.items():
        for r, c in _cells_of(ref):
            ws.cell(r, c).number_format = fmt


def workbook(m, answers=False):
    wb = Workbook()
    guide = wb.active
    guide.title = '과제'
    data = wb.create_sheet(DATA)
    dash = wb.create_sheet(DASH)

    # 데이터
    data.append(m['head'])
    for row in m['rows']:
        data.append(row)
    for j in range(1, len(m['head']) + 1):
        c = data.cell(1, j)
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = HEAD_FILL
        c.alignment = Alignment(horizontal='center')
        data.column_dimensions[get_column_letter(j)].width = 13
    for col, fmt in m['formats'].items():
        for r in range(2, len(m['rows']) + 2):
            data[f'{col}{r}'].number_format = fmt
    data.freeze_panes = 'A2'
    data.auto_filter.ref = f'A1:{get_column_letter(len(m["head"]))}{len(m["rows"]) + 1}'

    # 대시보드 틀
    for a, v in m['labels'].items():
        cell = dash[a]
        cell.value = v
        if a != 'A1' and not (isinstance(v, (int, dt.date))):
            cell.fill = LABEL_FILL
            cell.font = Font(bold=True)
    dash['A1'].font = Font(bold=True, size=16, color='1F6E43')
    for col in 'ABCDEFG':
        dash.column_dimensions[col].width = 16
    dash.column_dimensions['A'].width = 24
    for chk in m['checks']:
        if chk['kind'] in ('cell', 'range'):
            for r, c in _cells_of(chk['at']):
                dash.cell(r, c).fill = INPUT_FILL
                dash.cell(r, c).border = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
        if chk['kind'] == 'dv':
            dash[chk['at']].fill = INPUT_FILL
    _apply_formats(dash, m['dash_formats'])

    # 과제 안내
    guide['A1'] = m['title']
    guide['A1'].font = Font(bold=True, size=15, color='1F6E43')
    guide['A2'] = m['desc']
    guide['A3'] = '노란 칸에 수식을 넣고, 차트·조건부 서식을 추가한 뒤 저장해서 사이트에 올리면 채점됩니다.'
    guide['A5'], guide['B5'], guide['C5'] = '번호', '할 일', '위치 · 힌트'
    for c in ('A5', 'B5', 'C5'):
        guide[c].font = Font(bold=True, color='FFFFFF')
        guide[c].fill = HEAD_FILL
    for i, chk in enumerate(m['checks'], 1):
        guide.cell(5 + i, 1, i)
        guide.cell(5 + i, 2, chk['label'])
        where = f"{DASH}!{chk['at']}  " if chk.get('at') else ''
        guide.cell(5 + i, 3, where + chk['hint'])
    guide.column_dimensions['A'].width = 6
    guide.column_dimensions['B'].width = 26
    guide.column_dimensions['C'].width = 90

    if answers:
        _fill_answers(m, wb)
        wb.calculation.fullCalcOnLoad = True
        wb.active = 2
    else:
        wb.active = 0
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


def _fill_answers(m, wb):
    dash = wb[DASH]
    for chk in m['checks']:
        if chk['kind'] not in ('cell', 'range'):
            continue
        ast = fx.parse(chk['answer'])
        cells = _cells_of(chk['at'])
        r0, c0 = cells[0]
        for r, c in cells:
            dash.cell(r, c).value = '=' + fx.unparse(fx.shift(ast, r - r0, c - c0))
    for chk in m['checks']:
        if chk['kind'] == 'dv':
            dv = DataValidation(type='list', formula1='=$A$9:$A$12', allow_blank=False)
            dash.add_data_validation(dv)
            dv.add(chk['at'])
    for kind, ref, anchor, title in m.get('answer_charts', []):
        r1, c1, r2, c2 = fx.parse_range(ref)
        chart = {'bar': BarChart, 'line': LineChart, 'pie': PieChart}[kind]()
        chart.title = title
        chart.add_data(Reference(dash, min_col=c1 + 1, max_col=c2, min_row=r1, max_row=r2), titles_from_data=True)
        chart.set_categories(Reference(dash, min_col=c1, min_row=r1 + 1, max_row=r2))
        chart.width, chart.height = 14, 7
        dash.add_chart(chart, anchor)
    if m.get('answer_cf'):
        dash.conditional_formatting.add(m['answer_cf'], DataBarRule(start_type='min', end_type='max', color='5B9BD5'))


# ------------------------------------------------------------ 채점 ----------
def _find_sheet(wb, name):
    for ws in wb.worksheets:
        if ws.title.strip() == name:
            return ws
    return None


def _grade_unlimited(m, data):
    """올린 .xlsx 채점 → {'items': [...], 'score': n, 'total': n, 'notes': [...]}"""
    wb_f = xlsx.load(data, data_only=False)
    wb_v = xlsx.load(data, data_only=True)
    notes = []
    if _find_sheet(wb_f, DATA) is None or _find_sheet(wb_f, DASH) is None:
        raise xlsx.BadFile(f"'{DATA}' 와 '{DASH}' 시트가 있어야 합니다. 연습 파일을 내려받아 그 파일에서 작업해 주세요.")
    user = xlsx.to_book(wb_f, wb_v, today=BASE_DATE)
    ref = xlsx.to_book(wb_f, wb_v, today=BASE_DATE)
    ref_dash = ref.sheets[_find_sheet(wb_f, DASH).title]
    user_dash = user.sheets[_find_sheet(wb_f, DASH).title]
    ws_dash = _find_sheet(wb_f, DASH)

    # 기준 통합문서: 채점 칸은 비워 두고 정답을 차례로 채운다
    for chk in m['checks']:
        if chk['kind'] in ('cell', 'range'):
            for r, c in _cells_of(chk['at']):
                ref_dash.cells.pop((r, c), None)
    uncached = sum(1 for s in user.sheets.values() for v in s.cells.values()
                   if isinstance(v, fx.Formula) and v.cached is None)
    if uncached:
        notes.append('저장된 계산 결과가 없는 수식은 이 사이트의 계산기로 직접 계산했습니다.')

    items = []
    for chk in m['checks']:
        kind = chk['kind']
        item = {'id': chk['id'], 'label': chk['label'], 'hint': chk['hint'], 'at': chk.get('at'),
                'optional': chk.get('optional', False), 'cells': []}
        if kind in ('cell', 'range'):
            ast = fx.parse(chk['answer'])
            cells = _cells_of(chk['at'])
            r0, c0 = cells[0]
            all_ok, msgs = True, []
            for r, c in cells:
                exp = fx.evaluate(fx.shift(ast, r - r0, c - c0), ref, ref_dash.name, r, c)
                ref_dash.set(r, c, exp)
                raw = ws_dash.cell(r, c).value
                ftext = xlsx.formula_text(raw)
                got = user_dash.get(r, c)
                ok = got is not None and fx.same_value(got, exp, tol=1e-6)
                if ok and not ftext:
                    ok = False
                    msgs.append(f'{fx.addr(r, c)}: 값은 맞지만 수식이 아니라 직접 입력한 값입니다.')
                elif not ok and raw is None:
                    msgs.append(f'{fx.addr(r, c)}: 비어 있음')
                elif not ok and not ftext:
                    msgs.append(f'{fx.addr(r, c)}: 수식이 아니라 값({fx.display(got)})을 직접 입력했습니다 — '
                                f'기대한 결과 {fx.display(exp)}')
                elif not ok:
                    msgs.append(f'{fx.addr(r, c)}: 결과 {fx.display(got) if got is not None else "(없음)"} → '
                                f'기대한 결과 {fx.display(exp)}')
                all_ok &= ok
                item['cells'].append({'addr': fx.addr(r, c), 'formula': ftext or ('' if raw is None else str(raw)),
                                      'got': fx.display(got) if got is not None else '', 'expected': fx.display(exp),
                                      'ok': ok})
            item['ok'] = all_ok
            item['msgs'] = msgs[:3]
            item['answer'] = chk['answer']
        elif kind == 'chart':
            n = sum(len(ws._charts) for ws in wb_f.worksheets)
            item['ok'] = n >= chk.get('min', 1)
            item['msgs'] = [f'차트 {n}개 찾음']
        elif kind == 'cf':
            n = sum(len(cf.rules) for ws in wb_f.worksheets for cf in ws.conditional_formatting)
            item['ok'] = n >= 1
            item['msgs'] = [f'조건부 서식 규칙 {n}개 찾음']
        elif kind == 'dv':
            found = [d for d in ws_dash.data_validations.dataValidation
                     if d.type == 'list' and _in_sqref(chk['at'], str(d.sqref))]
            item['ok'] = bool(found)
            item['msgs'] = ['목록 유효성 검사 있음' if found else f'{chk["at"]} 에 목록 유효성 검사가 없습니다']
        elif kind == 'pivot':
            res = pivots.check(chk.get('pivot', {}), wb_f, wb_v, user)
            item['ok'], item['msgs'] = res['ok'], res['msgs']
        items.append(item)

    counted = [i for i in items if not i['optional']]
    score = sum(1 for i in counted if i['ok'])
    return {'items': items, 'score': score, 'total': len(counted), 'notes': notes}


def _in_sqref(at, sqref):
    r, c = fx.parse_addr(at)
    for part in sqref.split():
        r1, c1, r2, c2 = fx.parse_range(part.replace('$', ''))
        if r1 <= r <= r2 and c1 <= c <= c2:
            return True
    return False



GRADE_SECONDS = 25


def grade(m, data):
    """채점(계산 시간 제한 {GRADE_SECONDS}초 — 넘으면 BadFile 로 안내)."""
    try:
        with fx.time_limit(GRADE_SECONDS):
            return _grade_unlimited(m, data)
    except fx.TimeUp:
        raise xlsx.BadFile('파일 속 수식 계산이 너무 오래 걸려 채점을 멈췄습니다 — 아주 큰 범위·배열 수식을 줄여 다시 올려 주세요.')
