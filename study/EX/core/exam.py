"""컴활 실기 형식 모의고사: 문제 파일 만들기와 올린 답안 파일 채점.

시험 정의는 content/exams/*.json (형식은 content/exams/SCHEMA.md).
채점 항목(check) 종류: values · formula · style · comment · name · cf · dv · filter · sorted · subtotal · pivot ·
goalseek · scenario · datatable · chart · macro · vba · page · protect
"""
import datetime as dt
import io
import json
import re
import threading
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.styles.colors import COLOR_INDEX
from openpyxl.worksheet.formula import ArrayFormula

from . import formula as fx
from . import pivots, vba, xlsx

ROOT = Path(__file__).resolve().parent.parent
EXAM_DIR = ROOT / 'content' / 'exams'
TODAY = dt.date(2026, 10, 1)
LEVELS = {'c2': '컴활 2급', 'c1': '컴활 1급'}
SECTIONS = ['기본작업', '계산작업', '분석작업', '기타작업']

_lock = threading.Lock()
_cache = {'stamp': None, 'data': {}}


# ------------------------------------------------------------ 읽기 ----------
def _load():
    out = {}
    for p in sorted(EXAM_DIR.glob('*.json')):
        e = json.loads(p.read_text(encoding='utf-8'))
        e['_file'] = p.name
        out[e['id']] = e
    return out


def exams(level=None):
    with _lock:
        stamp = tuple((p.name, p.stat().st_mtime_ns) for p in sorted(EXAM_DIR.glob('*.json')))
        if stamp != _cache['stamp']:
            _cache['data'] = _load()
            _cache['stamp'] = stamp
        items = list(_cache['data'].values())
    items.sort(key=lambda e: (e['level'] != 'c2', e.get('round', 0), e['id']))
    return [e for e in items if not level or e['level'] == level]


def get(eid):
    return next((e for e in exams() if e['id'] == eid), None)


def task_points(task):
    return sum(c.get('points', 0) for c in task.get('checks', []))


def total_points(exam):
    return sum(task_points(t) for t in exam['tasks'])


def section_points(exam):
    out = {s: 0 for s in SECTIONS}
    for t in exam['tasks']:
        out[t['section']] = out.get(t['section'], 0) + task_points(t)
    return out


# ------------------------------------------------------------ 값 변환 -------
def _val(v):
    """JSON 값 → 계산기 값 ('@2026-03-01' 은 날짜 일련번호, '=…' 는 수식 문자열 그대로)."""
    if isinstance(v, str) and v.startswith('@') and fx.parse_date_text(v[1:]) is not None:
        return fx.parse_date_text(v[1:])
    return v


def _xl_val(v):
    """JSON 값 → 엑셀 파일에 쓸 값 (날짜는 date)."""
    if isinstance(v, str) and v.startswith('@') and fx.parse_date_text(v[1:]) is not None:
        return fx.serial_date(fx.parse_date_text(v[1:]))
    return v


def _cells(ref):
    r1, c1, r2, c2 = fx.parse_range(ref)
    return [(r, c) for r in range(r1, r2 + 1) for c in range(c1, c2 + 1)]


def sheet_cells(spec, conv=None):
    """시험 시트 정의 → {(r, c): 값}"""
    _val = conv or globals()['_val']
    cells = {}
    blocks = [{'at': spec.get('at', 'A1'), 'rows': spec.get('rows', [])}] + spec.get('blocks', [])
    for b in blocks:
        r0, c0 = fx.parse_addr(b['at'])
        for i, row in enumerate(b['rows']):
            for j, v in enumerate(row):
                if v is not None and v != '':
                    cells[(r0 + i, c0 + j)] = _val(v)
    for a, v in (spec.get('cells') or {}).items():
        cells[fx.parse_addr(a)] = _val(v)
    return cells


# ------------------------------------------------------------ 문제 파일 -----
THIN = Side(style='thin', color='000000')


def _rgb(c):
    c = c.lstrip('#').upper()
    return c if len(c) == 8 else 'FF' + c


def problem_workbook(exam):
    """시험 문제 파일(.xlsx). 시트·데이터·기본 서식·처음부터 있는 차트."""
    wb = Workbook()
    wb.remove(wb.active)
    for spec in exam['sheets']:
        ws = wb.create_sheet(spec['name'])
        for (r, c), v in sheet_cells(spec, _xl_val).items():
            cell = ws.cell(r, c)
            if isinstance(v, dt.date):
                cell.value = v
                cell.number_format = 'yyyy-mm-dd'
            else:
                cell.value = v
        for col, fmt in (spec.get('formats') or {}).items():
            refs = _cells(col) if any(ch.isdigit() for ch in col) else \
                [(r, fx.col_num(col)) for r in range(1, ws.max_row + 1)]
            for r, c in refs:
                if ws.cell(r, c).value is not None or any(ch.isdigit() for ch in col):
                    ws.cell(r, c).number_format = fmt
        for col, w in (spec.get('widths') or {}).items():
            ws.column_dimensions[col].width = w
        for ref in spec.get('merge', []):
            ws.merge_cells(ref)
        for ref in spec.get('bold', []):
            for r, c in _cells(ref):
                ws.cell(r, c).font = Font(bold=True)
        for ref, color in (spec.get('fills') or {}).items():
            for r, c in _cells(ref):
                ws.cell(r, c).fill = PatternFill('solid', fgColor=_rgb(color))
        for ref in spec.get('borders', []):
            for r, c in _cells(ref):
                ws.cell(r, c).border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
        for ref, how in (spec.get('align') or {}).items():
            for r, c in _cells(ref):
                ws.cell(r, c).alignment = Alignment(horizontal=how, vertical='center')
        if spec.get('title'):
            t = spec['title']
            ws[t['at']].font = Font(bold=True, size=t.get('size', 14))
        for ch in spec.get('charts', []):
            _add_chart(ws, ch)
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


def _add_chart(ws, spec):
    kind = spec.get('type', 'col')
    chart = {'col': BarChart, 'bar': BarChart, 'line': LineChart, 'pie': PieChart}[kind]()
    if kind == 'bar':
        chart.type = 'bar'
    r1, c1, r2, c2 = fx.parse_range(spec['data'])
    chart.add_data(Reference(ws, min_col=c1, max_col=c2, min_row=r1, max_row=r2), titles_from_data=True)
    cr1, cc1, cr2, _ = fx.parse_range(spec['cats'])
    chart.set_categories(Reference(ws, min_col=cc1, min_row=cr1, max_row=cr2))
    if spec.get('title'):
        chart.title = spec['title']
    chart.width, chart.height = spec.get('width', 15), spec.get('height', 8)
    ws.add_chart(chart, spec.get('at', 'H3'))


def original_book(exam):
    """문제 파일 그대로의 계산기 통합문서(기대 값 계산용)."""
    sheets = []
    for spec in exam['sheets']:
        cells = {k: (fx.Formula(v) if isinstance(v, str) and v.startswith('=') else v)
                 for k, v in sheet_cells(spec).items()}
        sheets.append(fx.Sheet(spec['name'], cells))
    return fx.Book(sheets, today=TODAY, names=exam.get('names'))


# ------------------------------------------------------------ 채점 도우미 ---
def _norm_text(v):
    return re.sub(r'\s+', ' ', str(v)).strip() if v is not None else ''


def _same(got, exp, tol=1e-6):
    if isinstance(exp, str) or isinstance(got, str):
        return _norm_text(got) == _norm_text(exp)
    if got is None:
        got = 0 if exp == 0 else None
    if got is None:
        return exp is None
    return fx.same_value(got, exp, tol)


def _color(c):
    """openpyxl 색 → 'RRGGBB'. 테마 색은 'T4'·'T4+0.40'(테마 번호·밝기), 자동 색은 None."""
    if c is None:
        return None
    kind = getattr(c, 'type', None)
    if kind == 'theme':
        tint = round(float(getattr(c, 'tint', 0) or 0), 2)
        return f'T{c.theme}' + (f'{tint:+.2f}' if tint else '')
    if kind == 'indexed':
        i = c.indexed
        return COLOR_INDEX[i][-6:].upper() if isinstance(i, int) and 0 <= i < 64 else None
    if kind == 'auto':
        return None
    try:
        rgb = c.rgb
    except (AttributeError, ValueError):
        return None
    if isinstance(rgb, str) and len(rgb) >= 6:
        return rgb[-6:].upper()
    return None


THEME_SLOTS = ('lt1', 'dk1', 'lt2', 'dk2', 'accent1', 'accent2', 'accent3', 'accent4', 'accent5', 'accent6',
               'hlink', 'folHlink')                 # 셀 색의 theme 번호 순서(0=배경1 흰색, 1=텍스트1 검정 …)
HALIGN_KO = {'general': '일반', 'left': '왼쪽', 'center': '가운데', 'right': '오른쪽', 'fill': '채우기',
             'centerContinuous': '선택 영역의 가운데로', 'distributed': '균등 분할', 'justify': '양쪽 맞춤'}
VALIGN_KO = {'top': '위쪽', 'center': '가운데', 'bottom': '아래쪽', 'justify': '양쪽 맞춤', 'distributed': '균등 분할'}


def theme_colors(wb):
    """통합 문서 테마의 색 → {'T0': 'FFFFFF', 'T1': '000000', 'T4': '4F81BD', ...}"""
    xml = getattr(wb, 'loaded_theme', None)
    if not xml:
        return {}
    text = xml.decode('utf-8', 'replace') if isinstance(xml, bytes) else str(xml)
    out = {}
    for i, slot in enumerate(THEME_SLOTS):
        m = re.search(rf'<a:{slot}>(.*?)</a:{slot}>', text, re.S)
        if not m:
            continue
        c = re.search(r'(?:lastClr|val)="([0-9A-Fa-f]{6})"', m.group(1))
        if c:
            out[f'T{i}'] = c.group(1).upper()
    return out


def same_color(got, want, theme=None):
    """got: _color() 값(RRGGBB 또는 T4 같은 테마 색), want: 'RRGGBB'. 같은 색이면 테마 색으로 골라도 맞음."""
    want = want.upper() if want else want
    if got == want:
        return True
    return bool(got and theme and theme.get(got) == want)


def _same_fmt(user_fmt, want_fmt, values=()):
    """표시 형식: 실제 칸 값(있으면)과 양수 예시가 같게 보이면 같은 형식(예: #,###"원" = #,##0"원" — 0 이 없을 때)."""
    if user_fmt == want_fmt:
        return True
    nums = [v for v in values if fx.is_num(v)][:50]
    samples = nums + [1234.5] if nums else (1234.5, 0, -56.78, 0.256, 46000)
    try:
        return all(fx.format_value(v, user_fmt).strip() == fx.format_value(v, want_fmt).strip() for v in samples)
    except Exception:  # noqa: BLE001 — 계산기가 모르는 서식
        return _fmt_sample(user_fmt) == _fmt_sample(want_fmt)


def _ws(wb, name):
    for ws in wb.worksheets:
        if ws.title.strip() == name.strip():
            return ws
    return None


def _sheet_range(text, default):
    """'시트!A1:B2' 또는 'A1:B2' → (시트, 범위)"""
    m = re.match(r"^(?:'([^']+)'|([^!]+))!(.+)$", text)
    if m:
        return (m.group(1) or m.group(2)), m.group(3).replace('$', '')
    return default, text.replace('$', '')


class Ctx:
    """채점 한 번에 쓰는 자료(파일 두 벌, 계산기 통합문서, 기대 값용 원본)."""

    def __init__(self, exam, data, filename):
        self.exam = exam
        self.data = data
        self.filename = filename
        self.wb_f = xlsx.load(data, data_only=False)
        self.wb_v = xlsx.load(data, data_only=True)
        self.book = xlsx.to_book(self.wb_f, self.wb_v, today=TODAY)
        self.orig = original_book(exam)
        self._vba = None

    def ws(self, name):
        ws = _ws(self.wb_f, name)
        if ws is None:
            raise MissingSheet(name)
        return ws

    def user_value(self, sheet, r, c):
        try:
            return self.book.sheet(sheet).get(r, c)
        except fx.XLErr:
            return None

    @property
    def theme(self):
        if not hasattr(self, '_theme'):
            self._theme = theme_colors(self.wb_f)
        return self._theme

    @property
    def vba(self):
        if self._vba is None:
            srcs = vba.sources(self.data)
            self._vba = {'sources': srcs, 'procs': vba.procedures(srcs), 'buttons': vba.buttons(self.data),
                         'cells': vba.button_cells(self.data)}
        return self._vba


class MissingSheet(Exception):
    pass


# ------------------------------------------------------------ 항목별 채점 ---
def check_values(ctx, sheet, chk):
    ws = ctx.ws(sheet)
    cells = _cells(chk['range'])
    exp_rows = chk['expect']
    r0, c0 = cells[0]
    bad = []
    for i, row in enumerate(exp_rows):
        for j, e in enumerate(row):
            r, c = r0 + i, c0 + j
            got = ctx.user_value(ws.title, r, c)
            want = _val(e)
            if fx.is_num(want) and isinstance(got, str) and got.strip():
                bad.append(f'{fx.addr(r, c)}: "{got}" 가 텍스트로 입력되었습니다 — 숫자로(Val·CLng 등) 넣어야 합니다')
            elif not _same(got, want):
                bad.append(f'{fx.addr(r, c)}: {fx.display(got) if got is not None else "(빈칸)"} → {fx.display(want)}')
    return not bad, bad[:4]


def _expected_formula(ctx, sheet, chk):
    """원본 통합문서에 정답 수식을 채워 기대 값 계산(앞 항목 정답이 뒤 항목에 쓰일 수 있게 차례로)."""
    book = ctx.orig
    ast = fx.parse(chk['answer'])
    cells = _cells(chk['range'])
    r0, c0 = cells[0]
    out = []
    s = book.sheet(sheet)
    for r, c in cells:
        v = fx.evaluate(fx.shift(ast, r - r0, c - c0), book, s.name, r, c)
        if isinstance(v, fx.Arr):
            v = v.rows[0][0]
        s.set(r, c, v)
        out.append(((r, c), v))
    return out


def check_formula(ctx, sheet, chk):
    ws = ctx.ws(sheet)
    bad = []
    need = [n.upper() for n in chk.get('require', [])]
    forbid = [n.upper() for n in chk.get('forbid', [])]
    for (r, c), exp in _expected_formula(ctx, sheet, chk):
        raw = ws.cell(r, c).value
        text = xlsx.formula_text(raw)
        got = ctx.user_value(ws.title, r, c)
        a = fx.addr(r, c)
        if text is None:
            if chk.get('allow_value') and _same(got, exp):
                continue
            bad.append(f'{a}: 수식이 아닙니다' if raw is not None else f'{a}: 비어 있음')
            continue
        try:
            used = fx.functions_used(fx.parse(text))
        except fx.FormulaError:
            used = None
        if used is None:                     # 이 채점기가 못 읽는 수식(표 이름 참조 등) — 함수 검사 대신 값만
            used = {n.upper() for n in re.findall(r'([A-Za-z][A-Za-z0-9.]*)\s*\(', text)}
        same_as = {'RANK.EQ': 'RANK', 'QUARTILE.INC': 'QUARTILE', 'PERCENTILE.INC': 'PERCENTILE',
                   'STDEV.S': 'STDEV', 'VAR.S': 'VAR', 'MODE.SNGL': 'MODE'}   # 예전 이름도 같은 함수
        miss = [n for n in need if n not in used and same_as.get(n) not in used]
        if miss:
            bad.append(f'{a}: {", ".join(miss)} 함수를 써야 합니다')
            continue
        if any(n in used for n in forbid):
            bad.append(f'{a}: {", ".join(n for n in forbid if n in used)} 함수는 쓰지 않아야 합니다')
            continue
        if chk.get('array') and not isinstance(raw, ArrayFormula):
            bad.append(f'{a}: 배열 수식(Ctrl+Shift+Enter)으로 입력해야 합니다')
            continue
        if not _same(got, exp, chk.get('tol', 1e-6)):
            bad.append(f'{a}: {fx.display(got) if got is not None else "(값 없음)"} → {fx.display(exp)}')
    return not bad, bad[:4]


def _fmt_sample(fmt):
    out = []
    for v in (1234.5, 0, -56.78, 0.256, 46000):
        try:
            out.append(fx.format_value(v, fmt))
        except Exception:  # noqa: BLE001 — 계산기가 모르는 서식은 글자 그대로 비교
            out.append(fmt)
    return out


def check_style(ctx, sheet, chk):
    ws = ctx.ws(sheet)
    bad = []
    cells = _cells(chk['range'])
    if chk.get('merge'):
        if not any(str(m) == chk['range'] for m in ws.merged_cells.ranges):
            bad.append(f"{chk['range']} 을 병합해야 합니다")
    f = chk.get('font') or {}
    for r, c in cells[:1] if chk.get('merge') else cells:
        cell = ws.cell(r, c)
        a = fx.addr(r, c)
        if 'name' in f and (cell.font.name or '') != f['name']:
            bad.append(f"{a}: 글꼴 {cell.font.name} → {f['name']}")
        if 'size' in f and float(cell.font.sz or 11) != float(f['size']):
            bad.append(f"{a}: 크기 {cell.font.sz} → {f['size']}")
        if 'bold' in f and bool(cell.font.b) != f['bold']:
            bad.append(f"{a}: 굵게 {'해제' if not f['bold'] else ''}")
        if 'italic' in f and bool(cell.font.i) != f['italic']:
            bad.append(f'{a}: 기울임꼴')
        if 'underline' in f and (cell.font.u or None) != (f['underline'] or None):
            bad.append(f"{a}: 밑줄({f['underline']})")
        if 'color' in f and not same_color(_color(cell.font.color) or ('000000' if f['color'].upper() == '000000' else None),
                                           f['color'], ctx.theme):
            bad.append(f"{a}: 글꼴 색")
        if 'fill' in chk:
            got = _color(cell.fill.fgColor) if cell.fill.patternType == 'solid' else None
            if not same_color(got, chk['fill'], ctx.theme):
                bad.append(f'{a}: 채우기 색')
        if 'halign' in chk and (cell.alignment.horizontal or 'general') != chk['halign']:
            have = HALIGN_KO.get(cell.alignment.horizontal or 'general', cell.alignment.horizontal)
            bad.append(f"{a}: 가로 맞춤 '{have}' → '{HALIGN_KO.get(chk['halign'], chk['halign'])}'")
        if 'valign' in chk and (cell.alignment.vertical or 'bottom') != chk['valign']:
            bad.append(f"{a}: 세로 맞춤 → '{VALIGN_KO.get(chk['valign'], chk['valign'])}'")
        if 'wrap' in chk and bool(cell.alignment.wrap_text) != chk['wrap']:
            bad.append(f'{a}: 자동 줄 바꿈')
        if 'numfmt' in chk and not _same_fmt(cell.number_format, chk['numfmt'],
                                             [ctx.user_value(ws.title, rr, cc) for rr, cc in cells]):
            bad.append(f"{a}: 표시 형식 {cell.number_format} → {chk['numfmt']}")
        if chk.get('border') == 'all':
            b = cell.border
            if not all(getattr(b, s).style for s in ('left', 'right', 'top', 'bottom')):
                bad.append(f'{a}: 모든 테두리')
        if len(bad) > 6:
            break
    if 'height' in chk:
        rows = sorted({r for r, _ in cells})
        if any(abs((ws.row_dimensions[r].height or 15) - chk['height']) > 0.5 for r in rows):
            bad.append(f"행 높이 {chk['height']}")
    return not bad, bad[:4]


def check_comment(ctx, sheet, chk):
    ws = ctx.ws(sheet)
    cm = ws[chk['cell']].comment
    if cm is None:
        return False, [f"{chk['cell']} 에 메모가 없습니다"]
    text = cm.text or ''
    body = text.split(':\n', 1)[1] if ':\n' in text[:40] else text      # '작성자:\n내용'
    ok = _norm_text(chk['text']) in _norm_text(body) or _norm_text(chk['text']) == _norm_text(text)
    return ok, [] if ok else [f"메모 내용 '{_norm_text(body)[:20]}' → '{chk['text']}'"]


def check_name(ctx, sheet, chk):
    names = dict(xlsx._defined_names(ctx.wb_f))
    for n, ref in names.items():
        if n.lower() == chk['name'].lower():
            want = _sheet_range(chk['ref'], sheet)
            got = _sheet_range(ref.lstrip('='), sheet)
            ok = want[0].strip("'") == got[0].strip("'") and want[1].upper() == got[1].upper()
            return ok, [] if ok else [f"이름 '{n}' 의 범위 {ref} → {chk['ref']}"]
    return False, [f"이름 '{chk['name']}' 이 정의되어 있지 않습니다"]


def _row_mask(ctx, sheet, rng, formula_text, book=None):
    """조건 수식(첫 행 기준)을 범위 각 행에 옮겨 계산 → [True/False] (조건부 서식·고급 필터 기대 값)."""
    book = book or ctx.orig
    r1, c1, r2, _ = fx.parse_range(rng)
    ast = fx.parse(formula_text)
    out = []
    for k, r in enumerate(range(r1, r2 + 1)):
        v = fx.evaluate(fx.shift(ast, k, 0), book, sheet, r, c1)
        try:
            out.append(fx.to_bool(v))
        except fx.XLErr:
            out.append(False)
    return out


def _cell_mask(ctx, sheet, rng, formula_text, book=None):
    """조건부 서식 규칙을 범위의 모든 칸에 옮겨 계산(엑셀은 칸마다 상대 참조를 옮겨 적용) → [[True/False]]"""
    book = book or ctx.orig
    r1, c1, r2, c2 = fx.parse_range(rng)
    ast = fx.parse(formula_text)
    out = []
    for r in range(r1, r2 + 1):
        row = []
        for c in range(c1, c2 + 1):
            v = fx.evaluate(fx.shift(ast, r - r1, c - c1), book, sheet, r, c)
            try:
                row.append(fx.to_bool(v))
            except fx.XLErr:
                row.append(False)
        out.append(row)
    return out


def check_cf(ctx, sheet, chk):
    ws = ctx.ws(sheet)
    want_rng = chk['range'].replace('$', '').upper()
    cands = []
    for cf in ws.conditional_formatting:
        sq = str(cf.sqref).replace('$', '').upper()
        for rule in cf.rules:
            cands.append((sq, rule))
    if not cands:
        return False, ['조건부 서식이 없습니다']
    msgs = []
    want_mask = _cell_mask(ctx, sheet, want_rng, chk['formula'])
    for sq, rule in cands:
        if sq != want_rng:
            msgs.append(f'적용 범위 {sq} → {want_rng}')
            continue
        if rule.type != 'expression' or not rule.formula:
            msgs.append('수식을 사용하는 규칙이어야 합니다')
            continue
        try:
            got_mask = _cell_mask(ctx, sheet, want_rng, '=' + rule.formula[0], book=ctx.book)
        except fx.FormulaError:
            msgs.append(f'규칙 수식을 읽을 수 없습니다: {rule.formula[0]}')
            continue
        if got_mask != want_mask:
            msgs.append(f'규칙 수식 ={rule.formula[0]} 이 서식을 줄 행이 다릅니다($ 고정 확인)')
            continue
        dxf = rule.dxf
        fmt_bad = []
        if 'font_color' in chk and not same_color(_color(dxf.font.color if dxf and dxf.font else None),
                                                  chk['font_color'], ctx.theme):
            fmt_bad.append('글꼴 색')
        if 'bold' in chk and bool(dxf and dxf.font and dxf.font.b) != chk['bold']:
            fmt_bad.append('굵게')
        if 'italic' in chk and bool(dxf and dxf.font and dxf.font.i) != chk['italic']:
            fmt_bad.append('기울임꼴')
        if 'fill' in chk:
            fill = dxf.fill if dxf else None
            got = _color(fill.bgColor) or _color(fill.fgColor) if fill else None
            if not same_color(got, chk['fill'], ctx.theme):
                fmt_bad.append('채우기 색')
        if fmt_bad:
            msgs.append('서식이 다릅니다: ' + ', '.join(fmt_bad))
            continue
        return True, []
    return False, msgs[:3]


def check_dv(ctx, sheet, chk):
    ws = ctx.ws(sheet)
    want = chk['range'].replace('$', '').upper()
    fields = {'type': 'type', 'operator': 'operator', 'formula1': 'formula1', 'formula2': 'formula2',
              'error_title': 'errorTitle', 'error': 'error', 'prompt_title': 'promptTitle', 'prompt': 'prompt',
              'style': 'errorStyle'}
    for d in ws.data_validations.dataValidation:
        if str(d.sqref).replace('$', '').upper() != want:
            continue
        bad = []
        for key, attr in fields.items():
            if key not in chk:
                continue
            got = getattr(d, attr)
            exp = chk[key]
            if key == 'operator' and got is None:
                got = 'between'
            if key == 'style' and got is None:
                got = 'stop'
            if key in ('formula1', 'formula2'):
                got = (got or '').lstrip('=').replace('$', '').strip('"')
                exp = str(exp).lstrip('=').replace('$', '').strip('"')
            if _norm_text(got) != _norm_text(exp):
                bad.append(f'{key}: {got} → {exp}')
        if chk.get('show_error', True) is True and d.showErrorMessage is False:
            bad.append('오류 메시지 표시')
        return not bad, bad[:4]
    return False, [f'{want} 에 유효성 검사가 없습니다']


def _block(ws_book, sheet, top_left):
    """top_left 부터 빈 행이 나올 때까지의 표 → 행 목록(첫 행은 머리글)."""
    r0, c0 = fx.parse_addr(top_left)
    s = ws_book.sheet(sheet)
    width = 0
    while s.get(r0, c0 + width) not in (None, ''):
        width += 1
    rows = []
    r = r0
    while any(s.get(r, c0 + j) not in (None, '') for j in range(width)) and r < r0 + 5000:
        rows.append([s.get(r, c0 + j) for j in range(width)])
        r += 1
    return rows


def check_filter(ctx, sheet, chk):
    src = chk['source']
    r1, c1, r2, c2 = fx.parse_range(src)
    s = ctx.orig.sheet(sheet)
    head = [s.get(r1, c) for c in range(c1, c2 + 1)]
    mask = _row_mask(ctx, sheet, fx.addr(r1 + 1, c1) + ':' + fx.addr(r2, c1), chk['answer'])
    cols = chk.get('columns') or head
    idx = [[_norm_text(h) for h in head].index(_norm_text(cn)) for cn in cols]
    exp = [cols] + [[s.get(r, c1 + j) for j in idx] for r, ok in zip(range(r1 + 1, r2 + 1), mask) if ok]
    got = _block(ctx.book, ctx.ws(sheet).title, chk['out'])
    bad = []
    if chk.get('criteria_at'):
        cr, cc = fx.parse_addr(chk['criteria_at'])
        if ctx.user_value(sheet, cr, cc) in (None, ''):
            bad.append(f"조건을 {chk['criteria_at']} 셀부터 입력해야 합니다")
    if not got:
        return False, bad + [f"{chk['out']} 에 결과가 없습니다"]
    if [_norm_text(x) for x in got[0]] != [_norm_text(x) for x in exp[0]]:
        bad.append('결과 머리글(추출할 필드)이 다릅니다: ' + ', '.join(_norm_text(x) for x in got[0]))
    elif len(got) != len(exp):
        bad.append(f'추출된 행 {len(got) - 1}개 → {len(exp) - 1}개')
    elif any(not _same(a, b) for gr, er in zip(got, exp) for a, b in zip(gr, er)):
        bad.append('추출된 내용이 다릅니다')
    return not bad, bad[:3]


def _sort_key(v):
    rank = 0 if fx.is_num(v) and not isinstance(v, bool) else (1 if isinstance(v, str) else 2)
    return rank, (v.lower() if isinstance(v, str) else (v if v is not None else 0))


def check_sorted(ctx, sheet, chk):
    r1, c1, r2, c2 = fx.parse_range(chk['range'])
    s0 = ctx.orig.sheet(sheet)
    head = [_norm_text(s0.get(r1, c)) for c in range(c1, c2 + 1)]
    rows = [[s0.get(r, c) for c in range(c1, c2 + 1)] for r in range(r1 + 1, r2 + 1)]
    for name, order in reversed(chk['keys']):
        k = head.index(_norm_text(name))
        if isinstance(order, list):
            pos = {v: i for i, v in enumerate(order)}
            rows.sort(key=lambda row: pos.get(row[k], len(pos)))
        else:
            rows.sort(key=lambda row: _sort_key(row[k]), reverse=order == 'desc')
    got = [[ctx.user_value(sheet, r, c) for c in range(c1, c2 + 1)] for r in range(r1 + 1, r2 + 1)]
    for i, (a, b) in enumerate(zip(got, rows)):
        if any(not _same(x, y) for x, y in zip(a, b)):
            return False, [f'{r1 + 1 + i}행부터 정렬 순서가 다릅니다']
    return True, []


SUBTOTAL_NO = {'sum': 9, 'count': 3, 'average': 1, 'max': 4, 'min': 5, 'product': 6, 'counta': 3, 'countnums': 2}
SUBTOTAL_NAME = {9: '합계', 3: '개수', 1: '평균', 4: '최대값', 5: '최소값', 6: '곱', 2: '숫자 개수'}


def check_subtotal(ctx, sheet, chk):
    """부분합: 그룹 기준 정렬 + 요구한 함수·필드의 SUBTOTAL 수식이 각 그룹마다 있고 값이 맞는지."""
    ws = ctx.ws(sheet)
    r1, c1, r2, c2 = fx.parse_range(chk['range'])
    s0 = ctx.orig.sheet(sheet)
    head = [_norm_text(s0.get(r1, c)) for c in range(c1, c2 + 1)]
    gcol = c1 + head.index(_norm_text(chk['group']))
    groups = []
    for r in range(r1 + 1, r2 + 1):
        g = s0.get(r, gcol)
        if g not in groups:
            groups.append(g)
    # 사용자 시트에서 SUBTOTAL 수식 모으기
    found = {}
    for row in ws.iter_rows(min_row=r1, max_row=ws.max_row):
        for cell in row:
            text = xlsx.formula_text(cell.value)
            if text and 'SUBTOTAL' in text.upper():
                m = re.search(r'SUBTOTAL\(\s*(\d+)\s*,\s*\$?([A-Z]+)\$?(\d+)\s*:\s*\$?[A-Z]+\$?(\d+)', text.upper())
                if m:
                    found.setdefault((int(m.group(1)) % 100, fx.col_num(m.group(2))), []).append(
                        (cell.row, int(m.group(3)), int(m.group(4))))
    bad = []

    def group_rows(no, col):                          # 그룹 하나를 묶은 부분합 칸(전체 합계 칸 제외)의 (행, 첫 데이터 행)
        return sorted((row, a) for row, a, b in found.get((no, col), []) if b - a + 1 < (r2 - r1))
    if chk.get('ordered') and chk['items']:           # 정렬: 그룹이 오름차순으로 이어져 있어야 한다
        it0 = chk['items'][0]
        keys = [ctx.user_value(ws.title, a, gcol) for _, a in group_rows(SUBTOTAL_NO[it0['func']],
                                                                       c1 + head.index(_norm_text(it0['fields'][0])))]
        keys = [k for k in keys if k is not None]
        if keys and keys != sorted(keys, key=lambda k: (not fx.is_num(k), str(k) if not fx.is_num(k) else k)):
            bad.append(f"{chk['group']} 그룹이 오름차순이 아닙니다({', '.join(str(k) for k in keys)}) — 먼저 오름차순 정렬하세요")
    if chk.get('above'):                               # 적용 순서: 나중에 만든 부분합 줄이 앞에 만든 줄 위에 있어야 한다
        ab = chk['above']
        mine = group_rows(SUBTOTAL_NO[chk['items'][0]['func']], c1 + head.index(_norm_text(chk['items'][0]['fields'][0])))
        other = group_rows(SUBTOTAL_NO[ab['func']], c1 + head.index(_norm_text(ab['field'])))
        if mine and other and len(mine) == len(other) and any(m[0] > o[0] for m, o in zip(mine, other)):
            bad.append('부분합을 적용한 순서가 문제와 다릅니다 — 문제에 적힌 순서대로 만들고 두 번째는 "새로운 값으로 대치"를 해제하세요')
    for item in chk['items']:
        no = SUBTOTAL_NO[item['func']]
        for field in item['fields']:
            col = c1 + head.index(_norm_text(field))
            cells = found.get((no, col), [])
            if len(cells) < len(groups):
                bad.append(f"{field} {SUBTOTAL_NAME[no]} 부분합이 {len(cells)}곳 — 그룹 {len(groups)}개(+전체)여야 합니다")
                continue
            # 각 부분합 칸이 한 그룹 행만 묶었는지
            for row, a, b in cells:
                keys = {ctx.user_value(ws.title, rr, gcol) for rr in range(a, b + 1)
                        if not xlsx.formula_text(ws.cell(rr, col).value)}
                keys.discard(None)
                if len(keys) > 1 and b - a + 1 < (r2 - r1):
                    bad.append(f'{fx.addr(row, col)}: 여러 그룹이 섞였습니다 — 먼저 {chk["group"]} 기준으로 정렬하세요')
                    break
            else:
                # 그룹마다 원래 자료로 계산한 값이 부분합 칸 중 하나와 같아야 한다
                got = [ctx.user_value(ws.title, row, col) for row, _, _ in cells]
                for g in groups:
                    vals = [s0.get(r, col) for r in range(r1 + 1, r2 + 1) if s0.get(r, gcol) == g]
                    exp = _subtotal_value(no, vals)
                    hit = next((i for i, v in enumerate(got) if exp is not None and _same(v, exp)), None)
                    if hit is None:
                        bad.append(f"{field} {SUBTOTAL_NAME[no]}: '{g}' 그룹 값 {fx.display(exp)} 이 없습니다")
                        break
                    got.pop(hit)
    return not bad, bad[:3]


def _subtotal_value(no, vals):
    nums = [v for v in vals if fx.is_num(v)]
    if no == 3:
        return sum(1 for v in vals if v not in (None, ''))
    if no == 2:
        return len(nums)
    if no == 9:
        return sum(nums)
    if not nums:
        return None
    if no == 1:
        return sum(nums) / len(nums)
    if no == 4:
        return max(nums)
    if no == 5:
        return min(nums)
    if no == 6:
        out = 1
        for v in nums:
            out *= v
        return out
    return None


def check_pivot(ctx, sheet, chk):
    res = pivots.check({k: v for k, v in chk.items() if k not in ('kind', 'points', 'label')}, ctx.wb_f, ctx.wb_v,
                       ctx.book)
    return res['ok'], res['msgs']


def check_goalseek(ctx, sheet, chk):
    ws = ctx.ws(sheet)
    r, c = fx.parse_addr(chk['changing'])
    if xlsx.formula_text(ws.cell(r, c).value):
        return False, [f"{chk['changing']} 는 값이어야 합니다(목표값 찾기가 바꾸는 셀)"]
    tr, tc = fx.parse_addr(chk['cell'])
    got = ctx.user_value(ws.title, tr, tc)
    tol = chk.get('tol', max(0.5, abs(chk['value']) * 1e-4))
    ok = fx.is_num(got) and abs(got - chk['value']) <= tol
    return ok, [] if ok else [f"{chk['cell']} = {fx.display(got)} → {chk['value']} 이 되어야 합니다"]


def check_scenario(ctx, sheet, chk):
    ws = ctx.ws(sheet)
    sc = getattr(ws, 'scenarios', None)
    have = {s.name: s for s in (sc.scenario if sc else [])}
    bad = []
    order = [a.replace('$', '').upper() for a in _expand(chk['changing'])]
    want_cells = sorted(order)
    for name, vals in chk['scenarios']:
        s = have.get(name)
        if s is None:
            bad.append(f"시나리오 '{name}' 이 없습니다")
            continue
        cells = [(ic.r.replace('$', '').upper(), ic.val) for ic in s.inputCells]
        if sorted(a for a, _ in cells) != want_cells:
            bad.append(f"'{name}' 의 변경 셀 {', '.join(a for a, _ in cells)} → {chk['changing']}")
            continue
        got = dict(cells)
        for a, v in zip(order if len(vals) == len(order) else [], vals):
            try:
                g = float(got[a])
            except (TypeError, ValueError):
                g = got[a]
            if not _same(g, v):
                bad.append(f"'{name}' {a} 값 {got[a]} → {v}")
    if chk.get('summary', True):
        names = [w.title for w in ctx.wb_f.worksheets]
        summ = [n for n in names if '시나리오 요약' in n]
        if not summ:
            bad.append("'시나리오 요약' 시트가 없습니다")
        elif chk.get('summary_before') and chk['summary_before'] in names and \
                names.index(summ[0]) != names.index(chk['summary_before']) - 1:
            bad.append(f"'시나리오 요약' 시트는 '{chk['summary_before']}' 시트 바로 앞에 있어야 합니다")
    return not bad, bad[:3]


def _expand(ref):
    out = []
    for part in ref.split(','):
        for r, c in _cells(part.strip().replace('$', '')):
            out.append(fx.addr(r, c))
    return out


def check_datatable(ctx, sheet, chk):
    ws = ctx.ws(sheet)
    r1, c1, r2, c2 = fx.parse_range(chk['range'])        # 결과 칸(머리 행·열 제외)
    master = ws.cell(r1, c1).value
    if type(master).__name__ != 'DataTableFormula':
        return False, [f"{fx.addr(r1, c1)} 에 데이터 표가 없습니다([가상 분석] > [데이터 표])"]
    bad = []
    two = str(master.dt2D) in ('1', 'True')
    if chk.get('row_input') and chk.get('col_input'):
        if not two:
            bad.append('행 입력 셀과 열 입력 셀을 모두 지정해야 합니다')
        elif (master.r1 or '').replace('$', '').upper() != chk['row_input'] or \
                (master.r2 or '').replace('$', '').upper() != chk['col_input']:
            bad.append(f"입력 셀: 행 {master.r1}, 열 {master.r2} → 행 {chk['row_input']}, 열 {chk['col_input']}")
    else:
        want = (chk.get('row_input') or chk.get('col_input')).upper()
        if (master.r1 or '').replace('$', '').upper() != want:
            bad.append(f'입력 셀 {master.r1} → {want}')
    if bad:
        return False, bad
    # 값: 엑셀이 저장한 결과를 원본 수식으로 다시 계산한 기대 값과 비교
    corner_r, corner_c = r1 - 1, c1 - 1
    s0 = ctx.orig.sheet(sheet)
    wv = ctx.wb_v[ws.title]
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            book = original_book(ctx.exam)
            sh = book.sheet(sheet)
            if chk.get('row_input') and chk.get('col_input'):
                rr, rc = fx.parse_addr(chk['row_input'])
                cr_, cc_ = fx.parse_addr(chk['col_input'])
                sh.set(rr, rc, s0.get(corner_r, c))
                sh.set(cr_, cc_, s0.get(r, corner_c))
                base = (corner_r, corner_c)
            elif chk.get('col_input'):
                cr_, cc_ = fx.parse_addr(chk['col_input'])
                sh.set(cr_, cc_, s0.get(r, corner_c))
                base = (corner_r, c)
            else:
                rr, rc = fx.parse_addr(chk['row_input'])
                sh.set(rr, rc, s0.get(corner_r, c))
                base = (r, corner_c)
            corner = s0.raw(*base)
            if not isinstance(corner, fx.Formula):
                corner = ctx.book.sheet(ws.title).raw(*base)
            if not isinstance(corner, fx.Formula):
                return False, [f'{fx.addr(*base)} 에 계산 수식이 있어야 합니다']
            exp = fx.evaluate(corner.ast, book, sheet, *base)
            got = wv.cell(r, c).value
            if got is not None and not _same(got, exp, 1e-6):
                return False, [f'{fx.addr(r, c)}: {fx.display(got)} → {fx.display(exp)}']
    return True, []


FILES = Path(__file__).resolve().parent.parent / 'content' / 'exams' / 'files'


def problem_file(exam):
    """내려받을 문제 파일 → (바이트, 확장자). 폼·ActiveX 단추가 있는 1급은 Excel 로 만든 .xlsm(tools/exam_problem_files.py)."""
    p = FILES / f"{exam['id']}.xlsm"
    if (exam.get('forms') or exam.get('commands')) and p.exists():
        return p.read_bytes(), 'xlsm'
    return problem_workbook(exam), 'xlsx'


def data_file(exam, name):
    """외부 데이터 자료(csv) → UTF-8(BOM) 바이트. 없으면 None."""
    spec = (exam.get('data_files') or {}).get(name)
    if not spec:
        return None
    import csv as _csv
    bio = io.StringIO()
    w = _csv.writer(bio, lineterminator='\r\n')
    for row in spec['rows']:
        w.writerow(['' if v is None else (v[1:] if isinstance(v, str) and v.startswith('@') else v) for v in row])
    return bio.getvalue().encode('utf-8-sig')


def _key_name(n):
    return re.sub(r'[\s-]', '', n)


CHART_TYPES = {'barChart': 'col', 'bar3DChart': 'col', 'lineChart': 'line', 'line3DChart': 'line', 'pieChart': 'pie',
               'pie3DChart': 'pie', 'doughnutChart': 'doughnut', 'areaChart': 'area', 'scatterChart': 'scatter',
               'radarChart': 'radar'}


def _ref_text(ref, book):
    """셀에 연결된 제목·계열 이름(strRef) → 그 셀의 글자."""
    if book is not None and ref.f:
        try:
            v = fx.evaluate(fx.parse('=' + ref.f), book, book.sheets and next(iter(book.sheets)), 1, 1)
            if not isinstance(v, fx.XLErr):
                return _norm_text(fx.display(v))
        except (fx.FormulaError, fx.XLErr):
            pass
    cache = ref.strCache
    if cache is not None and cache.pt:
        return _norm_text(cache.pt[0].v)
    return _norm_text(ref.f)


def _rich_text(t, book=None):
    if t is None:
        return None
    tx = getattr(t, 'tx', None)
    if tx is not None and getattr(tx, 'strRef', None) is not None:
        return _ref_text(tx.strRef, book)       # 셀과 연결한 제목(예: =차트작업!$A$1)
    try:
        return ''.join(r.t for p in t.tx.rich.p for r in (p.r or []))
    except AttributeError:
        return ''


def _series_name(s, book):
    if s.tx is None:
        return ''
    if s.tx.strRef is not None:
        return _ref_text(s.tx.strRef, book)
    return _norm_text(getattr(s.tx, 'v', '') or '')


def _axis_title(part, which, book=None):
    axis = getattr(part, which, None)
    return _rich_text(axis.title, book) if axis is not None and axis.title else None


def describe_charts(ctx, sheet, charts=None):
    ws = ctx.ws(sheet)
    out = []
    for ch in (ws._charts if charts is None else charts):
        parts = [ch] + [x for x in getattr(ch, '_charts', []) if x is not ch]
        series, kinds = [], set()
        ax = lambda part: getattr(getattr(part, 'y_axis', None), 'axId', None)  # noqa: E731 — 원형은 축 없음
        first_ax = ax(parts[0]) if parts else None
        for part in parts:
            kind = CHART_TYPES.get(part.tagname, part.tagname)
            if kind == 'col' and getattr(part, 'barDir', 'col') == 'bar':
                kind = 'bar'
            kinds.add(kind)
            for s in part.series:
                name = _series_name(s, ctx.book)
                if any(x['name'] == name for x in series):
                    continue
                series.append({'name': name, 'kind': kind, 'secondary': ax(part) != first_ax,
                               'labels': bool(s.dLbls and (s.dLbls.showVal or s.dLbls.showPercent or s.dLbls.showCatName)),
                               'trend': bool(s.trendline), 'val': s.val.numRef.f if s.val and s.val.numRef else ''})
        out.append({'kinds': kinds, 'title': _rich_text(ch.title, ctx.book), 'series': series,
                    'y_title': _axis_title(parts[0], 'y_axis', ctx.book),
                    'x_title': _axis_title(parts[0], 'x_axis', ctx.book),
                    'legend': ch.legend.position if ch.legend is not None else None,
                    'grouping': getattr(parts[0], 'grouping', None)})
    return out


def check_chart(ctx, sheet, chk):
    charts = describe_charts(ctx, sheet)
    if not charts:
        return False, ['차트가 없습니다']
    ch = charts[min(chk.get('index', 0), len(charts) - 1)]
    bad = []
    names = [s['name'] for s in ch['series']]
    if 'type' in chk:
        main = {s['kind'] for s in ch['series'] if s['name'] not in chk.get('line_series', [])}
        if chk['type'] not in main and not (chk['type'] == 'col' and main == {'col'}):
            bad.append(f"차트 종류: {'/'.join(sorted(main))} → {chk['type']}")
    if 'title' in chk and _norm_text(ch['title']) != _norm_text(chk['title']):
        bad.append(f"차트 제목 '{ch['title'] or ''}' → '{chk['title']}'")
    if 'series' in chk and [_norm_text(n) for n in chk['series']] != names:
        bad.append(f"데이터 계열 {', '.join(names)} → {', '.join(chk['series'])}")
    for n in chk.get('line_series', []):
        s = next((x for x in ch['series'] if x['name'] == n), None)
        if not s or s['kind'] != 'line':
            bad.append(f"'{n}' 계열을 꺾은선형으로")
    for n in chk.get('secondary', []):
        s = next((x for x in ch['series'] if x['name'] == n), None)
        if not s or not s['secondary']:
            bad.append(f"'{n}' 계열을 보조 축으로")
    for n in chk.get('labels', []):
        s = next((x for x in ch['series'] if x['name'] == n), None)
        if not s or not s['labels']:
            bad.append(f"'{n}' 계열에 데이터 레이블")
    for n in chk.get('trend', []):
        s = next((x for x in ch['series'] if x['name'] == n), None)
        if not s or not s['trend']:
            bad.append(f"'{n}' 계열에 추세선")
    if 'y_title' in chk and _norm_text(ch['y_title']) != _norm_text(chk['y_title']):
        bad.append(f"세로 축 제목 → '{chk['y_title']}'")
    if 'x_title' in chk and _norm_text(ch['x_title']) != _norm_text(chk['x_title']):
        bad.append(f"가로 축 제목 → '{chk['x_title']}'")
    if 'legend' in chk and ch['legend'] != chk['legend']:
        pos = {'b': '아래쪽', 't': '위쪽', 'r': '오른쪽', 'l': '왼쪽', None: '없음'}
        bad.append(f"범례 위치 {pos.get(ch['legend'], ch['legend'])} → {pos.get(chk['legend'], chk['legend'])}")
    if 'grouping' in chk and ch['grouping'] != chk['grouping']:
        bad.append(f"차트 종류(누적 여부) → {chk['grouping']}")
    return not bad, bad[:4]


def check_macro(ctx, sheet, chk):
    info = ctx.vba
    if not info['sources']:
        return False, ['매크로가 없습니다 — Excel 매크로 사용 통합 문서(.xlsm)로 저장하세요']
    bad = []
    if chk['name'].lower() not in info['procs']:
        bad.append(f"'{chk['name']}' 매크로가 없습니다")
    if chk.get('button'):
        btn = [(t, m) for t, m in info['buttons'] if _norm_text(t) == _norm_text(chk['button'])]
        if not btn:
            bad.append(f"'{chk['button']}' 단추가 없습니다")
        elif not any(m.lower().split('.')[-1] == chk['name'].lower() for _, m in btn):
            bad.append(f"'{chk['button']}' 단추에 '{chk['name']}' 매크로를 연결하세요")
        elif chk.get('near'):                        # 단추 위치: 지정한 셀(범위)에서 2칸 안
            r1, c1, r2, c2 = fx.parse_range(chk['near']) if ':' in chk['near'] else (*fx.parse_addr(chk['near']),
                                                                                    *fx.parse_addr(chk['near']))
            at = [(r, c) for t, m, r, c in info.get('cells', []) if _norm_text(t) == _norm_text(chk['button'])]
            if at and not any(r1 - 2 <= r <= r2 + 2 and c1 - 2 <= c <= c2 + 2 for r, c in at):
                bad.append(f"'{chk['button']}' 단추가 {chk['near']} 근처가 아니라 {fx.addr(*at[0])} 에 있습니다")
    body = info['procs'].get(chk['name'].lower())
    for pat in chk.get('patterns', []):               # 매크로 코드에 지시한 일이 들어 있는지(결과만 손으로 넣은 경우 걸러냄)
        if body is not None and not re.search(pat, re.sub(r"'.*", '', body), re.I):
            bad.append(f"'{chk['name']}' 매크로 코드에 지시한 내용이 없습니다({pat.split('|')[0]} …)")
            break
    return not bad, bad


def check_vba(ctx, sheet, chk):
    info = ctx.vba
    if not info['sources']:
        return False, ['VBA 코드가 없습니다 — .xlsm 으로 저장하세요']
    body = info['procs'].get(chk['proc'].lower())
    if body is None:
        return False, [f"프로시저 '{chk['proc']}' 가 없습니다"]
    flat = re.sub(r'\s+', ' ', re.sub(r"'.*", '', body))
    miss = [p for p in chk.get('patterns', []) if not re.search(p, flat, re.I)]
    return not miss, [f'필요한 코드가 없습니다: {p}' for p in miss[:2]]


def check_page(ctx, sheet, chk):
    ws = ctx.ws(sheet)
    bad = []
    if 'orientation' in chk and ws.page_setup.orientation != chk['orientation']:
        bad.append(f"용지 방향 → {'가로' if chk['orientation'] == 'landscape' else '세로'}")
    if chk.get('center_h') and not ws.print_options.horizontalCentered:
        bad.append('페이지 가운데 맞춤(가로)')
    if chk.get('center_v') and not ws.print_options.verticalCentered:
        bad.append('페이지 가운데 맞춤(세로)')
    if 'print_area' in chk:
        got = (ws.print_area or '').split('!')[-1].replace('$', '')
        if got.upper() != chk['print_area'].upper():
            bad.append(f"인쇄 영역 {got or '없음'} → {chk['print_area']}")
    if 'title_rows' in chk and (ws.print_title_rows or '').replace('$', '') != chk['title_rows'].replace('$', ''):
        bad.append(f"인쇄 제목(반복할 행) → {chk['title_rows']}")
    for key, part in (('header_center', ws.oddHeader.center), ('footer_center', ws.oddFooter.center),
                      ('header_right', ws.oddHeader.right), ('footer_right', ws.oddFooter.right)):
        if key in chk and _norm_text(part.text) != _norm_text(chk[key]):
            bad.append(f"{'머리글' if 'header' in key else '바닥글'} → {chk[key]}")
    fit = bool(ws.sheet_properties.pageSetUpPr and ws.sheet_properties.pageSetUpPr.fitToPage)
    if 'fit_width' in chk and (not fit or (ws.page_setup.fitToWidth or 1) != chk['fit_width']):
        bad.append(f"한 페이지에 맞춤(너비 {chk['fit_width']})")
    return not bad, bad[:4]


def check_protect(ctx, sheet, chk):
    ws = ctx.ws(sheet)
    bad = []
    if not ws.protection.sheet:
        bad.append('시트 보호가 되어 있지 않습니다')
    for r, c in _cells(chk['unlocked']) if chk.get('unlocked') else []:
        if ws.cell(r, c).protection.locked:
            bad.append(f"{chk['unlocked']} 의 '잠금'을 해제해야 합니다")
            break
    for r, c in _cells(chk['hidden']) if chk.get('hidden') else []:
        if not ws.cell(r, c).protection.hidden:
            bad.append(f"{chk['hidden']} 에 '숨김'을 지정해야 합니다")
            break
    for opt, attr in (('allow_select_locked', 'selectLockedCells'), ('allow_format_cells', 'formatCells')):
        if opt in chk and (not getattr(ws.protection, attr)) != chk[opt]:
            bad.append(f'보호 옵션 {opt}')
    return not bad, bad[:3]


CHECKS = {'values': check_values, 'formula': check_formula, 'style': check_style, 'comment': check_comment,
          'name': check_name, 'cf': check_cf, 'dv': check_dv, 'filter': check_filter, 'sorted': check_sorted,
          'subtotal': check_subtotal, 'pivot': check_pivot, 'goalseek': check_goalseek, 'scenario': check_scenario,
          'datatable': check_datatable, 'chart': check_chart, 'macro': check_macro, 'vba': check_vba,
          'page': check_page, 'protect': check_protect}


# ------------------------------------------------------------ 채점 ----------
def _same_origin(ctx):
    """문제 파일의 글자 상수 중 40% 이상이 수험자 파일의 같은 시트에 남아 있는지(다른 시험의 파일이면 거의 없음)."""
    total = kept = 0
    for name, sh in ctx.orig.sheets.items():
        wu = _ws(ctx.wb_f, name)
        have = set()
        if wu is not None:
            for row in wu.iter_rows():
                for c in row:
                    if isinstance(c.value, str):
                        have.add(c.value.strip())
        for v in sh.cells.values():
            if isinstance(v, str) and len(v.strip()) >= 2 and not v.startswith('='):
                total += 1
                kept += v.strip() in have
    return total < 8 or kept / total >= 0.4


def _grade_unlimited(exam, data, filename='답안.xlsx'):
    ctx = Ctx(exam, data, filename)
    if not _same_origin(ctx):
        raise xlsx.BadFile(f"이 모의고사({exam['title']})의 문제 파일이 아닌 것 같습니다 — 문제 파일의 글자가 거의 남아 있지 않습니다. "
                           '이 문제지에서 받은 파일에 답을 넣어 올려 주세요.')
    want = [s['name'] for s in exam['sheets']]
    have = {_key_name(n) for n in ctx.wb_f.sheetnames}
    if sum(1 for n in want if _key_name(n) in have) * 2 < len(want):
        raise xlsx.BadFile(f"이 모의고사({exam['title']})의 답안 파일이 아닌 것 같습니다 — 시트 "
                           f"{', '.join(want[:4])}… 가 없습니다. 이 문제지에서 받은 파일에 답을 넣어 올려 주세요.")
    tasks, score = [], 0
    for t in exam['tasks']:
        items, got = [], 0
        for chk in t['checks']:
            sheet = chk.get('sheet', t['sheet'])
            try:
                ok, msgs = CHECKS[chk['kind']](ctx, sheet, chk)
            except MissingSheet as e:
                ok, msgs = False, [f"'{e.args[0]}' 시트가 없습니다 — 문제 파일의 시트 이름을 바꾸지 마세요"]
            except (fx.FormulaError, fx.XLErr, KeyError, ValueError, IndexError, AttributeError, TypeError) as e:
                ok, msgs = False, [f'확인 중 문제가 생겼습니다: {type(e).__name__} {e}'[:160]]
            pts = chk.get('points', 0)
            got += pts if ok else 0
            items.append({'label': chk['label'], 'ok': ok, 'points': pts, 'msgs': msgs})
        score += got
        tasks.append({'no': t['no'], 'section': t['section'], 'title': t['title'], 'points': task_points(t),
                      'got': got, 'items': items})
    total = total_points(exam)
    sections = []
    for s in SECTIONS:
        ts = [x for x in tasks if x['section'] == s]
        if ts:
            sections.append({'name': s, 'points': sum(x['points'] for x in ts), 'got': sum(x['got'] for x in ts)})
    return {'score': score, 'total': total, 'passed': score >= exam.get('pass', 70) * total / 100,
            'sections': sections, 'tasks': tasks, 'has_vba': bool(ctx.vba['sources'])}


def validate(exam):
    """시험 정의 검사(형식·시트 이름·정답 수식 계산)."""
    errs = []
    names = {s['name'] for s in exam.get('sheets', [])}
    for key in ('id', 'level', 'title', 'minutes', 'sheets', 'tasks'):
        if key not in exam:
            errs.append(f'{key} 없음')
    if exam.get('level') not in LEVELS:
        errs.append('level 은 c2 또는 c1')
    for t in exam.get('tasks', []):
        tag = f"{t.get('no')}"
        if t.get('section') not in SECTIONS:
            errs.append(f'{tag}: section')
        if t.get('sheet') not in names:
            errs.append(f"{tag}: 시트 '{t.get('sheet')}' 없음")
        for c in t.get('checks', []):
            if c.get('kind') not in CHECKS:
                errs.append(f"{tag}: 알 수 없는 kind {c.get('kind')}")
            if not c.get('label') or not isinstance(c.get('points'), (int, float)):
                errs.append(f'{tag}: label·points 필요')
            if c.get('kind') == 'formula':
                try:
                    ctx = type('C', (), {'orig': original_book(exam)})()
                    for (r, cc), v in _expected_formula(ctx, c.get('sheet', t['sheet']), c):
                        if isinstance(v, fx.XLErr) and not c.get('allow_error'):
                            errs.append(f'{tag}: 정답이 {fx.addr(r, cc)} 에서 {v.code}')
                except Exception as e:  # noqa: BLE001
                    errs.append(f'{tag}: 정답 수식 오류 {e}')
    if exam.get('level') == 'c2' and total_points(exam) != 100:
        errs.append(f'배점 합계 {total_points(exam)} (100 이어야 함)')
    if exam.get('level') == 'c1' and total_points(exam) != 100:
        errs.append(f'배점 합계 {total_points(exam)} (100 이어야 함)')
    return errs



GRADE_SECONDS = 25


def grade(exam, data, filename='답안.xlsx'):
    """채점(계산 시간 제한 {GRADE_SECONDS}초 — 넘으면 BadFile 로 안내)."""
    try:
        with fx.time_limit(GRADE_SECONDS):
            return _grade_unlimited(exam, data, filename)
    except fx.TimeUp:
        raise xlsx.BadFile('파일 속 수식 계산이 너무 오래 걸려 채점을 멈췄습니다 — 아주 큰 범위·배열 수식을 줄여 다시 올려 주세요.')
