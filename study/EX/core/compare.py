"""두 파일 비교 채점: '문제(실습)' 파일과 '정답(완성)' 파일의 차이 = 해야 할 일.
그 차이를 시트별 항목으로 나누고, 수험자 파일이 정답과 같아졌는지 본다(공식 예제·교재 실습 파일 공용).
"""
import datetime as dt
import re

from openpyxl.worksheet.worksheet import Worksheet

from . import formula as fx
from . import pivots, vba, xlsx
from .exam import _cell_mask, _color, _fmt_sample, _norm_text, describe_charts

# 실제 시험 시트 이름일 때의 배점(공식 문제지 기준). 그 밖의 파일은 시트마다 같은 배점으로 100점.
SHEET_POINTS = {
    'c2': {'기본작업-1': 5, '기본작업-2': 10, '기본작업-3': 5, '계산작업': 40, '분석작업-1': 10, '분석작업-2': 10,
           '매크로작업': 10, '차트작업': 10, '기타작업-1': 10, '기타작업-2': 10},
    'c1': {'기본작업-1': 10, '기본작업-2': 5, '계산작업': 30, '분석작업-1': 10, '분석작업-2': 10, '기타작업-1': 10,
           '기타작업-2': 10, '기타작업-3': 15},
}


def section_of(sheet):
    for s in ('기본작업', '계산작업', '분석작업'):
        if sheet.startswith(s):
            return s
    return '기타작업'


def _key(name):
    return name.replace('-', '').replace(' ', '')


def exam_level(sheetnames):
    """시트 이름이 실제 시험 구성과 같으면 'c2' / 'c1', 아니면 None."""
    names = {_key(n) for n in sheetnames}
    if not {'계산작업', '분석작업1'} <= names:
        return None
    if '기본작업3' in names:            # 2급: 기본작업 시트 3개
        return 'c2'
    if '기타작업3' in names or '기타작업1' in names:
        return 'c1'
    return None


def _groups(cells):
    """좌표 집합 → 상하좌우로 이어진 덩어리들(각각 정렬된 좌표 목록)."""
    left = set(cells)
    out = []
    while left:
        start = left.pop()
        stack, comp = [start], [start]
        while stack:
            r, c = stack.pop()
            for nb in ((r + 1, c), (r - 1, c), (r, c + 1), (r, c - 1)):
                if nb in left:
                    left.remove(nb)
                    stack.append(nb)
                    comp.append(nb)
        out.append(sorted(comp))
    out.sort()
    return out


def _span(comp):
    rs = [r for r, _ in comp]
    cs = [c for _, c in comp]
    a, b = fx.addr(min(rs), min(cs)), fx.addr(max(rs), max(cs))
    return a if a == b else f'{a}:{b}'


def _raw(ws, r, c):
    return ws.cell(r, c).value if ws is not None else None


SHORT_DATE = {'mm-dd-yy': 'yyyy-mm-dd'}     # 기본 형식 14('간단한 날짜')는 한국어 엑셀에서 yyyy-mm-dd 로 보인다


class Fmt:
    """표시 형식 비교: 양수 예시가 같으면서 0·음수 표시도 같아야 같은 형식.
    회계·쉼표 스타일(_-, _(, '* ')은 0·음수 표시 방식이 여러 가지라 양수 예시만 본다."""

    def __init__(self, fmt):
        fmt = SHORT_DATE.get(fmt, fmt)
        self.fmt = fmt
        s = [x.strip() for x in _fmt_sample(fmt)]
        self.pos = (s[0], s[3], s[4])
        self.all = tuple(s)
        self.acct = any(t in fmt for t in ('* ', '_(', '_-'))

    def __eq__(self, other):
        if not isinstance(other, Fmt):
            return NotImplemented
        return self.pos == other.pos and (self.acct or other.acct or self.all == other.all)

    def __hash__(self):
        return hash(self.pos)


def _font_color(c):
    v = _color(c)
    return None if v in ('T1', '000000') else v       # 자동·검정·'텍스트 1'은 같은 검정


def _style_key(cell):
    f, a, b = cell.font, cell.alignment, cell.border
    fill = _color(cell.fill.fgColor) if cell.fill is not None and cell.fill.patternType == 'solid' else None
    return {'글꼴': f.name, '크기': float(f.sz or 11), '굵게': bool(f.b), '기울임꼴': bool(f.i), '밑줄': f.u or None,
            '글꼴 색': _font_color(f.color), '채우기': fill, '가로 맞춤': a.horizontal or 'general',
            '세로 맞춤': a.vertical or 'bottom', '줄 바꿈': bool(a.wrap_text),
            '표시 형식': Fmt(cell.number_format),
            '테두리': tuple(bool(getattr(b, s).style) for s in ('left', 'right', 'top', 'bottom'))}


UNDERLINE_KO = {'single': '실선', 'double': '이중 실선', 'singleAccounting': '회계용 실선',
                'doubleAccounting': '회계용 이중 실선'}


def _show_style(k, v):
    """채점 항목·힌트에 보일 말(내부 값 → 엑셀 화면의 이름)."""
    from .describe import color_name                 # describe 가 이 모듈을 불러오므로 여기서
    from .exam import HALIGN_KO, VALIGN_KO
    if k == '표시 형식':
        if not isinstance(v, Fmt):
            return ''
        if v.fmt == '@':
            return "'텍스트'"
        plain = re.sub(r'"[^"]*"|\[[^\]]*\]', '', v.fmt.lower())
        sample = v.all[3] if '%' in plain else v.all[4] if re.search(r'[ymd]', plain) else v.all[0]
        return f'{v.fmt} (예: {sample})'
    if k in ('글꼴 색', '채우기'):
        return '없음' if v is None and k == '채우기' else color_name(v)
    if k == '가로 맞춤':
        return HALIGN_KO.get(v, v)
    if k == '세로 맞춤':
        return VALIGN_KO.get(v, v)
    if k == '크기':
        return f'{v:g}'
    if k == '테두리' and isinstance(v, tuple):
        names = [n for n, on in zip(('왼쪽', '오른쪽', '위', '아래'), v) if on]
        return '모든 테두리' if len(names) == 4 else ('없음' if not names else '·'.join(names))
    if k == '밑줄':
        return UNDERLINE_KO.get(v, '없음') if v else '없음'
    if k == '테두리':
        return '있음' if any(v) else '없음'
    if isinstance(v, bool):
        return '예' if v else '아니요'
    return '' if v is None else str(v)


def _same_val(got, exp):
    if isinstance(exp, str) or isinstance(got, str):
        return _norm_text(got) == _norm_text(exp)
    if exp is None:
        return got in (None, '')
    if got is None:
        return False
    if isinstance(exp, dt.datetime) or isinstance(got, dt.datetime) or isinstance(exp, dt.date):
        return fx.same_value(xlsx.plain(got), xlsx.plain(exp), 1e-6)
    return fx.same_value(got, exp, 1e-6)


class Wb:
    def __init__(self, data):
        self.data = data
        self.f = xlsx.load(data, data_only=False)
        self.v = xlsx.load(data, data_only=True)
        self.book = xlsx.to_book(self.f, self.v)
        self._vba = None

    def ws(self, name):
        ws = self.f[name] if name in self.f.sheetnames else None
        return ws if isinstance(ws, Worksheet) else None        # 차트 시트는 None

    def wv(self, name):
        ws = self.v[name] if name in self.v.sheetnames else None
        return ws if isinstance(ws, Worksheet) else None

    @property
    def sheets(self):
        """셀이 있는 시트 이름들(차트 시트 제외)."""
        return [w.title for w in self.f.worksheets]

    def value(self, sheet, r, c):
        try:
            return self.book.sheet(sheet).get(r, c)
        except fx.XLErr:
            return None

    @property
    def custom_heights(self):
        if not hasattr(self, '_custom_heights'):
            self._custom_heights = xlsx.custom_height_rows(self.data)
        return self._custom_heights

    @property
    def raw_charts(self):
        if not hasattr(self, '_raw_charts'):
            self._raw_charts = xlsx.raw_charts(self.data)
        return self._raw_charts

    @property
    def procs(self):
        if self._vba is None:
            self._vba = vba.procedures(vba.sources(self.data))
        return self._vba


VOLATILE = ('TODAY', 'NOW', 'RAND', 'RANDBETWEEN')


def _volatile_funcs(text):
    return {n for n in VOLATILE if re.search(r'\b' + n + r'\s*\(', text or '', re.I)}


def _cells_items(sheet, src, ans, user):
    """값·수식이 바뀐 칸 → 덩어리마다 한 항목."""
    ws_s, ws_a = src.ws(sheet), ans.ws(sheet)
    ws_u = user.ws(sheet)
    changed = set()
    for row in ws_a.iter_rows():
        for cell in row:
            if cell.value != _raw(ws_s, cell.row, cell.column):
                changed.add((cell.row, cell.column))
    if ws_s is not None:
        for row in ws_s.iter_rows():
            for cell in row:
                if cell.value is not None and ws_a.cell(cell.row, cell.column).value is None:
                    changed.add((cell.row, cell.column))
    items = []
    wv_a = ans.wv(sheet)
    for comp in _groups(changed):
        bad = []
        formulas = sum(1 for r, c in comp if xlsx.formula_text(ws_a.cell(r, c).value) is not None
                       or type(ws_a.cell(r, c).value).__name__ == 'DataTableFormula')
        for r, c in comp:
            a_raw = ws_a.cell(r, c).value
            exp = wv_a.cell(r, c).value
            a_formula = xlsx.formula_text(a_raw) is not None
            dt_table = type(a_raw).__name__ == 'DataTableFormula'
            if ws_u is None:
                bad.append('시트 없음')
                break
            u_raw = ws_u.cell(r, c).value
            got = user.value(ws_u.title, r, c)
            if got is None and u_raw is not None and xlsx.formula_text(u_raw) is None \
                    and type(u_raw).__name__ != 'DataTableFormula':
                got = xlsx.plain(u_raw)
            volatile = _volatile_funcs(xlsx.formula_text(a_raw)) if a_formula else set()
            if volatile:
                uf = xlsx.formula_text(u_raw)
                if uf is None:
                    bad.append(f'{fx.addr(r, c)}: 수식이 아니라 값입니다')
                elif not volatile <= _volatile_funcs(uf):
                    bad.append(f"{fx.addr(r, c)}: {', '.join(sorted(volatile))} 함수를 써야 합니다")
            elif dt_table and type(u_raw).__name__ != 'DataTableFormula':
                bad.append(f'{fx.addr(r, c)}: 데이터 표가 아닙니다')
            elif a_formula and xlsx.formula_text(u_raw) is None and got not in (None, ''):
                bad.append(f'{fx.addr(r, c)}: 수식이 아니라 값입니다')
            elif not _same_val(got, xlsx.plain(exp) if exp is not None else None):
                bad.append(f'{fx.addr(r, c)}: {fx.display(got) if got not in (None, "") else "(빈칸)"} → '
                           f'{fx.display(xlsx.plain(exp)) if exp not in (None, "") else "(빈칸)"}')
            if len(bad) >= 3:
                break
        label = ('수식 결과 ' if formulas * 2 >= len(comp) else '셀 내용 ') + _span(comp)
        r0, c0 = comp[0]
        a0 = ws_a.cell(r0, c0).value
        hint = (f'{fx.addr(r0, c0)} 정답 수식: {xlsx.formula_text(a0)}' if xlsx.formula_text(a0) else
                f"{fx.addr(r0, c0)} 정답 값: {fx.display(xlsx.plain(wv_a.cell(r0, c0).value)) if wv_a.cell(r0, c0).value is not None else '(빈칸)'}")
        items.append({'label': label, 'ok': not bad, 'msgs': bad, 'hint': hint})
    return items


def _style_items(sheet, src, ans, user):
    ws_s, ws_a, ws_u = src.ws(sheet), ans.ws(sheet), user.ws(sheet)
    if ws_s is None:
        return []
    changes = {}
    for row in ws_a.iter_rows():
        for cell in row:
            ka = _style_key(cell)
            ks = _style_key(ws_s.cell(cell.row, cell.column))
            diff = tuple(sorted(k for k in ka if ka[k] != ks[k]))
            if diff:
                changes[(cell.row, cell.column)] = diff
    items = []
    by_attr = {}
    for pos, diff in changes.items():
        ka = _style_key(ws_a.cell(*pos))
        for k in diff:
            by_attr.setdefault((k, ka[k]), set()).add(pos)
    for attr, val in sorted(by_attr, key=lambda k: (min(by_attr[k]), k[0])):
        for comp in _groups(by_attr[(attr, val)]):
            bad, first = [], None
            for r, c in comp:
                if ws_u is None:
                    bad.append('시트 없음')
                    break
                if _style_key(ws_u.cell(r, c))[attr] != _style_key(ws_a.cell(r, c))[attr]:
                    bad.append(fx.addr(r, c))
                    if first is None:
                        first = (fx.addr(r, c), _show_style(attr, _style_key(ws_u.cell(r, c))[attr]),
                                 _show_style(attr, _style_key(ws_a.cell(r, c))[attr]))
            want = _show_style(attr, _style_key(ws_a.cell(*comp[0]))[attr])
            msg = f"{', '.join(bad[:4])}{' 등' if len(bad) > 4 else ''} 이(가) 다릅니다"
            if first:
                msg += f" — {first[0]}: 현재 {first[1] or '(기본)'} → 정답 {first[2] or '(기본)'}"
            items.append({'label': f'{attr} {_span(comp)}', 'ok': not bad, 'hint': f'{attr}: {want}'.rstrip(': '),
                          'msgs': [msg] if bad else []})
    # 병합
    m_src = {str(m) for m in ws_s.merged_cells.ranges}
    m_usr = {str(m) for m in ws_u.merged_cells.ranges} if ws_u is not None else set()
    for m in sorted({str(m) for m in ws_a.merged_cells.ranges} - m_src):
        items.append({'label': f'셀 병합 {m}', 'ok': m in m_usr, 'msgs': [] if m in m_usr else [f'{m} 병합 필요'],
                      'hint': f'{m} 을 병합'})
    # 행 높이(엑셀이 글꼴·테두리에 맞춰 스스로 정한 높이는 빼고, 직접 지정한 행만)
    own = ans.custom_heights.get(sheet, set())
    for r, d in ws_a.row_dimensions.items():
        hs = ws_s.row_dimensions[r].height if r in ws_s.row_dimensions else None
        if d.height and r in own and (hs is None or abs(d.height - hs) > 0.5):
            hu = ws_u.row_dimensions[r].height if ws_u is not None and r in ws_u.row_dimensions else None
            ok = hu is not None and abs(hu - d.height) <= 0.5
            items.append({'label': f'{r}행 높이 {d.height:g}', 'ok': ok, 'msgs': [] if ok else [f'현재 {hu or "기본"}'],
                          'hint': f'{r}행 높이를 {d.height:g} 으로'})
    return items


def _rules(ws):
    out = []
    for cf in ws.conditional_formatting:
        for rule in cf.rules:
            out.append((str(cf.sqref).replace('$', '').upper(), rule))
    return out


def _rule_key(sq, rule):
    return sq, rule.type, tuple(_norm_f(f) for f in (rule.formula or ()))


def _norm_f(f):
    return (f or '').replace('$', '').replace(' ', '').upper()


def broken_rule(rule):
    """정답 파일에 남은 깨진 규칙: 행 번호가 엑셀 끝(1048576 근처)이거나 수식에 둥근 따옴표."""
    text = ' '.join(rule.formula or [])
    return bool(re.search(r'\$?[A-Z]+\$?10[0-9]{5}', text) or re.search('[“”‘’]', text))


RULE_ATTRS = (('operator', None), ('rank', None), ('percent', False), ('bottom', False), ('aboveAverage', True),
              ('equalAverage', False), ('stdDev', None), ('text', None), ('timePeriod', None))


def _dxf(d):
    font = (_font_color(d.font.color) if d and d.font else None, bool(d and d.font and d.font.b),
            bool(d and d.font and d.font.i))
    fill = (_color(d.fill.bgColor) or _color(d.fill.fgColor)) if d and d.fill else None
    return font, fill


def _scale(rule):
    """색조·데이터 막대·아이콘 집합의 설정."""
    if rule.colorScale is not None:
        return ('colorScale', tuple(c.type for c in rule.colorScale.cfvo),
                tuple(_color(c) for c in rule.colorScale.color))
    if rule.dataBar is not None:
        return ('dataBar', tuple(c.type for c in rule.dataBar.cfvo), _color(rule.dataBar.color))
    if rule.iconSet is not None:
        return ('iconSet', rule.iconSet.iconSet, bool(rule.iconSet.reverse), tuple(c.type for c in rule.iconSet.cfvo))
    return None


def _formula_funcs(text):
    try:
        return set(fx.functions_used(fx.parse('=' + text)))
    except fx.FormulaError:
        return set()


def _rule_same(rule, r, ans, user, sheet, sq):
    """정답 규칙과 수험자 규칙이 같은 일을 하는지 → (같음, 다르면 이유)."""
    if r.type != rule.type:
        return False, '규칙 종류가 다릅니다'
    if rule.type == 'expression' and rule.formula:
        uf = (r.formula or ['FALSE'])[0]
        if uf.replace(' ', '').upper() != rule.formula[0].replace(' ', '').upper():
            try:
                want = _cell_mask(_Ctx(ans), sheet, sq.split()[0], '=' + rule.formula[0], ans.book)
                got = _cell_mask(_Ctx(user), sheet, sq.split()[0], '=' + uf, user.book)
            except fx.FormulaError:
                return False, f'규칙 수식 ={uf} 을(를) 계산할 수 없습니다'
            if want != got:
                return False, f'규칙 수식 ={uf} 의 결과가 다릅니다'
        need = _formula_funcs(rule.formula[0])         # 지문에 "AND 함수 사용" 처럼 적힌 함수
        miss = need - _formula_funcs(uf)
        if miss:
            return False, f"{', '.join(sorted(miss))} 함수를 사용해 규칙을 만드세요(현재 ={uf})"
    else:
        for attr, default in RULE_ATTRS:
            a = getattr(rule, attr, None)
            u = getattr(r, attr, None)
            if (default if a is None else a) != (default if u is None else u):
                return False, f'규칙 조건({attr})이 다릅니다: {u} → {a}'
        if rule.text is None and [_norm_f(f) for f in rule.formula or []] != [_norm_f(f) for f in r.formula or []]:
            return False, f"조건 값 {', '.join(r.formula or []) or '-'} → {', '.join(rule.formula or [])}"
        if _scale(rule) != _scale(r):
            return False, '색조·데이터 막대·아이콘 설정이 다릅니다'
    if _scale(rule) is None and _dxf(rule.dxf) != _dxf(r.dxf):
        return False, '서식(글꼴 색·스타일·채우기)이 다릅니다'
    return True, ''


RULE_KO = {'cellIs': '셀 값 비교', 'expression': '수식', 'top10': '상위·하위 항목', 'aboveAverage': '평균 초과·미만',
           'containsText': '특정 텍스트 포함', 'colorScale': '색조', 'dataBar': '데이터 막대', 'iconSet': '아이콘 집합',
           'duplicateValues': '중복 값', 'uniqueValues': '고유 값', 'timePeriod': '날짜 발생'}


OP_KO = {'greaterThan': '보다 큼', 'lessThan': '보다 작음', 'between': '사이', 'notBetween': '사이 아님', 'equal': '같음',
         'notEqual': '같지 않음', 'greaterThanOrEqual': '크거나 같음', 'lessThanOrEqual': '작거나 같음'}


def _rule_hint(rule):
    t = rule.type
    if t == 'expression' and rule.formula:
        return f'규칙: ={rule.formula[0]}'
    if t == 'cellIs':
        return f"셀 값 {', '.join(rule.formula or [])} {OP_KO.get(rule.operator, rule.operator or '')}"
    if t == 'top10':
        return f"{'하위' if rule.bottom else '상위'} {rule.rank}{'%' if rule.percent else '개'} 항목"
    if t == 'aboveAverage':
        return '평균 ' + ('초과' if rule.aboveAverage is not False else '미만') + ('(같은 값 포함)' if rule.equalAverage else '')
    if t == 'containsText':
        return f"'{rule.text}' 포함"
    return f"규칙 종류: {RULE_KO.get(t, t)}"


def _cf_items(sheet, src, ans, user):
    ws_s, ws_a, ws_u = src.ws(sheet), ans.ws(sheet), user.ws(sheet)
    old = {_rule_key(*x) for x in _rules(ws_s)} if ws_s is not None else set()
    items, seen = [], set()
    for sq, rule in _rules(ws_a):
        k = _rule_key(sq, rule)
        if k in old or k in seen or broken_rule(rule):
            continue
        seen.add(k)
        label = f'조건부 서식 {sq}'
        hint = _rule_hint(rule)
        cands = [r for s, r in (_rules(ws_u) if ws_u is not None else []) if s == sq]
        if not cands:
            items.append({'label': label, 'ok': False, 'msgs': [f'{sq} 범위의 규칙이 없습니다'], 'hint': hint})
            continue
        ok, msg = False, '규칙 내용이 다릅니다'
        for r in cands:
            ok, why = _rule_same(rule, r, ans, user, sheet, sq)
            if ok:
                break
            msg = why
        items.append({'label': label, 'ok': ok, 'msgs': [] if ok else [msg], 'hint': hint})
    return items


class _Ctx:
    def __init__(self, wb):
        self.orig = wb.book


def _dv_items(sheet, src, ans, user):
    ws_s, ws_a, ws_u = src.ws(sheet), ans.ws(sheet), user.ws(sheet)
    key = lambda d: (str(d.sqref).upper(), d.type, (d.formula1 or '').replace('$', ''),  # noqa: E731
                     (d.formula2 or '').replace('$', ''), d.operator or 'between')
    old = {key(d) for d in ws_s.data_validations.dataValidation} if ws_s is not None else set()
    have = {key(d): d for d in ws_u.data_validations.dataValidation} if ws_u is not None else {}
    items = []
    for d in ws_a.data_validations.dataValidation:
        k = key(d)
        if k in old:
            continue
        u = have.get(k)
        bad = []
        if u is None:
            bad.append('유효성 검사 범위·조건이 다릅니다')
        else:
            for attr, name in (('errorTitle', '오류 제목'), ('error', '오류 메시지'), ('promptTitle', '설명 제목'),
                               ('prompt', '설명 메시지'), ('errorStyle', '스타일')):
                if _norm_text(getattr(d, attr)) != _norm_text(getattr(u, attr)):
                    bad.append(name)
        items.append({'label': f'데이터 유효성 {k[0]}', 'ok': not bad, 'msgs': bad,
                      'hint': f"제한 대상 {d.type}, 조건 {d.operator or 'between'} {d.formula1 or ''} {d.formula2 or ''}".strip()})
    return items


CHART_KIND_KO = {'col': '세로 막대형', 'bar': '가로 막대형', 'line': '꺾은선형', 'pie': '원형', 'area': '영역형',
                 'scatter': '분산형', 'doughnut': '도넛형', 'radar': '방사형'}
LEGEND_KO = {'t': '위쪽', 'b': '아래쪽', 'l': '왼쪽', 'r': '오른쪽', 'tr': '오른쪽 위'}
GROUP_KO = {'clustered': '묶은형', 'stacked': '누적형', 'percentStacked': '100% 누적형', 'standard': '표준'}


def _chart_show(label, v):
    """차트 항목 값을 학습자가 읽을 말로(내부 코드·파이썬 목록 모양 없이)."""
    if v in (None, '', [], ()):
        return '없음'
    if label == '범례':
        return LEGEND_KO.get(v, str(v))
    if label == '누적 여부':
        return GROUP_KO.get(v, str(v))
    if isinstance(v, (list, tuple, set)):
        parts = []
        for x in v:
            if isinstance(x, tuple) and len(x) == 2:
                parts.append(f"{x[0]}: {CHART_KIND_KO.get(x[1], x[1])}")
            else:
                parts.append(str(x))
        return ', '.join(parts) or '없음'
    return str(v)


def _chart_items(sheet, src, ans, user):
    try:
        a = _charts_of(ans, sheet)
    except KeyError:
        return []
    if not a:
        return []
    s = _charts_of(src, sheet) if src.ws(sheet) is not None else []
    u = _charts_of(user, sheet) if user.ws(sheet) is not None else []
    items = []
    for i, ca in enumerate(a):
        cs = s[i] if i < len(s) else None
        cu = u[i] if i < len(u) else None

        def add(label, getter):
            want = getter(ca)
            if cs is not None and getter(cs) == want:
                return
            got = getter(cu) if cu is not None else None
            if cs is None and want in ('', [], None, ()) :
                return                                   # 새 차트에서 정답에도 없는 항목은 과제가 아님
            items.append({'label': f'차트: {label}', 'ok': got == want, 'hint': f'정답: {_chart_show(label, want)}',
                          'msgs': [] if got == want else [f'{_chart_show(label, got)} → {_chart_show(label, want)}'
                                                          if cu is not None else '차트가 없습니다']})
        add('데이터 계열', lambda c: [x['name'] for x in c['series']])
        add('차트 종류', lambda c: sorted({(x['name'], x['kind']) for x in c['series']}))
        add('보조 축', lambda c: sorted(x['name'] for x in c['series'] if x['secondary']))
        add('데이터 레이블', lambda c: sorted(x['name'] for x in c['series'] if x['labels']))
        add('추세선', lambda c: sorted(x['name'] for x in c['series'] if x['trend']))
        add('차트 제목', lambda c: _norm_text(c['title']))
        add('세로 축 제목', lambda c: _norm_text(c['y_title']))
        add('가로 축 제목', lambda c: _norm_text(c['x_title']))
        add('범례', lambda c: c['legend'])
        add('누적 여부', lambda c: c['grouping'])
        hid_a = {x['name'] for x in ca['series'] if not x.get('visible', True)}
        hid_u = {x['name'] for x in cu['series'] if not x.get('visible', True)} if cu is not None else set()
        if hid_u - hid_a:
            items.append({'label': '차트: 보이는 계열', 'ok': False, 'hint': '정답: 모든 계열이 보임',
                          'msgs': [f"{', '.join(sorted(hid_u - hid_a))} 계열이 채우기·선 없음으로 보이지 않습니다"]})
    return items


def _charts_of(wb, sheet):
    """openpyxl 이 읽은 차트가 파일 속 차트보다 적으면(그룹 안 차트 등) XML 에서 직접 읽은 것으로."""
    ws = wb.ws(sheet)
    if ws is None:
        raise KeyError(sheet)
    raw = wb.raw_charts.get(sheet, [])
    charts = raw if len(raw) > len(ws._charts) else None
    return describe_charts(_Ctx2(wb), sheet, charts)


class _Ctx2:
    def __init__(self, wb):
        self.wb = wb
        self.book = wb.book

    def ws(self, name):
        ws = self.wb.ws(name)
        if ws is None:
            raise KeyError(name)
        return ws


def _pivot_items(sheet, src, ans, user):
    ws_a = ans.ws(sheet)
    if not getattr(ws_a, '_pivots', None):
        return []
    want = [p for p in pivots.describe(ans.f) if p['sheet'] == sheet]
    got = pivots.describe(user.f)
    items = []
    for p in want:
        r1, c1, _, _ = fx.parse_range(p['ref'])
        spec = {'rows': p['rows'], 'cols': p['cols'], 'filters': p['filters'], 'values': p['values'],
                'sheet': sheet, 'layout': p['layout']}
        best = None
        for g in got:
            ok, why = pivots.match(spec, g)
            if ok and fx.parse_range(g['ref'])[:2] == (r1, c1):
                best = (True, [])
                break
            if best is None or len(why) < len(best[1]):
                best = (False, why or [f"위치 {g['ref']} → {p['ref']}"])
        best = best or (False, ['피벗 테이블이 없습니다'])
        hint = (f"행: {', '.join(p['rows']) or '-'} · 열: {', '.join(p['cols']) or '-'} · 필터: {', '.join(p['filters']) or '-'} · 값: "
                + ', '.join(f"{f} {pivots.FUNC_NAMES.get(fn, fn)}" for f, fn in p['values']) + f" · 위치 {sheet}!{p['ref'].split(':')[0]}")
        items.append({'label': f"피벗 테이블 {p['ref']}", 'ok': best[0], 'msgs': best[1][:3], 'hint': hint})
    return items


def _misc_items(sheet, src, ans, user):
    ws_s, ws_a, ws_u = src.ws(sheet), ans.ws(sheet), user.ws(sheet)
    items = []
    if ws_s is None:
        return items

    def page(ws):
        if ws is None:
            return {}
        return {'용지 방향': 'landscape' if ws.page_setup.orientation == 'landscape' else 'portrait',
                '가로 가운데': bool(ws.print_options.horizontalCentered),
                '세로 가운데': bool(ws.print_options.verticalCentered),
                '인쇄 영역': (ws.print_area or '').split('!')[-1].replace('$', ''),
                '반복할 행': (ws.print_title_rows or '').replace('$', ''),
                '머리글': '|'.join(_norm_text(getattr(ws.oddHeader, p).text) for p in ('left', 'center', 'right')),
                '바닥글': '|'.join(_norm_text(getattr(ws.oddFooter, p).text) for p in ('left', 'center', 'right')),
                '시트 보호': bool(ws.protection.sheet),
                '보기': ws.sheet_view.view or 'normal'}
    ps, pa, pu = page(ws_s), page(ws_a), page(ws_u)
    for k in pa:
        if pa[k] != ps.get(k):
            items.append({'label': k, 'ok': pu.get(k) == pa[k], 'hint': f'정답: {pa[k]}',
                          'msgs': [] if pu.get(k) == pa[k] else [f"{pu.get(k)} → {pa[k]}"]})
    # 자동 필터·현재 위치 고급 필터: 숨겨진 행
    hid = lambda ws: {r for r, d in ws.row_dimensions.items() if d.hidden} if ws is not None else set()  # noqa: E731
    if hid(ws_a) != hid(ws_s) or (ws_a.auto_filter.ref and ws_a.auto_filter.ref != ws_s.auto_filter.ref):
        want, got = hid(ws_a), hid(ws_u)
        ok = want == got and (not ws_a.auto_filter.ref or bool(ws_u is not None and ws_u.auto_filter.ref))
        msg = []
        if not ok:
            msg = [f'숨겨진 행 {len(got)}개 → {len(want)}개' if want != got else '자동 필터가 설정되어 있지 않습니다']
        items.append({'label': '필터 결과(보이는 행)', 'ok': ok, 'msgs': msg,
                      'hint': f"자동 필터 범위 {ws_a.auto_filter.ref or '-'}, 숨겨진 행 {len(want)}개"})
    # 잠금·숨김
    if pa['시트 보호'] and not ps['시트 보호']:
        bad, changed = [], 0
        for row in ws_a.iter_rows():
            for cell in row:
                p = cell.protection
                o = ws_s.cell(cell.row, cell.column).protection
                if (p.locked, p.hidden) == (o.locked, o.hidden):
                    continue                     # 정답에서 바꾼 칸만 본다
                changed += 1
                q = ws_u.cell(cell.row, cell.column).protection if ws_u is not None else None
                if q is None or (p.locked, p.hidden) != (q.locked, q.hidden):
                    bad.append(cell.coordinate)
        if changed:
            items.append({'label': '셀 잠금·숨김', 'ok': not bad, 'msgs': [f"{', '.join(bad[:4])} 설정 다름"] if bad else []})
    # 메모
    for row in ws_a.iter_rows():
        for cell in row:
            if cell.comment and (ws_s.cell(cell.row, cell.column).comment is None):
                cu = ws_u.cell(cell.row, cell.column).comment if ws_u is not None else None
                want = _norm_text(cell.comment.text.split(':\n', 1)[-1])
                ok = cu is not None and want in _norm_text(cu.text)
                items.append({'label': f'메모 {cell.coordinate}', 'ok': ok, 'hint': f"메모 내용: {want}",
                              'msgs': [] if ok else [f"'{want}' 메모 필요"]})
    # 시나리오
    scen = lambda ws: {s.name: s for s in (ws.scenarios.scenario if ws is not None and ws.scenarios else [])}  # noqa: E731
    sc_a, sc_s, sc_u = scen(ws_a), scen(ws_s), scen(ws_u)
    new = sorted(set(sc_a) - set(sc_s))
    if new:
        msgs = []
        for name in new:
            if name not in sc_u:
                msgs.append(f"'{name}' 시나리오가 없습니다")
                continue
            want = {ic.r.replace('$', '').upper(): ic.val for ic in sc_a[name].inputCells}
            got = {ic.r.replace('$', '').upper(): ic.val for ic in sc_u[name].inputCells}
            if set(want) != set(got):
                msgs.append(f"'{name}' 변경 셀 {', '.join(sorted(got))} → {', '.join(sorted(want))}")
                continue
            for cell, v in want.items():
                if not _same_val(_num(got[cell]), _num(v)):
                    msgs.append(f"'{name}' {cell} 값 {got[cell]} → {v}")
                    break
        items.append({'label': '시나리오 ' + ', '.join(new), 'ok': not msgs, 'msgs': msgs[:3],
                      'hint': '; '.join(f"{n}: " + ', '.join(f'{ic.r}={ic.val}' for ic in sc_a[n].inputCells)
                                        for n in new)})
    return items


def _num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return v


def _book_items(src, ans, user):
    """통합 문서 전체: 이름 정의, 새 시트(시나리오 요약 등), 매크로·VBA."""
    items = []
    names_a = dict(xlsx._defined_names(ans.f))
    names_s = dict(xlsx._defined_names(src.f))
    names_u = {k.lower(): v for k, v in xlsx._defined_names(user.f)}
    for n, ref in names_a.items():
        if names_s.get(n) == ref:
            continue
        got = names_u.get(n.lower())
        ok = got is not None and got.replace('$', '').replace("'", '').upper() == ref.replace('$', '').replace("'", '').upper()
        where = ref.lstrip('=').split('!')[0].strip("'") if '!' in ref else None
        items.append(('기본작업', {'label': f'이름 정의 {n}', 'ok': ok, 'hint': f'{n} = {ref}', 'sheet': where,
                                   'msgs': [] if ok else [f'{got or "없음"} → {ref}']}))
    for s in ans.sheets:
        if s not in src.f.sheetnames:
            ok = s in user.f.sheetnames
            if ok:
                idx_a, idx_u = ans.f.sheetnames.index(s), user.f.sheetnames.index(s)
                nxt = ans.f.sheetnames[idx_a + 1] if idx_a + 1 < len(ans.f.sheetnames) else None
                if nxt and nxt in user.f.sheetnames and user.f.sheetnames.index(nxt) != idx_u + 1:
                    ok = False
            items.append(('분석작업', {'label': f"'{s}' 시트(위치 포함)", 'ok': ok, 'msgs': [] if ok else ['시트가 없거나 위치가 다릅니다']}))
    # VBA: 정답에서 새로 생겼거나 내용이 바뀐 프로시저
    pa, ps, pu = ans.procs, src.procs, user.procs
    for name, body in pa.items():
        if name in ps and _similar(ps[name], body) >= 0.6:
            continue                             # 문제 파일에 이미 있던(거의 같은) 코드
        if len(_code_lines(body)) <= 1:
            continue
        sec = '계산작업' if body.lstrip().lower().startswith(('public function', 'function')) else '기타작업'
        u = pu.get(name)
        if u is None:
            items.append((sec, {'label': f'VBA {name}', 'ok': False, 'msgs': ['프로시저가 없습니다(.xlsm 으로 저장)']}))
            continue
        score = _similar(u, body)
        lit_u = _literals(u)
        miss = _same_result_formulas(_literals(body) - lit_u, lit_u, _literals(body), ans.book)
        ok = score >= 0.95 or (score >= 0.3 and not miss)
        msgs = []
        if not ok:
            msgs.append(f'정답 코드와 {round(score * 100)}% 일치(같은 문장 기준)')
            if miss:
                msgs.append('빠졌거나 다른 값: ' + ', '.join(sorted(miss)[:4]))
        items.append((sec, {'label': f'VBA {name}', 'ok': ok, 'code': '\n'.join(_code_lines(body)), 'msgs': msgs,
                            'macro': name}))
    b_a = set(vba.buttons(ans.data)) - set(vba.buttons(src.data))
    b_u = {(_norm_text(t), m.lower().split('.')[-1]) for t, m in vba.buttons(user.data)}
    for text, macro in sorted(b_a):
        ok = (_norm_text(text), macro.lower().split('.')[-1]) in b_u
        items.append(('기타작업', {'label': f"단추 '{text}' → {macro}", 'ok': ok,
                                   'msgs': [] if ok else ['단추가 없거나 매크로가 연결되지 않았습니다']}))
    return items


def _similar(got, want):
    """정답 코드의 문장 중 몇 %가 들어 있는지(공백·대소문자·주석 무시)."""
    w = set(_code_lines(want))
    return len(w & set(_code_lines(got))) / max(1, len(w))


RANGE_RE = re.compile(r'\$?([a-z]{1,3})\$?(\d+)(?::\$?([a-z]{1,3})\$?(\d+))?$')


COLOR_INDEX_RGB = {'6': '65535', '3': '255', '4': '65280', '5': '16711680', '2': '16777215', '1': '0'}


def _literals(body):
    """코드의 핵심 값: 쓰거나 서식을 바꾼 셀들(범위 글자를 모두 합친 칸 목록), 그 밖의 문자열(수식·메시지), 색·서식 줄의 값."""
    out, cells = set(), set()
    lines = _code_lines(body)
    # 기록한 매크로 끝의 Range("..").Select 는 커서를 놓은 것일 뿐 — 그 줄 뒤에 쓰기·서식이 없으면 셀 목록에서 뺀다
    tail = [lines.pop()] if lines and re.match(r'end (sub|function)', lines[-1]) else []
    while lines and re.fullmatch(r'(range\(".*"\)|cells\(.*\)|\[.*\])\.select', lines[-1].replace(' ', '')):
        lines = lines[:-1]
    lines += tail
    for line in lines:
        for s in re.findall(r'"((?:[^"]|"")*)"', line):
            s = s.replace('""', '"').replace(' ', '')
            m = RANGE_RE.match(s)
            if m:
                r1, c1 = int(m.group(2)), fx.col_num(m.group(1).upper())
                r2, c2 = (int(m.group(4)), fx.col_num(m.group(3).upper())) if m.group(3) else (r1, c1)
                if (r2 - r1 + 1) * (c2 - c1 + 1) <= 10000:
                    cells |= {fx.addr(r, c) for r in range(r1, r2 + 1) for c in range(c1, c2 + 1)}
            elif s:
                out.add(s)
        if re.search(r'\.(color|colorindex|themecolor|numberformat\w*|bold|italic|size|name)\s*=', line) \
                and not re.search(r'tintandshade|patterncolor', line):
            key, val = line.split('=', 1)[0].split('.')[-1], re.sub(r'\s+', '', line.split('=', 1)[1])
            if key.strip().lower() == 'colorindex':
                key, val = 'color', COLOR_INDEX_RGB.get(val, val)
            out.add(re.sub(r'\s+', '', key) + '=' + val)
    if cells:
        out.add('셀 ' + ','.join(sorted(cells, key=fx.parse_addr)))
    return out


R1C1_RE = re.compile(r'(?<![a-z0-9_.])r(\[-?\d+\]|\d+)?c(\[-?\d+\]|\d+)?(?![a-z0-9_(])', re.I)


def r1c1_to_a1(formula, row, col):
    """R1C1 수식(=rc[-2]+rc[-1]) → (row, col) 칸 기준 A1 수식. 따옴표 안 글자는 그대로."""
    def one(m):
        def part(g, base):
            if g is None:
                return base, False
            if g.startswith('['):
                return base + int(g[1:-1]), False
            return int(g), True
        r, r_abs = part(m.group(1), row)
        c, c_abs = part(m.group(2), col)
        if r < 1 or c < 1:
            raise fx.FormulaError('범위를 벗어난 R1C1 참조')
        return ('$' if c_abs else '') + fx.addr(1, c)[:-1] + ('$' if r_abs else '') + str(r)
    pieces = re.split(r'("[^"]*")', formula)
    return ''.join(x if x.startswith('"') else R1C1_RE.sub(one, x) for x in pieces)


def _at(formula, r, c, first):
    """범위에 넣은 수식이 (r, c) 칸에서 갖는 모양: R1C1 은 그 칸 기준, A1 은 첫 칸에서 채운 것처럼 이동."""
    if any(R1C1_RE.search(x) for x in re.split(r'"[^"]*"', formula)):
        return r1c1_to_a1(formula, r, c)
    return '=' + fx.unparse(fx.shift(fx.parse(formula), r - first[0], c - first[1]))


def _formula_cells(lits):
    for x in lits:
        if x.startswith('셀 '):
            return [fx.parse_addr(a) for a in x[2:].split(',')]
    return []


def _same_result_formulas(miss, user_lits, ans_lits, book):
    """빠진 수식 값 중, 수험자 코드의 다른 수식이 매크로가 다루는 칸들에서 같은 결과를 내면 맞은 것으로 뺀다."""
    cells = _formula_cells(ans_lits)
    if not cells or not any(x.startswith('=') for x in miss):
        return miss
    spare = [x for x in user_lits if x.startswith('=') and x not in ans_lits]
    left = set(miss)
    for want in [x for x in miss if x.startswith('=')]:
        hit = next((u for u in spare if _same_results(want, u, cells, book)), None)
        if hit is not None:
            left.discard(want)
            spare.remove(hit)
    return left


def _same_results(f1, f2, cells, book):
    for sheet in book.sheets:
        vals = []
        try:
            for r, c in cells[:200]:
                a = fx.evaluate_text(_at(f1, r, c, cells[0]), book, sheet, r, c)
                b = fx.evaluate_text(_at(f2, r, c, cells[0]), book, sheet, r, c)
                if isinstance(a, (fx.XLErr, fx.Arr)) or isinstance(b, (fx.XLErr, fx.Arr)):
                    raise fx.FormulaError('계산 불가')
                vals.append((a, b))
        except (fx.FormulaError, fx.XLErr, ValueError, ZeroDivisionError):
            continue
        if vals and any(a not in (None, '', 0) for a, _ in vals) and all(_same_val(b, a) for a, b in vals):
            return True
    return False


def _code_lines(body):
    lines = []
    for line in body.splitlines():
        line = re.sub(r"'.*$", '', line).strip()
        if not line or line.lower().startswith('attribute '):
            continue
        lines.append(re.sub(r'\s+', ' ', line).lower())
    return lines


def _grade_unlimited(src_bytes, ans_bytes, user_bytes, level=None):
    """문제 파일·정답 파일·수험자 파일 → 시트별 항목 채점. level 이 없으면 시트 이름으로 시험 구성인지 판단."""
    src = Wb(src_bytes)
    ans = Wb(ans_bytes)
    user = Wb(user_bytes)
    level = level or exam_level(ans.sheets)
    points = SHEET_POINTS[level] if level else None
    sheets, book_items = [], _book_items(src, ans, user)
    for sheet in ans.sheets:
        if src.ws(sheet) is None:
            continue
        section = section_of(sheet)
        fns = (_cells_items, _style_items, _cf_items, _dv_items, _pivot_items, _chart_items, _misc_items)
        if level and section == '계산작업':       # 시험의 계산작업은 수식만(정답 파일의 강조 서식 등은 채점 안 함)
            fns = (_cells_items,)
        items = []
        for fn in fns:
            items += fn(sheet, src, ans, user)
        pts = {_key(k): v for k, v in points.items()}.get(_key(sheet), 0) if points else 0
        sheets.append({'name': sheet, 'section': section, 'points': pts, 'items': items})
    if not sheets:
        raise xlsx.BadFile('문제 파일과 정답 파일에 같은 이름의 시트가 없습니다.')
    by_name = {s['name']: s for s in sheets}
    macro_sheet = next((s for s in sheets if _key(s['name']) in ('매크로작업', '기타작업2')), None)
    if level == 'c2':
        macro_sheet = by_name.get('매크로작업', macro_sheet)
    button_macros = {m.split('!')[-1].split('.')[-1].lower() for _, m in vba.buttons(ans.data)}
    # 통합 문서 단위 항목: 이름 정의는 이름이 가리키는 시트, 매크로(단추로 실행)는 매크로 시트, VBA 프로그래밍은 마지막 기타작업
    for sec, it in book_items:
        target = by_name.get(it.pop('sheet', None) or '')
        macro = it.pop('macro', None)
        if it['label'].startswith('단추'):
            target = macro_sheet
        elif macro and sec == '기타작업':
            if level == 'c2' or macro.lower() in button_macros:
                target = macro_sheet
            target = target or next((s for s in reversed(sheets) if s['section'] == '기타작업'), None)
        if target is None:
            target = next((s for s in sheets if s['section'] == sec), sheets[0])
        target['items'].append(it)
    active = [s for s in sheets if s['items']]
    if not active:
        raise xlsx.BadFile('문제 파일과 정답 파일이 같아서 채점할 항목이 없습니다.')
    mine = set(user.sheets)
    if not _same_origin(src, user):
        raise xlsx.BadFile('올린 파일이 이 문제의 실습 파일이 아닌 것 같습니다 — 문제 파일의 내용(글자)이 거의 남아 있지 않습니다. '
                           '이 문제에서 받은 실습 파일에 풀이를 넣어 올려 주세요.')
    if not any(s['name'] in mine for s in active):
        raise xlsx.BadFile(f"올린 파일에 이 문제의 시트({', '.join(s['name'] for s in active[:3])})가 없습니다 — "
                           "다른 문제의 파일이 아닌지 확인하세요.")
    if not points:      # 시험 구성이 아닌 실습 파일: 할 일이 있는 시트끼리 100점을 나눈다(반올림 나머지는 마지막 시트)
        share = round(100 / len(active), 1)
        for s in sheets:
            s['points'] = share if s in active else 0
        active[-1]['points'] = round(100 - share * (len(active) - 1), 1)
    for s in sheets:
        n = len(s['items'])
        good = sum(1 for i in s['items'] if i['ok'])
        s['got'] = s['points'] if n and good == n else round(s['points'] * good / n, 1) if n else 0
    score = round(sum(s['got'] for s in sheets), 1)
    total = sum(s['points'] for s in sheets)
    sections = []
    for sec in ('기본작업', '계산작업', '분석작업', '기타작업'):
        ss = [s for s in sheets if s['section'] == sec]
        if ss:
            sections.append({'name': sec, 'points': sum(s['points'] for s in ss), 'got': round(sum(s['got'] for s in ss), 1)})
    if not level:                                   # 시험 구성이 아니면 영역 대신 시트별로
        sections = [{'name': s['name'], 'points': s['points'], 'got': s['got']} for s in sheets if s['items']]
        total = 100
    return {'score': round(score, 1), 'total': round(total, 1), 'passed': score >= 70 * total / 100,
            'sheets': sheets, 'sections': sections, 'has_vba': bool(user.procs), 'level': level}



GRADE_SECONDS = 25


def _same_origin(src, user):
    """문제 파일의 글자 상수 중 40% 이상이 수험자 파일의 같은 시트에 남아 있는지(위치가 밀려도 되게 값 집합으로)."""
    total = kept = 0
    for ws in src.f.worksheets:
        wu = user.ws(ws.title)
        have = set()
        if wu is not None:
            for row in wu.iter_rows():
                for c in row:
                    if isinstance(c.value, str):
                        have.add(c.value.strip())
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and len(c.value.strip()) >= 2 and not c.value.startswith('='):
                    total += 1
                    kept += c.value.strip() in have
    return total < 8 or kept / total >= 0.4


def grade(src_bytes, ans_bytes, user_bytes, level=None):
    """채점(계산 시간 제한 {GRADE_SECONDS}초 — 넘으면 BadFile 로 안내)."""
    try:
        with fx.time_limit(GRADE_SECONDS):
            return _grade_unlimited(src_bytes, ans_bytes, user_bytes, level)
    except fx.TimeUp:
        raise xlsx.BadFile('파일 속 수식 계산이 너무 오래 걸려 채점을 멈췄습니다 — 아주 큰 범위·배열 수식을 줄여 다시 올려 주세요.')
