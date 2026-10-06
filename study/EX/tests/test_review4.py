"""4차 점검(필기·1급·새 문제)에서 확인된 문제들이 다시 생기지 않는지."""
import io
import re

import openpyxl

from core import content, written
from core import exam as ex
from core import formula as fx


def test_formula_fixes():
    b = fx.Book([fx.Sheet('S', {(r, 2): r * 10 for r in range(1, 11)})])
    assert fx.evaluate_text('=SUM(OFFSET(B10,0,0,-3,1))', b, 'S', 1, 5) == 270        # 음수 높이 = 위로
    assert fx.evaluate_text('=SUM(SIGN({-2,0,5}))', b, 'S', 1, 5) == 0
    assert fx.format_value(0.395833, 'h시간 m분') == '9시간 30분'                      # 글자 사이에 있어도 분
    assert fx.has_reference(fx.parse('=ROW()-1'))


def test_check_accepts_common_forms():
    ok = lambda pid, a: content.check(content.get(pid), a)['ok']   # noqa: E731
    assert ok('basic-101', '=ROWS($H$2:H2)') and ok('basic-101', '=ROW()-1')
    assert ok('array-101', '{=FREQUENCY($D$2:$D$13,$J$2:$J$5)}')                     # 범위에 한 번에 넣는 배열 수식
    assert ok('lookup-103', '=SUM(OFFSET(B2,COUNT(B2:B13)-1,0,-3,1))')
    assert not ok('date-104', '=MAX(0,(C2-B2)*24-9)*$K$1*1.5')                       # 반올림 지시
    r = content.check(content.get('lookup-102'), '=SUM(G2:INDEX(G:G,I2+1))')
    assert '아직 계산하지 못합니다' in r['error']


def test_exam_values_must_be_numbers():
    e = ex.get('c1-02')
    wb = openpyxl.load_workbook(io.BytesIO(ex.problem_workbook(e)))
    ws = wb['기타작업-3']
    for c, v in zip('ABCDEF', [1006, '윤서아', '전기차', '6', '74000', 444000]):    # 대여일수·일요금이 글자
        ws[f'{c}9'] = v
    bio = io.BytesIO()
    wb.save(bio)
    ctx = ex.Ctx(e, bio.getvalue(), 'x.xlsx')
    chk = next(c for t in e['tasks'] for c in t['checks'] if c['label'].startswith('② 등록 결과'))
    ok, msgs = ex.check_values(ctx, '기타작업-3', chk)
    assert not ok and '텍스트' in msgs[0]


def test_written_mock_paper_and_submit(client):
    ids = written.build_mock('c2', seed=3)
    assert written.valid_paper('c2', ids) and not written.valid_paper('c2', ids[:-1] + ids[:1])
    page = client.get('/written/mock/c2?q=' + ','.join(ids)).get_data(as_text=True)
    assert re.search(r'name="qids" value="([^"]+)"', page).group(1) == ','.join(ids)        # 이어 풀기
    comp = [written.get(i)['topic'] for i in ids if i.startswith('pc-')]
    assert comp.count('운영체제') >= comp.count('최신 기술')                                  # 문제 은행 비율
    h = {'X-CSRF-Token': client.csrf}
    assert client.post('/written/api/answer', json={'id': ids[0], 'picked': True}, headers=h).status_code == 400
    r = client.post('/written/mock/c2/submit', data={'_csrf': client.csrf, 'qids': ','.join([ids[0]] * 5 + ['db-0001']),
                                                     'seconds': '-99999'})
    res = client.get(r.headers['Location']).get_data(as_text=True)
    assert '-' not in re.search(r'평균 [\d.]+점', res).group(0)
    data = client.get('/api/backup').json['data']
    assert len(data['written_attempts']) == 1 and data['written_results'][0]['seconds'] == 0
    assert client.get('/written/practice?subject=database&level=c2').status_code == 302


def test_round5_fixes(tmp_path):
    import zipfile
    from core import analyze, compare, describe, xlsx
    # 합계·메모 행은 자료가 아님 / 두 줄 머리글 / 글자로 저장된 숫자·날짜
    t = analyze.to_table('S', [['지점', '매출', None], [None, '1분기', '2분기'], ['서울', 1, 2], ['합계', 1, 2], ['※ 기준', None, None]])
    assert [c['name'] for c in t['columns']] == ['지점', '매출 1분기', '매출 2분기'] and len(t['rows']) == 1 and t['skipped'] == 2
    assert analyze._csv_value('2026.01.05') is not None and analyze._csv_value('₩1,200원') == 1200
    # 휘발 함수와 모듈 이름
    assert compare._volatile_funcs('=YEAR(TODAY())-1') == {'TODAY'}
    assert compare._literals('Sub a()\n Range("A1").Value = 1\n Range("E19").Select\nEnd Sub') == \
        compare._literals('Sub a()\n Range("A1").Value = 1\nEnd Sub')
    assert compare._literals('Sub a()\n Range("A1").Interior.ColorIndex = 6\nEnd Sub') == \
        compare._literals('Sub a()\n Range("A1").Interior.Color = 65535\nEnd Sub')
    # 직접 지정한 행 높이만
    import openpyxl
    wb = openpyxl.Workbook()
    wb.active.row_dimensions[2].height = 30
    p = tmp_path / 'h.xlsx'
    wb.save(p)
    assert xlsx.custom_height_rows(p.read_bytes()).get('Sheet') == {2}
    with zipfile.ZipFile(p) as z:
        assert b'customHeight="1"' in z.read('xl/worksheets/sheet1.xml')
    # 시간 제한
    import pytest
    big = fx.parse('=SUMPRODUCT(SEQUENCE(2000)*TRANSPOSE(SEQUENCE(2000)))')
    with pytest.raises(fx.TimeUp):
        with fx.time_limit(0.3):
            fx.evaluate(big, fx.Book([fx.Sheet('S', {})]), 'S')
    del describe


def test_round6_ux(client):
    import io
    import openpyxl
    from core import compare, library
    page = client.get('/').get_data(as_text=True)
    assert 'skip-link' in page and 'id="main"' in page and '맞힌 문제' in page and '컴활 실기 실습' in page
    # 다른 문제의 파일은 0점으로 저장하지 않고 안내
    def book(sheet, val):
        wb = openpyxl.Workbook()
        wb.active.title = sheet
        wb.active['A1'] = val
        bio = io.BytesIO()
        wb.save(bio)
        return bio.getvalue()
    import pytest
    from core import xlsx
    with pytest.raises(xlsx.BadFile):
        compare.grade(book('가', None), book('가', 1), book('다른', 1))
    # 서식 메시지는 현재 → 정답
    from openpyxl.styles import Font
    def styled(bold):
        wb = openpyxl.Workbook()
        wb.active.title = 'S'
        wb.active['A1'] = 1
        wb.active['A1'].font = Font(b=bold)
        bio = io.BytesIO()
        wb.save(bio)
        return bio.getvalue()
    res = compare.grade(styled(False), styled(True), styled(False))
    msg = [m for s in res['sheets'] for i in s['items'] for m in i['msgs']][0]
    assert '현재' in msg and '정답' in msg
    assert compare._chart_show('범례', 'b') == '아래쪽' and compare._chart_show('차트 종류', {('B', 'col')}) == 'B: 세로 막대형'
    del library
