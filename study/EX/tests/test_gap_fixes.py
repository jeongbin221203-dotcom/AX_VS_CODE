"""채점 빈틈 보완: 피벗 원본 범위 · 고급 필터 조건 영역 · 보이지 않는 차트 계열 · 조건부 서식 함수 · 서식 코드 표시."""
import io
import json
from pathlib import Path

import openpyxl

from core import compare
from core import exam as ex
from core import pivots

FIX = Path(__file__).parent / 'fixtures'
PV = {'filters': [], 'rows': [], 'cols': [], 'values': [], 'sheet': 'S', 'ref': 'A1:B2',
      'source': ('분석작업-1', 'A3:F39', None)}


def test_pivot_source_range_must_match():
    assert pivots.match({'source': 'A3:F39'}, PV)[0]
    assert pivots.match({'source': '$A$3:$F$39'}, PV)[0]
    ok, why = pivots.match({'source': 'A3:F30'}, dict(PV, source=('분석작업-1', 'A3:F25', None)))
    assert not ok and 'A3:F25' in why[0]


def test_exam_answers_still_pass_with_source_checks():
    for eid, fix in (('c1-01', 'c1-01_answer_excel.xlsm'), ('c2-01', 'c2-01_answer_excel.xlsx')):
        res = ex.grade(ex.get(eid), (FIX / fix).read_bytes(), fix)
        bad = [(t['no'], i['label'], i['msgs']) for t in res['tasks'] for i in t['items']
               if not i['ok'] and '매크로' not in i['label']]
        assert bad == [], (eid, bad)


def test_filter_criteria_range_checked():
    """조건 영역([A25:A26])의 머리글 아래 칸이 비어 있으면(조건을 다른 곳에 썼으면) 틀림."""
    e = ex.get('c1-01')
    wb = openpyxl.load_workbook(FIX / 'c1-01_answer_excel.xlsm', keep_vba=True)
    ws = wb['기본작업-1']
    assert ws['A26'].value
    ws['C40'] = ws['A26'].value                     # 조건을 영역 밖으로 옮기고
    ws['A26'] = None
    bio = io.BytesIO()
    wb.save(bio)
    res = ex.grade(e, bio.getvalue(), '답.xlsm')
    msgs = [m for t in res['tasks'] for i in t['items'] for m in i['msgs']]
    assert any('[A25:A26]' in m for m in msgs), msgs


def test_hidden_series_detected():
    from openpyxl.chart.series import Series
    from openpyxl.chart.shapes import GraphicalProperties
    s = Series()
    assert not ex._series_hidden(s, 'col')
    s.spPr = GraphicalProperties(noFill=True)
    assert not ex._series_hidden(s, 'col')           # 채우기만 없고 윤곽선이 보이면 보이는 계열
    s.spPr.ln.noFill = True
    assert ex._series_hidden(s, 'col')


def test_cf_required_functions():
    assert compare._formula_funcs('AND(OR($B3="a",$B3="b"),$H3>=LARGE($H$3:$H$22,5))') == {'AND', 'OR', 'LARGE'}
    assert compare._formula_funcs('$B3*$H3') == set()
    assert compare._formula_funcs('((') == set()


def test_format_label_is_readable():
    assert ex._fmt_label('General') == '일반'
    assert ex._fmt_label('@') == "'텍스트'"
    assert ex._fmt_label('#,##0').startswith('#,##0 (예: 1,23')


def test_exam_json_extras_valid():
    for eid in ('c1-01', 'c1-02', 'c2-02'):
        assert ex.validate(ex.get(eid)) == []
    checks = [c for t in ex.get('c1-02')['tasks'] for c in t['checks'] if c.get('kind') == 'filter']
    assert checks[0]['criteria_head'] == '조건'
    json.dumps(checks)
