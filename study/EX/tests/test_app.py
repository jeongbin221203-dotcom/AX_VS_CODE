import io

from openpyxl import Workbook

from core import build, content

FIX = __import__('pathlib').Path(__file__).parent / 'fixtures'


def post_json(client, url, body):
    return client.post(url, json=body, headers={'X-CSRF-Token': client.csrf})


def test_pages_render(client):
    for url in ['/', '/learn', '/learn/lookup', '/review', '/functions', '/functions/VLOOKUP', '/functions/RANK',
                '/analyze/', '/build/', '/build/sales', '/p/basic-001', '/p/basic-004']:
        r = client.get(url)
        assert r.status_code == 200, url
        assert "script-src 'self'" in r.headers['Content-Security-Policy']


def test_every_problem_page_renders(client):
    for p in content.problems():
        assert client.get(f"/p/{p['id']}").status_code == 200, p['id']


def test_check_records_attempt_and_dashboard(client):
    r = post_json(client, '/api/check', {'pid': 'basic-001', 'answer': '=SUM(G2:G16)'})
    assert r.json['ok'] and r.json['answer'] == '=SUM(G2:G16)'
    r = post_json(client, '/api/check', {'pid': 'basic-004', 'answer': 0})
    assert r.json['ok'] is False and r.json['answer'] == 1
    page = client.get('/').get_data(as_text=True)
    assert '참조 전환 키' in page            # 다시 볼 문제
    assert '50' in page                      # 첫 시도 정답률 50%
    assert '오답 노트' in client.get('/review').get_data(as_text=True)


def test_try_mode_not_recorded(client):
    r = post_json(client, '/api/check', {'pid': 'basic-001', 'answer': '=SUM(G2:G16)', 'mode': 'try'})
    assert 'ok' not in r.json
    assert '이어서 풀기' not in client.get('/').get_data(as_text=True)


def test_parse_error_message(client):
    r = post_json(client, '/api/check', {'pid': 'basic-001', 'answer': '=SUM(G2:G16'})
    assert '읽을 수 없습니다' in r.json['error']


def test_csrf_required(client):
    r = client.post('/api/check', json={'pid': 'basic-001', 'answer': '=1'})
    assert r.status_code == 400


def test_star_and_track(client):
    assert post_json(client, '/api/star', {'pid': 'basic-001'}).json['on'] is True
    assert '★' in client.get('/learn/basic').get_data(as_text=True)
    client.post('/track', data={'track': 'c1', '_csrf': client.csrf, 'next': '/learn'})
    page = client.get('/learn').get_data(as_text=True)
    assert 'selected>컴활 1급' in page.replace('" selected', ' selected').replace('"c1" ', '')


def test_next_redirects_to_unsolved(client):
    r = client.get('/next')
    assert r.status_code == 302 and '/p/' in r.headers['Location']


def test_analyze_sample_flow(client):
    r = client.post('/analyze/sample', data={'_csrf': client.csrf})
    assert r.status_code == 302
    url = r.headers['Location']
    page = client.get(url).get_data(as_text=True)
    assert 'SUMIFS(' in page and '피벗 표' in page and '월별' in page
    page = client.get(url + '?g=4&a=avg&c=2&f=2&fv=서울').get_data(as_text=True)
    assert '평균' in page
    x = client.get(url + '/export')
    assert x.status_code == 200 and x.data[:2] == b'PK'


def test_analyze_upload_csv_and_bad_file(client):
    data = '이름,부서,점수\n가,A,10\n나,B,20\n다,A,30\n'.encode('cp949')
    r = client.post('/analyze/upload', data={'_csrf': client.csrf, 'file': (io.BytesIO(data), 'a.csv')})
    assert r.status_code == 302 and '/analyze/' in r.headers['Location']
    page = client.get(r.headers['Location']).get_data(as_text=True)
    assert '부서' in page
    r = client.post('/analyze/upload', data={'_csrf': client.csrf, 'file': (io.BytesIO(b'not zip'), 'b.xlsx')},
                    follow_redirects=True)
    assert '.xlsx' in r.get_data(as_text=True)


def test_analyze_header_not_first_row(client):
    wb = Workbook()
    ws = wb.active
    ws['A1'] = '2026년 실적 보고'
    ws.append([])
    ws.append(['팀', '실적'])
    for t, v in [('1팀', 10), ('2팀', 20), ('1팀', 5)]:
        ws.append([t, v])
    bio = io.BytesIO()
    wb.save(bio)
    r = client.post('/analyze/upload', data={'_csrf': client.csrf, 'file': (io.BytesIO(bio.getvalue()), 'r.xlsx')})
    page = client.get(r.headers['Location']).get_data(as_text=True)
    assert '머리글 3행' in page and '$B$4:$B$6' in page


def test_build_download_and_grade(client):
    r = client.get('/build/sales/file')
    assert r.status_code == 200 and r.data[:2] == b'PK'
    ans = build.workbook(build.mission('sales'), answers=True)
    r = client.post('/build/sales/submit', data={'_csrf': client.csrf, 'file': (io.BytesIO(ans), 'done.xlsx')})
    assert r.status_code == 302
    page = client.get(r.headers['Location']).get_data(as_text=True)
    assert '11<small> / 11' in page
    assert '최고 <b>11/11</b>' in client.get('/').get_data(as_text=True)


def test_build_rejects_wrong_workbook(client):
    wb = Workbook()
    bio = io.BytesIO()
    wb.save(bio)
    r = client.post('/build/hr/submit', data={'_csrf': client.csrf, 'file': (io.BytesIO(bio.getvalue()), 'x.xlsx')},
                    follow_redirects=True)
    assert '시트가 있어야 합니다' in r.get_data(as_text=True)


def _public_app(tmp_path):
    from app import create_app
    return create_app({'DATA_DIR': str(tmp_path), 'DATABASE': str(tmp_path / 'ex.db'), 'SECRET_KEY': 't',
                       'PUBLIC': True, 'OWNER_TOKEN': 'k' * 40, 'TESTING_NO_CSRF': True})


def test_public_users_are_isolated(tmp_path):
    app = _public_app(tmp_path)
    a, b = app.test_client(), app.test_client()
    a.post('/api/check', json={'pid': 'basic-001', 'answer': '=SUM(G2:G16)'})
    a.post('/analyze/sample')
    r = a.post('/exam/c2-01/submit', data={'file': (open(FIX / 'c2-01_answer_excel.xlsx', 'rb'), 'x.xlsx')})
    result_url = r.headers['Location']
    a.post('/track', data={'track': 'c1'})
    assert '1<small> / ' in a.get('/').get_data(as_text=True) or '푼 문제' in a.get('/').get_data(as_text=True)
    page_b = b.get('/').get_data(as_text=True)
    assert '샘플_상반기매출' not in page_b and '최고 <b>94점</b>' not in page_b
    assert b.get(result_url).status_code == 404
    assert '샘플_상반기매출' not in b.get('/analyze/').get_data(as_text=True)
    assert 'selected>컴활 1급' not in b.get('/learn').get_data(as_text=True).replace('" selected', ' selected')


def test_owner_devices_share_records(tmp_path):
    app = _public_app(tmp_path)
    phone, pc = app.test_client(), app.test_client()
    phone.get('/me/' + 'k' * 40)
    pc.get('/me/' + 'k' * 40)
    phone.post('/api/check', json={'pid': 'basic-004', 'answer': 1})
    assert '참조 전환 키' not in pc.get('/review').get_data(as_text=True)          # 맞혀서 오답 노트에 없음
    assert '1<small> / ' in pc.get('/').get_data(as_text=True)                    # 푼 문제 1
