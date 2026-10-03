import io

import openpyxl
import pytest

from core import build


@pytest.mark.parametrize('key', list(build.MISSIONS))
def test_blank_scores_zero_and_answer_full(key):
    m = build.mission(key)
    assert build.grade(m, build.workbook(m))['score'] == 0
    res = build.grade(m, build.workbook(m, answers=True))
    assert res['score'] == res['total']


def _edit(m, fn):
    wb = openpyxl.load_workbook(io.BytesIO(build.workbook(m, answers=True)))
    fn(wb[build.DASH])
    bio = io.BytesIO()
    wb.save(bio)
    return build.grade(m, bio.getvalue())


def item(res, iid):
    return next(i for i in res['items'] if i['id'] == iid)


def test_typed_value_is_not_accepted():
    m = build.mission('sales')
    ref = build.grade(m, build.workbook(m, answers=True))
    total = item(ref, 'total')['cells'][0]['expected'].replace(',', '')
    res = _edit(m, lambda ws: ws.__setitem__('B3', int(total)))
    t = item(res, 'total')
    assert not t['ok'] and '직접 입력' in t['msgs'][0]
    # 비중은 B3 값을 쓰므로 여전히 맞다
    assert item(res, 'share')['ok']


def test_missing_dollar_fails_fill():
    m = build.mission('sales')

    def broken(ws):
        for r in range(9, 14):
            ws[f'B{r}'] = f'=SUMIFS(데이터!I2:I121,데이터!C2:C121,A{r})'.replace('I2:I121', f'I{r - 7}:I{r + 112}') \
                .replace('C2:C121', f'C{r - 7}:C{r + 112}')
    res = _edit(m, broken)
    assert not item(res, 'region')['ok']


def test_alternative_formula_accepted():
    m = build.mission('sales')

    def alt(ws):
        for r in range(9, 15):
            ws[f'F{r}'] = f'=SUMPRODUCT((MONTH(데이터!$B$2:$B$121)=E{r})*데이터!$I$2:$I$121)'
        ws['B5'] = '=B3/B4'
    res = _edit(m, alt)
    assert item(res, 'month')['ok'] and item(res, 'avg')['ok']


def test_dropdown_check():
    m = build.mission('inventory')
    res = build.grade(m, build.workbook(m, answers=True))
    assert item(res, 'dv')['ok']

    def no_dv(ws):
        ws.data_validations.dataValidation = []
    assert not item(_edit(m, no_dv), 'dv')['ok']


def test_changed_dropdown_value_still_graded_correctly():
    m = build.mission('inventory')
    res = _edit(m, lambda ws: ws.__setitem__('B15', '전자'))
    assert item(res, 'pick')['ok']


FIXTURES = __import__('pathlib').Path(__file__).parent / 'fixtures'


def test_pivot_contents_graded_from_real_excel_file():
    """엑셀(COM)로 만든 피벗: 행 지역 · 열 분류 · 값 금액 합계."""
    m = build.mission('sales')
    res = build.grade(m, (FIXTURES / 'sales_pivot_excel.xlsx').read_bytes())
    p = item(res, 'pivot')
    assert p['ok'], p['msgs']
    assert res['score'] == res['total']


def test_pivot_wrong_fields_reported():
    from core import pivots
    wb_f = openpyxl.load_workbook(FIXTURES / 'sales_pivot_excel.xlsx')
    pv = pivots.describe(wb_f)[0]
    assert pv['rows'] == ['지역'] and pv['cols'] == ['분류'] and pv['values'] == [('금액', 'sum')]
    ok, why = pivots.match({'rows': ['분류'], 'values': [['금액', 'average']]}, pv)
    assert not ok and any('행 영역' in w for w in why) and any('평균' in w for w in why)
    ok, why = pivots.match({'sheet': 'PV', 'at': 'A3'}, pv)
    assert ok, why
