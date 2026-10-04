"""채점 검토(2026-10-04)에서 확인된 문제들이 다시 생기지 않는지."""
import io

import openpyxl
import pytest
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Font, PatternFill
from openpyxl.styles.colors import Color
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.scenario import InputCells, Scenario, ScenarioList

from core import compare, xlsx
from core import exam as ex


def _save(wb):
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


def _book(edit=None, sheets=('Data',)):
    wb = openpyxl.Workbook()
    wb.active.title = sheets[0]
    for name in sheets[1:]:
        wb.create_sheet(name)
    for name in sheets:
        ws = wb[name]
        ws['A1'], ws['B1'] = '이름', '금액'
        for i in range(2, 7):
            ws.cell(i, 1, f'p{i}')
            ws.cell(i, 2, i * 100)
    if edit:
        edit(wb)
    return _save(wb)


def _fails(res):
    return [i['label'] for s in res['sheets'] for i in s['items'] if not i['ok']]


def test_chartsheet_is_skipped():
    def chart(wb):
        ws = wb['Data']
        ws['C1'] = '완료'
        cs = wb.create_chartsheet('Chart1')
        ch = BarChart()
        ch.add_data(Reference(ws, min_col=2, min_row=1, max_row=6))
        cs.add_chart(ch)
    src = _book(lambda wb: chart(wb) or wb['Data'].__setitem__('C1', None))
    ans = _book(chart)
    assert compare.grade(src, ans, ans)['score'] == 100
    assert compare.grade(src, ans, src)['score'] == 0


def test_pair_without_common_sheet_or_tasks_is_badfile():
    with pytest.raises(xlsx.BadFile):
        compare.grade(_book(sheets=('Sheet1',)), _book(sheets=('Data',)), _book())
    same = _book()
    with pytest.raises(xlsx.BadFile):
        compare.grade(same, same, same)


def test_theme_fill_is_graded():
    def theme(wb):
        c = wb['Data']['A1']
        c.fill = PatternFill('solid', fgColor=Color(theme=4))
        c.font = Font(color=Color(theme=0), b=True)
    src, ans = _book(), _book(theme)
    labels = _fails(compare.grade(src, ans, src))
    assert any(x.startswith('채우기') for x in labels) and any(x.startswith('글꼴 색') for x in labels)

    def other_theme(wb):
        theme(wb)
        wb['Data']['A1'].fill = PatternFill('solid', fgColor=Color(theme=5))
    assert any(x.startswith('채우기') for x in _fails(compare.grade(src, ans, _book(other_theme))))


def test_orientation_unset_equals_portrait_and_accounting_format():
    def ans_edit(wb):
        ws = wb['Data']
        ws.page_setup.orientation = 'portrait'
        for r in range(2, 7):
            ws.cell(r, 2).number_format = '_(* #,##0_);_(* \\(#,##0\\);_(* "-"_);_(@_)'
    def user_edit(wb):
        for r in range(2, 7):
            wb['Data'].cell(r, 2).number_format = '_-* #,##0_-;\\-* #,##0_-;_-* "-"_-;_-@_-'   # 한국어 '쉼표 스타일'
    res = compare.grade(_book(), _book(ans_edit), _book(user_edit))
    assert res['score'] == 100, _fails(res)


def test_style_items_split_by_target_value():
    from openpyxl.styles import Alignment

    def src_edit(wb):
        wb['Data']['B2'].alignment = Alignment(vertical='center')
    def ans_edit(wb):
        for r in range(1, 7):
            wb['Data'].cell(r, 1).alignment = Alignment(vertical='center')
        wb['Data']['B2'].alignment = Alignment(vertical='bottom')
    labels = [i['label'] for s in compare.grade(_book(src_edit), _book(ans_edit), _book(src_edit))['sheets']
              for i in s['items']]
    assert '세로 맞춤 A1:A6' in labels and '세로 맞춤 B2' in labels


def test_cf_and_dv_conditions_are_compared():
    red = PatternFill('solid', bgColor='FFC7CE')

    def ans_edit(wb, op='lessThan', f2='300'):
        ws = wb['Data']
        ws.conditional_formatting.add('B2:B6', CellIsRule(operator=op, formula=['200'], fill=red))
        dv = DataValidation(type='whole', operator='between', formula1='0', formula2=f2)
        dv.add('B2:B6')
        ws.add_data_validation(dv)
    src, ans = _book(), _book(ans_edit)
    assert compare.grade(src, ans, ans)['score'] == 100
    wrong = _fails(compare.grade(src, ans, _book(lambda wb: ans_edit(wb, 'greaterThan', '999'))))
    assert '조건부 서식 B2:B6' in wrong and '데이터 유효성 B2:B6' in wrong


def test_scenario_values_are_compared():
    def ans_edit(wb, val='0.2'):
        wb['Data'].scenarios = ScenarioList(scenario=[Scenario(name='인상', count=1,
                                                               inputCells=[InputCells(r='B2', val=val)])])
    src, ans = _book(), _book(ans_edit)
    assert compare.grade(src, ans, _book(lambda wb: ans_edit(wb, '0.20')))['score'] == 100
    assert compare.grade(src, ans, _book(lambda wb: ans_edit(wb, '12345')))['score'] == 0


def test_section_points_add_up_to_total():
    def ans_edit(wb):
        for name in ('A', 'B', 'C'):
            wb[name]['C1'] = '완료'
    sheets = ('A', 'B', 'C')
    src, ans = _book(sheets=sheets), _book(ans_edit, sheets=sheets)
    res = compare.grade(src, ans, ans)
    assert res['score'] == 100 and round(sum(s['points'] for s in res['sections']), 1) == 100
    assert round(sum(s['got'] for s in res['sections']), 1) == res['score']


def test_vba_key_values_matter():
    want = '''Sub 채우기()
    Range("A3:E3").Select
    With Selection.Interior
        .Pattern = xlSolid
        .Color = 65535
        .TintAndShade = 0
    End With
End Sub'''
    wrong = want.replace('65535', '255')
    alt = 'Sub 채우기()\n    Range("A3:E3").Interior.Color = 65535\nEnd Sub'

    def ok(got):
        s = compare._similar(got, want)
        return s >= 0.95 or (s >= 0.3 and not compare._literals(want) - compare._literals(got))
    assert ok(want) and not ok(wrong) and ok(alt.replace('Range("A3:E3").Interior', 'Range("A3:E3").Select\n    Selection.Interior'))


def test_exam_scenario_pairs_values_in_task_order():
    e = ex.get('c2-02')
    wb = openpyxl.load_workbook(io.BytesIO(ex.problem_workbook(e)))
    wb['분석작업-1'].scenarios = ScenarioList(scenario=[Scenario(name='인상', count=2, inputCells=[
        InputCells(r='C3', val='0.2'), InputCells(r='B5', val='5')])])
    ctx = ex.Ctx(e, _save(wb), 'x.xlsx')
    chk = {'changing': 'C3,B5', 'scenarios': [['인상', [0.2, 5]]], 'summary': False}
    assert ex.check_scenario(ctx, '분석작업-1', chk) == (True, [])


def test_exam_subtotal_needs_group_values():
    e = ex.get('c2-01')
    t = next(t for t in e['tasks'] if t['no'] == '3-2')
    wb = openpyxl.load_workbook(io.BytesIO(ex.problem_workbook(e)))
    ws = wb['분석작업-2']
    for i, r in enumerate(range(20, 24)):          # 정렬도 안 하고 한 행씩 묶은 엉터리 부분합
        ws.cell(r, 5).value = f'=SUBTOTAL(4,E{4 + i}:E{4 + i})'
        ws.cell(r, 6).value = f'=SUBTOTAL(1,F{4 + i}:F{4 + i})'
    ctx = ex.Ctx(e, _save(wb), 'x.xlsx')
    assert not any(ex.check_subtotal(ctx, '분석작업-2', chk)[0] for chk in t['checks'])


def test_vba_formula_with_same_result_is_accepted():
    def book(edit=None):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = '매크로작업'
        for r in range(4, 9):
            ws.cell(r, 3, r * 10)
            ws.cell(r, 4, r)
        return compare.Wb(_save(wb))
    want = 'Sub 총점()\n    Range("E4").Select\n    ActiveCell.FormulaR1C1 = "=RC[-2]+RC[-1]"\n' \
           '    Selection.AutoFill Destination:=Range("E4:E8")\nEnd Sub'
    b = book().book

    def ok(got):
        s = compare._similar(got, want)
        la, lu = compare._literals(want), compare._literals(got)
        miss = compare._same_result_formulas(la - lu, lu, la, b)
        return s >= 0.95 or (s >= 0.3 and not miss)
    assert ok('Sub 총점()\n    Range("E4:E8").FormulaR1C1 = "=SUM(RC[-2]:RC[-1])"\nEnd Sub')
    assert ok('Sub 총점()\n    Range("E4:E8").Formula = "=C4+D4"\nEnd Sub')
    assert not ok(want.replace('RC[-2]+RC[-1]', 'RC[-2]*RC[-1]'))
    assert compare.r1c1_to_a1('=IF(RC[-1]>=80,"RC","")', 4, 5) == '=IF(D4>=80,"RC","")'
