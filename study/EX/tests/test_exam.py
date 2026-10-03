"""모의고사 채점: 실제 Excel(COM)로 만든 정답 파일을 고정 자료로 쓴다 (tools/exam_answer.py 로 다시 만들 수 있음)."""
import io
import json
from pathlib import Path

import openpyxl
import pytest

from core import exam as ex

FIX = Path(__file__).parent / 'fixtures'


def _items(res):
    return {(t['no'], i['label']): i for t in res['tasks'] for i in t['items']}


def test_all_exams_valid():
    for e in ex.exams():
        assert ex.validate(e) == [], e['id']


def test_c2_01_excel_answer_scores_all_but_macros():
    e = ex.get('c2-01')
    res = ex.grade(e, (FIX / 'c2-01_answer_excel.xlsx').read_bytes())
    vba_pts = sum(c['points'] for t in e['tasks'] for c in t['checks'] if c.get('needs_vba'))
    assert res['score'] == res['total'] - vba_pts
    fails = [k for k, v in _items(res).items() if not v['ok']]
    assert all('매크로' in label for _, label in fails), fails


def test_c2_01_problem_file_scores_almost_nothing():
    e = ex.get('c2-01')
    res = ex.grade(e, ex.problem_workbook(e))
    assert res['score'] == 0
    assert not res['passed']


def test_check_kinds_from_excel_file():
    """고급 필터·유효성·메모·페이지·보호·데이터 표·시나리오·목표값·통합·텍스트 나누기."""
    e = json.loads((FIX / 'checks_exam.json').read_text(encoding='utf-8'))
    res = ex.grade(e, (FIX / 'checks_answer_excel.xlsx').read_bytes())
    bad = {k: v['msgs'] for k, v in _items(res).items() if not v['ok']}
    assert bad == {}
    assert res['score'] == 100


def _modified(fn):
    wb = openpyxl.load_workbook(FIX / 'c2-01_answer_excel.xlsx')
    fn(wb)
    bio = io.BytesIO()
    wb.save(bio)
    return ex.grade(ex.get('c2-01'), bio.getvalue())


def test_wrong_function_and_typed_values_rejected():
    def edit(wb):
        ws = wb['계산작업']
        for r in range(4, 12):                      # IF 대신 IFS → 지시 함수 위반
            ws[f'F{r}'] = f'=IFS(AVERAGE(C{r}:E{r})>=85,"우수",AVERAGE(C{r}:E{r})>=70,"보통",TRUE,"노력")'
        ws['E24'] = 13                               # 값만 입력
    res = _items(_modified(edit))
    assert not res[('2-1', '평가 F4:F11')]['ok'] and 'IF' in res[('2-1', '평가 F4:F11')]['msgs'][0]
    assert not res[('2-3', '강남 평균 수량 E24')]['ok']


def test_cf_equivalent_formula_accepted_and_wrong_anchor_rejected():
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import Font

    def equivalent(wb):
        ws = wb['기본작업-3']
        ws.conditional_formatting = type(ws.conditional_formatting)()
        ws.conditional_formatting.add('A4:F15', FormulaRule(formula=['AND($E4>299,$B4="서울")'],
                                                            font=Font(bold=True, color='FF0070C0')))
    assert _items(_modified(equivalent))[('1-3', '조건부 서식(수식 규칙·굵게·파랑)')]['ok']

    def no_dollar(wb):
        ws = wb['기본작업-3']
        ws.conditional_formatting = type(ws.conditional_formatting)()
        ws.conditional_formatting.add('A4:F15', FormulaRule(formula=['AND(B4="서울",E4>=300)'],
                                                            font=Font(bold=True, color='FF0070C0')))
    item = _items(_modified(no_dollar))[('1-3', '조건부 서식(수식 규칙·굵게·파랑)')]
    assert not item['ok'] and '$' in item['msgs'][0]


def test_pages(client):
    assert client.get('/exam/').status_code == 200
    page = client.get('/exam/c2-01').get_data(as_text=True)
    assert '문제 2 계산작업' in page and 'data-exam-timer' in page
    r = client.get('/exam/c2-01/file')
    assert r.status_code == 200 and r.data[:2] == b'PK'


def test_submit_and_result(client):
    data = (FIX / 'c2-01_answer_excel.xlsx').read_bytes()
    r = client.post('/exam/c2-01/submit', data={'_csrf': client.csrf, 'seconds': '1830',
                                               'file': (io.BytesIO(data), '답안.xlsx')})
    assert r.status_code == 302
    page = client.get(r.headers['Location']).get_data(as_text=True)
    assert '94' in page and '합격선' in page and '30분 30초' in page
    assert '최고 <b>94점</b>' in client.get('/').get_data(as_text=True)


@pytest.mark.parametrize('name', ['x.txt', 'x.xls'])
def test_submit_rejects_other_files(client, name):
    r = client.post('/exam/c2-01/submit', data={'_csrf': client.csrf, 'file': (io.BytesIO(b'abc'), name)},
                    follow_redirects=True)
    assert '.xlsx 또는 .xlsm' in r.get_data(as_text=True)


def test_every_exam_page_and_file(client):
    for e in ex.exams():
        assert client.get(f"/exam/{e['id']}").status_code == 200, e['id']
        assert ex.grade(e, ex.problem_workbook(e))['score'] == 0, e['id']
    assert {e['level'] for e in ex.exams()} == {'c2', 'c1'}
