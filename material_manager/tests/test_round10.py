"""회사 엑셀을 올리면 그 양식을 기억해 내려받을 때 같은 모양으로 받는다 (core/excel_forms.py learn·adopt)."""
from __future__ import annotations

import io
import os
import re
import sys
import tempfile
from pathlib import Path

os.environ.setdefault("MM_DB_PATH", str(Path(tempfile.mkdtemp(prefix="mm_r10_")) / "test.db"))
os.environ.setdefault("MM_PW_ITERATIONS", "1000")
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import openpyxl  # noqa: E402
import pandas as pd  # noqa: E402
from openpyxl.styles import Border, Font, PatternFill, Side  # noqa: E402

import config  # noqa: E402
from core import db, excel_forms, seed, storage  # noqa: E402,F401
from test_app import app, login, post  # noqa: E402,F401

THIN = Side(style="thin")


def _company_xlsx(title_rows: int = 2, headers=None, rows=None, merge_title: bool = True) -> bytes:
    """회사에서 쓰는 자재 목록 엑셀: 제목·작성자 줄 + 머리글 + 데이터 (머리글이 title_rows+1 행)."""
    headers = headers or ["품번", "품명", "규격", "담당부서", "단위", "분류", "단가"]
    rows = rows or [["BRG-100", "베어링 6205", "25x52", "생산1팀", "EA", "부품", 1500],
                    ["BLT-200", "벨트 A-40", "A40", "생산2팀", "EA", "부품", 2300]]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "자재목록"
    if title_rows >= 1:
        ws["A1"] = "㈜한빛 자재 현황"
        ws["A1"].font = Font(bold=True, size=14)
        if merge_title:
            ws.merge_cells("A1:C1")
    if title_rows >= 2:
        ws["A2"] = "작성: {{작성자}}"
    hr = title_rows + 1
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=hr, column=i, value=h)
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="FFFF00")
    for r, row in enumerate(rows, start=hr + 1):
        for i, v in enumerate(row, start=1):
            c = ws.cell(row=r, column=i, value=v)
            c.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
            if headers[i - 1] == "단가":
                c.number_format = "#,##0"
    ws.column_dimensions["B"].width = 30
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _master_df() -> pd.DataFrame:
    cols = list(config.MATERIAL_COLS.values())
    row = {c: "" for c in cols}
    row.update({"자재코드": "X-1", "자재명": "시험 자재", "규격": "S", "단위": "EA", "분류": "부품", "단가": 1234.5, "공급처": "공급A"})
    return pd.DataFrame([row, {**row, "자재코드": "X-2", "자재명": "시험 자재2", "단가": 99000}])


# ── learn: 머리글 찾기·열 연결·데이터 비우기 ─────────────────────
def test_learn_detects_header_row_maps_columns_and_drops_data(app):  # noqa: F811
    data = _company_xlsx()
    out = excel_forms.learn("materials_master", data, "한빛.xlsx", "material_upload")
    assert out["ok"], out.get("problem")
    s, cfg = out["summary"], out["cfg"]
    assert s["header_row"] == 3 and s["sheet"] == "자재목록"                     # 제목·작성자 줄을 건너뛰고 머리글 행을 찾는다
    assert s["unmapped"] == ["담당부서"] and "안전재고" in s["extra"]            # 모르는 열은 비워 두고, 없던 항목은 덧붙인다
    by_header = {c["header"]: c for c in cfg["columns"]}
    assert by_header["품번"]["source"] == "자재코드" and by_header["품번"]["col"] == 1
    assert by_header["담당부서"]["source"] == "" and by_header["담당부서"]["col"] == 4
    assert by_header["단가"]["source"] == "단가" and by_header["단가"]["col"] == 7
    assert cfg["start_row"] == 4 and cfg["write_header"] is False
    wb = openpyxl.load_workbook(io.BytesIO(out["template"]))
    ws = wb.active
    texts = {str(c.value) for row in ws.iter_rows() for c in row if c.value is not None}
    assert "BRG-100" not in texts and "베어링 6205" not in texts and "1500" not in texts      # 데이터는 저장하지 않는다
    assert "㈜한빛 자재 현황" in texts and "품번" in texts                                  # 제목·머리글은 남는다
    assert ws["A3"].font.bold and ws["G4"].number_format == "#,##0" and ws.max_row == 4      # 첫 데이터 행 서식만 남음


def test_learn_rejects_bad_files(app):  # noqa: F811
    assert not excel_forms.learn("materials_master", b"a,b\n1,2", "x.csv")["ok"]
    assert not excel_forms.learn("materials_master", b"junk", "x.xlsx")["ok"]
    assert not excel_forms.learn("nope", _company_xlsx(), "x.xlsx")["ok"]
    unknown = _company_xlsx(headers=["가", "나", "다"], rows=[[1, 2, 3]])
    out = excel_forms.learn("materials_master", unknown, "x.xlsx")
    assert not out["ok"] and "2개 이상" in out["problem"]                                  # 연결되는 열이 없으면 기억하지 않음


# ── 내려받기: 회사 양식 그대로 ──────────────────────────────────
def test_export_uses_learned_layout(app):  # noqa: F811
    out = excel_forms.learn("materials_master", _company_xlsx(), "한빛.xlsx", "material_upload")
    excel_forms.adopt("materials_master", out["cfg"], out["template"], "한빛.xlsx", {"id": None, "name": "관리자", "ip": ""})
    assert excel_forms.load("materials_master")["has_template"]
    data = excel_forms.export("materials_master", _master_df(), {"user": "홍길동", "period": ""})
    ws = openpyxl.load_workbook(io.BytesIO(data)).active
    assert ws["A1"].value == "㈜한빛 자재 현황" and ws["A2"].value == "작성: 홍길동"     # 제목 유지·자리표시 채움
    assert [ws.cell(row=3, column=c).value for c in range(1, 8)] == ["품번", "품명", "규격", "담당부서", "단위", "분류", "단가"]
    assert ws["A4"].value == "X-1" and ws["B4"].value == "시험 자재" and ws["C4"].value == "S"
    assert ws["D4"].value is None and ws["D5"].value is None                                 # 연결 안 된 열은 비워 둔다 (밀리지 않음)
    assert ws["E4"].value == "EA" and ws["F4"].value == "부품" and ws["G4"].value == 1234.5
    assert ws["A5"].value == "X-2" and ws["G5"].value == 99000 and ws["G5"].number_format == "#,##0"   # 둘째 행도 같은 서식
    assert ws["A5"].border.left.style == "thin" and ws["A3"].font.bold
    extra = {ws.cell(row=3, column=c).value for c in range(8, 20)}
    assert {"안전재고", "공급처", "바코드"} <= extra                                          # 파일에 없던 항목은 오른쪽에 덧붙임
    supplier_col = next(c for c in range(8, 20) if ws.cell(row=3, column=c).value == "공급처")
    assert ws.cell(row=4, column=supplier_col).value == "공급A"
    # 설정 저장 화면을 거쳐도 열 위치가 유지된다 (save_export 가 col 을 보존)
    cfg = excel_forms.load("materials_master")
    assert excel_forms.save_export("materials_master", cfg["columns"], cfg["sheet"], cfg["header_row"], cfg["start_row"],
                                   cfg["start_col"], False, "", None) == ""
    assert {c["header"]: c.get("col") for c in excel_forms.load("materials_master")["columns"]}["단가"] == 7


def test_default_export_unchanged_without_learning(app):  # noqa: F811
    data = excel_forms.export("materials_master", _master_df(), {})
    ws = openpyxl.load_workbook(io.BytesIO(data)).active
    assert ws["A1"].value == "자재코드" and ws["A2"].value == "X-1"                           # 기억한 양식이 없으면 이전과 같다


# ── 화면 흐름 ───────────────────────────────────────────────────
def _upload(client, url, xlsx, name, follow_redirects=False, **form):
    return post(client, url, {"file": (io.BytesIO(xlsx), name), **form}, content_type="multipart/form-data",
                follow_redirects=follow_redirects)


def test_material_upload_remembers_layout_and_download_follows(app):  # noqa: F811
    seed.seed()
    client = login(app.test_client(), "admin")
    company = _company_xlsx(title_rows=0, headers=["품번", "품명", "규격", "단위", "분류", "단가"],
                            rows=[["NEW-1", "새 자재1", "R", "EA", "부품", 700], ["NEW-2", "새 자재2", "R", "EA", "부품", 800]])
    page = client.get("/data/").get_data(as_text=True)
    assert "내려받기 양식으로 기억" in page and "checked" in page.split("learn_form")[1][:80]       # 처음에는 기본으로 체크
    res = _upload(client, "/data/upload", company, "회사자재.xlsx", learn_form="1")
    html = res.get_data(as_text=True)
    assert "양식을 읽었습니다" in html and "자재 마스터" in html
    token = re.search(r'name="token" value="([0-9a-f]{32})"', html).group(1)
    assert not excel_forms.load("materials_master")["has_template"]                          # 반영 전에는 저장하지 않음
    res = post(client, "/data/upload/apply", {"token": token}, follow_redirects=True)
    assert "올린 파일의 양식으로 바뀌었습니다" in res.get_data(as_text=True)
    assert excel_forms.load("materials_master")["has_template"]
    assert excel_forms.staged(token) is None                                                  # 임시 보관은 지운다
    dl = client.get("/materials/?export=xlsx")
    assert dl.status_code == 200
    ws = openpyxl.load_workbook(io.BytesIO(dl.data)).active
    assert [ws.cell(row=1, column=c).value for c in range(1, 7)] == ["품번", "품명", "규격", "단위", "분류", "단가"]
    codes = {ws.cell(row=r, column=1).value for r in range(2, ws.max_row + 1)}
    assert {"NEW-1", "NEW-2"} <= codes


def test_upload_without_checkbox_does_not_change_layout(app):  # noqa: F811
    seed.seed()
    client = login(app.test_client(), "admin")
    company = _company_xlsx(title_rows=0, headers=["품번", "품명", "단위"], rows=[["Z-9", "시험", "EA"]])
    res = _upload(client, "/data/upload", company, "회사자재.xlsx")                            # learn_form 없음
    token = re.search(r'name="token" value="([0-9a-f]{32})"', res.get_data(as_text=True)).group(1)
    assert excel_forms.staged(token) is None
    post(client, "/data/upload/apply", {"token": token}, follow_redirects=True)
    assert not excel_forms.load("materials_master")["has_template"]


def test_partner_upload_remembers_layout_and_non_admin_cannot(app):  # noqa: F811
    seed.seed()
    xlsx = _company_xlsx(title_rows=0, headers=["거래처명", "구분", "사업자번호", "담당자", "메모"],
                         rows=[["새공급사", "공급처", "1248100998", "김구매", "첫 거래"]])
    manager = login(app.test_client(), "manager")                                             # 관리자(ADMIN) 아님 → 기억 안 함
    res = _upload(manager, "/partners/import", xlsx, "p.xlsx", learn_form="1")
    html = res.get_data(as_text=True)
    token = re.search(r'name="token" value="([0-9a-f]{32})"', html).group(1)
    assert "양식을 읽었습니다" not in html and excel_forms.staged(token) is None
    post(manager, "/partners/import/apply", {"token": token}, follow_redirects=True)
    assert not excel_forms.load("partners")["has_template"]
    admin = login(app.test_client(), "admin")
    xlsx2 = _company_xlsx(title_rows=0, headers=["거래처명", "구분", "사업자번호", "담당자", "메모"],
                          rows=[["다른공급사", "공급처", "1018100001", "이구매", ""]])
    res = _upload(admin, "/partners/import", xlsx2, "p2.xlsx", learn_form="1")
    html = res.get_data(as_text=True)
    assert "양식을 읽었습니다" in html and "거래처" in html
    token = re.search(r'name="token" value="([0-9a-f]{32})"', html).group(1)
    res = post(admin, "/partners/import/apply", {"token": token}, follow_redirects=True)
    assert "올린 파일의 양식으로 바뀌었습니다" in res.get_data(as_text=True)
    assert excel_forms.load("partners")["has_template"]


def test_admin_learn_button(app):  # noqa: F811
    client = login(app.test_client(), "admin")
    page = client.get("/admin/forms/materials_master").get_data(as_text=True)
    assert "회사 엑셀로 자동 설정" in page
    res = _upload(client, "/admin/forms/materials_master/learn", _company_xlsx(), "한빛.xlsx", follow_redirects=True)
    html = res.get_data(as_text=True)
    assert "양식을 기억했습니다" in html and "머리글 3행" in html
    assert excel_forms.load("materials_master")["has_template"]
    # 편집 화면에서 저장해도(열 위치 hidden 값) 레이아웃이 유지된다
    page = client.get("/admin/forms/materials_master").get_data(as_text=True)
    assert 'name="col" value="7"' in page
    bad = _upload(client, "/admin/forms/materials_master/learn", b"junk", "x.xlsx", follow_redirects=True)
    assert "손상되었습니다" in bad.get_data(as_text=True)
    assert _upload(client, "/admin/forms/nope/learn", _company_xlsx(), "x.xlsx").status_code == 404
