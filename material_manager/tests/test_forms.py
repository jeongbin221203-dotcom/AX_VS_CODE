"""엑셀 양식: 회사 양식 파일에 채워 내려받기 · 회사 엑셀 그대로 올리기."""
from __future__ import annotations

import io
import os
import sys
import tempfile
from pathlib import Path

os.environ.setdefault("MM_DB_PATH", str(Path(tempfile.mkdtemp(prefix="mm_forms_")) / "test.db"))
os.environ.setdefault("MM_PW_ITERATIONS", "1000")
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import openpyxl  # noqa: E402
import pandas as pd  # noqa: E402
import pytest  # noqa: E402
from openpyxl.styles import Border, Font, Side  # noqa: E402

from core import db, excel_forms, reconcile, seed  # noqa: E402
from test_app import PW, app, client, login, post  # noqa: E402,F401

ADMIN = {"id": 1, "name": "관리자", "role": "ADMIN", "ip": ""}


def company_template() -> bytes:
    """회사 재고 보고서 양식: 1행 제목, 2행 작성 정보, 4행 결재란 옆 머리글, 5행 첫 데이터 행 서식."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "재고보고"
    ws["A1"] = "{{제목}}"
    ws["A1"].font = Font(size=16, bold=True)
    ws["A2"] = "기준: {{기간}} / 작성: {{작성자}} / 출력일 {{오늘}} / {{건수}}건"
    ws["F1"], ws["G1"] = "담당", "팀장"                       # 결재란
    for i, h in enumerate(["품번", "품명", "재고수량", "금액"], start=1):
        ws.cell(row=4, column=i, value=h).font = Font(bold=True)
    thin = Side(style="thin")
    for i in range(1, 5):
        ws.cell(row=5, column=i).border = Border(bottom=thin)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.fixture()
def fresh():
    db.reset_database()
    seed.seed()


def test_export_into_company_template(fresh):
    assert excel_forms.upload_template("stock", company_template(), "재고보고.xlsx", ADMIN) == ""
    _, headers = excel_forms.template_headers("stock")
    assert headers[0] == "{{제목}}", "머리글 행을 정하기 전에는 1행을 읽는다"
    problem = excel_forms.save_export("stock", [
        {"source": "자재코드", "header": "품번"}, {"source": "자재명", "header": "품명"},
        {"source": "현재고", "header": "재고수량", "format": "#,##0.00"}, {"source": "재고금액", "header": "금액", "format": "₩#,##0"},
    ], "재고보고", header_row=4, start_row=5, start_col="A", write_header=False, title="월간 재고 보고", actor=ADMIN)
    assert problem == ""
    assert excel_forms.template_headers("stock")[1] == ["품번", "품명", "재고수량", "금액"]
    df = pd.DataFrame({"자재코드": ["A-1", "B-2", "=HACK()"], "자재명": ["가", "나", "다"], "현재고": [1.5, 20.0, 3.0],
                       "재고금액": [1500, 20000, 3000], "규격": ["x", "y", "z"]})
    wb = openpyxl.load_workbook(io.BytesIO(excel_forms.export("stock", df, {"user": "김자재", "period": "2026-09"})))
    ws = wb["재고보고"]
    assert ws["A1"].value == "월간 재고 보고" and ws["A1"].font.size == 16, "양식 서식 유지"
    assert "2026-09" in ws["A2"].value and "김자재" in ws["A2"].value and "3건" in ws["A2"].value
    assert ws["F1"].value == "담당", "결재란 유지"
    assert [ws.cell(row=5, column=i).value for i in range(1, 5)] == ["A-1", "가", 1.5, 1500]
    assert ws["C6"].number_format == "#,##0.00" and ws["D6"].number_format == "₩#,##0"
    assert ws["A6"].border.bottom.style == "thin", "첫 데이터 행 서식을 다음 행에 이어 씀"
    assert ws["A7"].value == "'=HACK()" and ws["A7"].data_type == "s", "사용자 데이터는 수식으로 실행 안 됨"
    assert ws.max_column == 7 and ws["E5"].value is None, "설정하지 않은 열(규격)은 빠짐"


def test_export_without_template_uses_order_and_headers(fresh):
    excel_forms.save_export("ledger", [{"source": "기말", "header": "월말재고"}, {"source": "자재코드", "header": "코드"}],
                            "", 1, 2, "B", True, "", ADMIN)
    df = pd.DataFrame({"자재코드": ["A"], "기말": [5.0], "기초": [1.0]})
    ws = openpyxl.load_workbook(io.BytesIO(excel_forms.export("ledger", df))).active
    assert [ws["B1"].value, ws["C1"].value, ws["B2"].value, ws["C2"].value] == ["월말재고", "코드", 5, "A"]
    assert ws["A1"].value is None, "시작 열 B"


def test_template_security(fresh):
    assert "xlsm" in excel_forms.upload_template("stock", company_template(), "x.xlsm", ADMIN)
    assert "손상" in excel_forms.upload_template("stock", b"PK\x03\x04not-really", "x.xlsx", ADMIN)
    assert "행" in excel_forms.save_export("stock", [], "", 5, 3, "A", True, "", ADMIN), "시작 행은 머리글 아래"
    assert "없는 항목" in excel_forms.save_export("stock", [{"source": "비밀번호"}], "", 1, 2, "A", True, "", ADMIN)


def test_import_company_sheet_with_aliases(fresh):
    assert excel_forms.save_import("material_upload", {"code": ["품번"], "name": ["품명"], "unit_price": ["표준단가"]},
                                   "자재", 3, ADMIN) == ""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "자재"
    ws["A1"] = "자재 등록 요청서"
    ws.append([])
    ws.append(["품번", "품 명", "규격", "표준단가", "비고(무시)"])
    ws.append(["new-9", "새 자재", "10kg", 3000, "메모"])
    buf = io.BytesIO()
    wb.save(buf)
    raw, problem = excel_forms.read_import("material_upload", buf.getvalue(), "요청서.xlsx", 100)
    assert not problem
    assert {"자재코드", "자재명", "규격", "단가"} <= set(raw.columns), "띄어쓰기 달라도 연결"
    assert raw.iloc[0]["자재코드"] == "new-9"
    assert "두 항목" in excel_forms.save_import("material_upload", {"code": ["품번"], "name": ["품번"]}, "", 1, ADMIN)


def test_sap_stock_import_alias(fresh):
    excel_forms.save_import("sap_stock", {"sap_qty": ["재고"]}, "", 1, ADMIN)
    csv = "자재,플랜트,저장 위치,재고\nPKG001,1000,0001,35\n".encode()
    raw, _ = excel_forms.read_import("sap_stock", csv, "mb52.csv", 100)
    df, problem = reconcile.normalize_sap_stock(raw)
    assert not problem and df.iloc[0]["sap_qty"] == 35 and df.iloc[0]["sap_matnr"] == "PKG001"


def test_forms_admin_ui_and_downloads(client, app):
    seed.seed()
    assert client.get("/admin/forms").status_code == 200
    res = post(client, "/admin/forms/stock/template",
               {"file": (io.BytesIO(company_template()), "재고보고.xlsx")}, content_type="multipart/form-data")
    assert res.status_code == 302
    page = client.get("/admin/forms/stock").get_data(as_text=True)
    assert "재고보고" in page
    post(client, "/admin/forms/stock", {"sheet": "재고보고", "header_row": "4", "start_row": "5", "start_col": "A",
                                        "write_header": "1", "title": "재고 보고", "source": ["자재코드", "현재고", ""],
                                        "header": ["품번", "재고수량", ""], "format": ["", "#,##0", ""]})
    ws = openpyxl.load_workbook(io.BytesIO(client.get("/stock/?export=xlsx").data))["재고보고"]
    assert ws["A1"].value == "재고 보고" and ws["A4"].value == "품번" and ws["B4"].value == "재고수량"
    assert ws["A5"].value == "LBL-001", "실제 재고현황 내려받기가 회사 양식으로"
    assert client.get("/admin/forms/stock/sample.xlsx").status_code == 200
    assert client.get("/admin/forms/material_upload/sample.xlsx").status_code == 200
    assert client.get("/admin/forms/sap_stock").status_code == 200
    post(client, "/admin/forms/stock/reset", {})
    ws = openpyxl.load_workbook(io.BytesIO(client.get("/stock/?export=xlsx").data)).active
    assert ws["A1"].value == "자재코드", "되돌리면 기본 양식"
    assert "FORM_UPDATE" in db.query_df("SELECT action FROM audit_log")["action"].tolist()
    manager = login(app.test_client(), "manager")
    assert manager.get("/admin/forms").status_code == 403


def test_material_upload_uses_form(client):
    excel_forms.save_import("material_upload", {"code": ["품번"], "name": ["품명"]}, "", 2, ADMIN)
    csv = "자재 등록 요청\n품번,품명,단위\nq-1,큐 자재,BOX\n".encode()
    res = post(client, "/data/upload", {"file": (io.BytesIO(csv), "req.csv")}, content_type="multipart/form-data")
    assert "1건 반영" in res.get_data(as_text=True)


def test_guess_columns_from_company_headers():
    sources = excel_forms.EXPORT_FORMS["stock"][1]
    assert [excel_forms.guess_source(h, sources) for h in ["품번", "품 명", "재고수량", "금액", "비고"]] ==         ["자재코드", "자재명", "현재고", "재고금액", ""]
