"""엑셀 양식: 회사에서 쓰는 양식 그대로 내려받고, 회사 엑셀 그대로 올린다.

내려받기(export)
  - 시스템관리자가 회사 .xlsx 양식(로고·제목·결재란·서식 포함)을 올려 두면 그 파일에 데이터를 채워 준다.
  - 열마다 '어느 항목을 · 어떤 머리글로 · 어떤 숫자 서식으로' 넣을지, 몇 행 · 몇 열부터 쓸지, 어느 시트인지 정한다.
  - 양식 안의 자리표시를 채운다: {{제목}} {{오늘}} {{작성자}} {{기간}} {{건수}}
  - 데이터 행은 양식의 첫 데이터 행 서식(테두리·글꼴·정렬)을 그대로 이어 쓴다.
  - 양식 파일이 없으면 설정한 열 순서·머리글로 새 파일을 만든다(설정도 없으면 기본 양식).
올리기(import)
  - 회사 엑셀의 열 이름(예: 품번·품명)을 이 시스템 항목에 연결(별칭)하고, 머리글이 몇 행에 있는지·어느 시트인지 정한다.

보안: 양식은 .xlsx만(매크로 파일 .xlsm 거부), 압축 폭탄 검사. 사용자 데이터는 계속 수식으로 실행되지 않게 처리한다.
양식 파일은 파일 저장소(forms/…), 설정은 DB(excel_forms)에 있어 서버가 여러 대여도 같은 양식을 쓴다.
"""

import io
import json
import re
from copy import copy
from datetime import date

import openpyxl
import pandas as pd
from openpyxl.utils import column_index_from_string, get_column_letter

import config
from core import audit, db, storage
from core.utils import neutralize_formula, now_str, xlsx_problem

# ── 양식 목록 ────────────────────────────────────────────────
EXPORT_FORMS: dict[str, tuple[str, list[str]]] = {
    "stock": ("재고현황 (자재별)", ["자재코드", "자재명", "규격", "단위", "분류", "현재고", "안전재고", "부족수량", "단가",
                                   "재고금액", "보관위치", "공급처"]),
    "stock_wh": ("재고현황 (창고별)", ["자재코드", "자재명", "플랜트", "창고", "창고명", "단위", "현재고", "단가", "재고금액"]),
    "stock_lots": ("로트별 재고", ["자재코드", "자재명", "창고", "로트", "유효기한", "남은 일수", "단위", "현재고"]),
    "history": ("거래이력", ["ID", "일자", "구분", "상태", "창고", "자재코드", "자재명", "로트", "단위", "수량", "단가", "금액",
                            "문서번호", "거래처/부서", "구매오더", "원가센터", "이동유형", "SAP", "자재문서", "담당자", "승인자",
                            "이동번호", "비고", "증빙", "등록일시"]),
    "ledger": ("수불부", ["자재코드", "자재명", "단위", "기초", "입고", "출고", "조정", "이동입고", "이동출고", "기말", "단가",
                         "기말금액"]),
    "valuation": ("재고평가", ["자재코드", "자재명", "플랜트", "단위", "기초수량", "기초금액", "입고금액", "출고원가", "조정금액",
                              "이동금액", "기말수량", "평가단가", "기말금액"]),
    "reconcile": ("재고대사", ["판정", "자재코드", "자재명", "창고", "SAP자재", "플랜트", "저장위치", "앱 재고", "미전기",
                              "SAP 예상", "SAP 재고", "차이"]),
    "ap_snapshot": ("월말 미지급·GR/IR 스냅샷", ["월", "발주번호", "공급처", "창고", "발주금액", "입고금액", "계산서 공급가액",
                                                  "GR/IR(입고−계산서)", "지급 상태"]),
    "audit": ("감사로그", ["ID", "일시", "사용자", "행위", "대상", "대상ID", "내용", "IP", "요청번호"]),
    "material_template": ("자재 업로드 양식 (빈 양식)", list(config.MATERIAL_COLS.values())),
    "statement_template": ("거래명세서 품목 양식 (빈 양식)", ["자재코드", "품명", "규격", "수량", "단가", "공급가액", "세액",
                                                         "로트", "유효기한", "발주번호", "발주품목", "비고"]),
    # 기준정보 내려받기 — 올리기 양식과 같은 열이라 고쳐서 그대로 다시 올린다
    "materials_master": ("자재 마스터 (내려받아 고쳐 다시 올리기)", list(config.MATERIAL_COLS.values())),
    "units": ("단위 환산", ["자재코드", "자재명", "기본단위", "단위", "배수", "단위바코드"]),
    "boms": ("BOM 전체", ["제품코드", "기준수량", "부품코드", "수량", "손실률", "출고창고", "메모"]),
    "mrp_demand_template": ("MRP 수요 일괄 등록 양식 (빈 양식)", ["제품코드", "수량", "납기", "메모"]),
    # 화면 목록 내려받기
    "mrp_plans": ("MRP 계획", ["구분", "단계", "자재코드", "자재명", "수량", "단위", "발주·착수일", "필요일", "리드타임", "창고", "근거",
                               "상태", "바꾼 번호"]),
    "work_orders": ("작업지시", ["작업지시", "상태", "출처", "제품코드", "제품명", "수량", "단위", "완료 예정", "부품 창고", "입고 창고",
                                 "투입 금액", "작업지시 번호", "등록자"]),
    "wip": ("재공품", ["작업지시", "제품코드", "제품명", "수량", "단위", "창고", "착수", "완료 예정", "공정", "재공 금액(투입)", "표준 재료비"]),
    "approvals": ("결재함", ["종류", "번호", "단계", "내용", "금액", "요청자", "요청일시", "대결", "기한 넘김"]),
    "my_requests": ("내 요청 현황", ["종류", "번호", "요청일시", "상태", "결재 단계", "금액", "내용", "마지막 처리", "반려 사유"]),
    "approval_history": ("결재 이력", ["처리일시", "종류", "번호", "결정", "금액", "내용", "요청자", "처리자(대결 표시)", "의견"]),
    "quality": ("데이터 점검", ["항목", "심각도", "내용", "고칠 곳"]),
    "purchase_requests": ("구매요청", ["요청번호", "창고", "상태", "결재 단계", "금액", "필요일", "사유", "요청자", "요청일시"]),
    "purchase_orders": ("발주", ["발주번호", "요청번호", "공급처", "창고", "상태", "SAP 발주번호", "금액", "작성자", "작성일시", "승인자"]),
    "payment_match": ("지급 대조", ["발주번호", "공급처", "발주 상태", "발주 금액", "입고 금액(발주 단가)", "계산서 공급가액", "차이",
                                     "지급", "사유"]),
    "variances": ("작업지시 원가 차이", ["작업지시", "제품코드", "제품명", "양품", "표준 원가", "실제 원가", "재료 가격차이",
                                       "재료 수량차이", "노무 차이", "경비 차이", "기타 차이", "차이 합계", "정산", "비고"]),
    "standard_costs": ("표준원가", ["자재코드", "자재명", "단위", "재료", "노무", "경비", "표준원가", "표준시간(분)", "산정일시", "산정자"]),
    "ppv": ("구매 가격차이", ["일자", "자재코드", "자재명", "수량", "단위", "입고 단가", "표준 단가", "가격차이", "발주번호", "거래처"]),
    "hometax": ("홈택스 매입 대사", ["결과", "승인번호", "작성일자", "공급자", "사업자번호", "공급가액(홈택스)", "세액(홈택스)",
                                   "공급가액(증빙)", "세액(증빙)", "증빙번호", "후보 입고", "품목"]),
    "documents": ("증빙", ["번호", "작성일자", "종류", "공급자", "사업자번호", "공급가액", "세액", "합계", "승인번호", "거래", "자재코드",
                           "자재명", "파일", "등록자", "등록일시"]),
    "partner_template": ("거래처 일괄 등록 양식 (빈 양식)", ["거래처코드", "거래처명", "구분", "사업자등록번호", "담당자", "전화",
                                                         "메일", "메모", "다른 이름"]),
    "bom_template": ("BOM 일괄 등록 양식 (빈 양식)", ["제품코드", "기준수량", "부품코드", "수량", "손실률", "출고창고", "메모"]),
    "partners": ("거래처", ["거래처코드", "거래처명", "구분", "사업자등록번호", "담당자", "전화", "메일", "다른 이름", "입고금액",
                          "출고금액", "거래 수", "마지막 거래", "상태"]),
    "productions": ("생산 투입 이력", ["생산번호", "일자", "제품코드", "제품명", "수량", "단위", "부품 창고", "입고 창고",
                                   "재료비", "작업지시", "등록자", "상태"]),
}

# 회사 양식에서 흔히 쓰는 이름 → 이 시스템 항목 (양식을 올리면 열을 자동으로 맞출 때, 올리기 기본 별칭에 쓴다)
SYNONYMS: dict[str, list[str]] = {
    "자재코드": ["품번", "품목코드", "자재번호", "코드", "Item Code", "Part No"],
    "자재명": ["품명", "품목명", "자재 이름", "품목", "Description", "Item Name"],
    "규격": ["사양", "스펙", "Spec"],
    "단위": ["UOM", "Unit"],
    "분류": ["자재그룹", "품목군", "카테고리", "Category"],
    "현재고": ["재고수량", "재고", "현재 재고", "보유수량", "Stock"],
    "안전재고": ["최소재고", "Safety Stock"],
    "단가": ["단위가격", "기준단가", "표준단가", "Price", "Unit Price"],
    "재고금액": ["금액", "재고 금액", "평가금액", "Amount"],
    "보관위치": ["로케이션", "Location", "Bin"],
    "공급처": ["거래처", "공급업체", "업체", "Vendor", "Supplier"],
    "창고": ["창고코드", "Warehouse"],
    "일자": ["날짜", "전기일", "거래일", "Date"],
    "수량": ["Qty", "Quantity"],
    "금액": ["Amount"],
    "문서번호": ["전표번호", "Ref No"],
    "SAP자재번호": ["SAP 자재", "SAP 품번"],
    "바코드": ["바코드번호", "EAN", "JAN", "Barcode", "GTIN"],
    "로트": ["배치", "LOT", "Batch"],
    "유효기한": ["사용기한", "Expiry"],
}

# 올리기: 항목 → (표준 머리글, 기본 별칭)
IMPORT_FORMS: dict[str, tuple[str, dict[str, tuple[str, list[str]]]]] = {
    "unit_upload": ("단위 환산 일괄 등록", {
        "code": ("자재코드", ["품번", "품목코드", "자재번호", "코드", "Item Code"]),
        "unit": ("단위", ["환산단위", "포장단위", "UOM", "Alt Unit"]),
        "factor": ("배수", ["환산수량", "입수", "기본단위 수량", "Factor"]),
        "barcode": ("단위바코드", ["상자바코드", "바코드", "Barcode", "GTIN"]),
    }),
    "mrp_demand_upload": ("MRP 수요 일괄 등록", {
        "code": ("제품코드", ["자재코드", "품번", "품목코드", "Item Code"]),
        "qty": ("수량", ["수요량", "계획수량", "Qty", "Quantity"]),
        "due": ("납기", ["납기일", "출하일", "필요일", "Due Date"]),
        "note": ("메모", ["비고", "고객", "주문번호", "Note"]),
    }),
    "partner_upload": ("거래처 일괄 등록", {
        "code": ("거래처코드", ["코드", "거래처 코드", "Vendor", "Customer", "LIFNR", "KUNNR"]),
        "name": ("거래처명", ["상호", "업체명", "거래처", "회사명", "Name"]),
        "kind": ("구분", ["거래처구분", "유형", "Type"]),
        "biz_no": ("사업자등록번호", ["사업자번호", "등록번호", "Business No"]),
        "contact": ("담당자", ["담당", "Contact"]),
        "phone": ("전화", ["전화번호", "연락처", "Tel", "Phone"]),
        "email": ("메일", ["이메일", "E-mail", "Email"]),
        "note": ("메모", ["비고", "Note"]),
        "aliases": ("다른 이름", ["별칭", "다른이름", "약칭"]),
    }),
    "bom_upload": ("BOM 일괄 등록", {
        "product": ("제품코드", ["모품목", "상위품목", "완제품코드", "제품", "Parent", "Material"]),
        "base_qty": ("기준수량", ["기준 수량", "Base Qty"]),
        "component": ("부품코드", ["자품목", "하위품목", "구성품", "부품", "Component"]),
        "qty": ("수량", ["소요량", "단위수량", "Qty", "Quantity"]),
        "scrap_pct": ("손실률", ["로스율", "불량률", "Scrap", "Scrap %"]),
        "wh": ("출고창고", ["창고", "저장위치", "Storage Location"]),
        "note": ("메모", ["비고", "Note"]),
    }),
    "material_upload": ("자재 일괄 업로드", {k: (v, SYNONYMS.get(v, [])) for k, v in config.MATERIAL_COLS.items()}),
    "sap_stock": ("SAP 재고 (재고 대사)", {
        "sap_matnr": ("SAP자재번호", ["자재", "자재번호", "Material", "MATNR"]),
        "plant": ("플랜트", ["Plant", "WERKS"]),
        "sloc": ("저장위치", ["저장 위치", "Storage Location", "SLoc", "LGORT"]),
        "sap_qty": ("수량", ["가용재고", "제한없는 사용", "Unrestricted", "LABST", "Quantity"]),
    }),
    # 거래명세서 품목 줄 (core/statements.py). 회사·공급처 명세서마다 머리글이 달라 별칭을 넉넉히 둔다.
    "statement_lines": ("거래명세서 품목 (거래명세서 입출고)", {
        "code": ("자재코드", ["품번", "품목코드", "자재번호", "코드", "Item Code", "Part No", "SAP자재번호"]),
        "name": ("품명", ["자재명", "품목명", "품목", "상품명", "Description", "Item Name"]),
        "spec": ("규격", ["사양", "스펙", "Spec"]),
        "qty": ("수량", ["Qty", "Quantity", "입고수량", "출고수량"]),
        "unit_price": ("단가", ["단위가격", "Price", "Unit Price"]),
        "supply": ("공급가액", ["공급가", "금액", "Amount"]),
        "tax": ("세액", ["부가세", "VAT", "Tax"]),
        "lot_no": ("로트", ["배치", "LOT", "Batch", "로트번호"]),
        "expiry_date": ("유효기한", ["사용기한", "Expiry"]),
        "po_no": ("발주번호", ["PO", "PO번호", "구매오더"]),
        "po_item": ("발주품목", ["PO품목", "품목번호"]),
        "note": ("비고", ["메모", "적요"]),
    }),
}

NUMBER_FORMATS = {"": "기본", "#,##0": "1,234", "#,##0.00": "1,234.56", "0": "1234", "₩#,##0": "₩1,234",
                  "0.0%": "12.3%", "yyyy-mm-dd": "2026-09-27 (날짜)", "@": "텍스트"}
PLACEHOLDERS = ("{{제목}}", "{{오늘}}", "{{작성자}}", "{{기간}}", "{{건수}}")
COL_RE = re.compile(r"^[A-Z]{1,3}$")
MAX_TEMPLATE_BYTES = 5 * 1024 * 1024


def label(form_key: str) -> str:
    return (EXPORT_FORMS.get(form_key) or IMPORT_FORMS.get(form_key) or (form_key,))[0]


def _key(form_key: str) -> str:
    return f"forms/{form_key}.xlsx"


# ── 설정 저장 ────────────────────────────────────────────────
def load(form_key: str) -> dict:
    row = None
    with db.get_conn() as conn:
        row = conn.execute("SELECT config, has_template, updated_by, updated_at FROM excel_forms WHERE form_key = ?",
                           (form_key,)).fetchone()
    if row is None:
        return {"has_template": False}
    cfg = json.loads(row["config"] or "{}")
    cfg.update(has_template=bool(row["has_template"]), updated_by=row["updated_by"], updated_at=row["updated_at"])
    return cfg


def _save(conn, form_key: str, cfg: dict, has_template: bool | None, actor: dict | None) -> None:
    keep = {k: v for k, v in cfg.items() if k not in ("has_template", "updated_by", "updated_at")}
    prev = conn.execute("SELECT has_template FROM excel_forms WHERE form_key = ?", (form_key,)).fetchone()
    tmpl = int(has_template if has_template is not None else (prev["has_template"] if prev else 0))
    conn.execute("INSERT INTO excel_forms (form_key, config, has_template, updated_by, updated_at) VALUES (?, ?, ?, ?, ?) "
                 "ON CONFLICT (form_key) DO UPDATE SET config = excluded.config, has_template = excluded.has_template, "
                 "updated_by = excluded.updated_by, updated_at = excluded.updated_at",
                 (form_key, json.dumps(keep, ensure_ascii=False), tmpl, (actor or audit.SYSTEM)["name"], now_str()))


def save_export(form_key: str, columns: list[dict], sheet: str, header_row: int, start_row: int, start_col: str,
                write_header: bool, title: str, actor: dict | None) -> str:
    """columns: [{"source": 기본 머리글, "header": 회사 머리글, "format": 숫자 서식}] — 순서대로. 문제 있으면 사유."""
    if form_key not in EXPORT_FORMS:
        return "없는 양식입니다."
    sources = EXPORT_FORMS[form_key][1]
    cols = []
    for c in columns:
        if not c.get("source"):
            continue
        if c["source"] not in sources:
            return f"'{c['source']}'는 이 양식에 없는 항목입니다."
        if c.get("format", "") not in NUMBER_FORMATS:
            return "숫자 서식을 목록에서 고르세요."
        cols.append({"source": c["source"], "header": (c.get("header") or c["source"]).strip()[:100],
                     "format": c.get("format", "")})
    start_col = (start_col or "A").strip().upper()
    if not COL_RE.match(start_col):
        return "시작 열은 A, B, C … 처럼 입력하세요."
    if not (1 <= header_row <= 200 and 1 <= start_row <= 1000):
        return "행 번호를 확인하세요."
    if write_header and start_row <= header_row:
        return "데이터 시작 행은 머리글 행보다 아래여야 합니다."
    cfg = {"columns": cols, "sheet": sheet.strip()[:31], "header_row": header_row, "start_row": start_row,
           "start_col": start_col, "write_header": write_header, "title": title.strip()[:100]}
    with db.transaction() as conn:
        _save(conn, form_key, cfg, None, actor)
        audit.record(conn, actor, "FORM_UPDATE", "excel_form", form_key, {"columns": len(cols), "sheet": cfg["sheet"]})
    return ""


def save_import(form_key: str, aliases: dict[str, list[str]], sheet: str, header_row: int, actor: dict | None) -> str:
    if form_key not in IMPORT_FORMS:
        return "없는 양식입니다."
    if not 1 <= header_row <= 200:
        return "머리글 행 번호를 확인하세요."
    fields = IMPORT_FORMS[form_key][1]
    clean = {k: [a.strip()[:100] for a in v if a.strip()] for k, v in aliases.items() if k in fields}
    seen: dict[str, str] = {}
    for field, names in clean.items():
        for n in names + [fields[field][0]]:
            other = seen.get(_norm(n))
            if other and other != field:
                return f"'{n}'가 두 항목({fields[other][0]}, {fields[field][0]})에 함께 지정됐습니다."
            seen[_norm(n)] = field
    with db.transaction() as conn:
        _save(conn, form_key, {"aliases": clean, "sheet": sheet.strip()[:31], "header_row": header_row}, None, actor)
        audit.record(conn, actor, "FORM_UPDATE", "excel_form", form_key, {"aliases": clean, "header_row": header_row})
    return ""


def upload_template(form_key: str, data: bytes, filename: str, actor: dict | None) -> str:
    """회사 양식 파일(.xlsx)을 올린다. 문제 있으면 사유."""
    if form_key not in EXPORT_FORMS:
        return "없는 양식입니다."
    if not filename.lower().endswith(".xlsx"):
        return "엑셀 통합 문서(.xlsx)만 올릴 수 있습니다. 매크로 파일(.xlsm)은 받지 않습니다."
    if len(data) > MAX_TEMPLATE_BYTES:
        return "양식 파일은 5MB 이하만 올릴 수 있습니다."
    problem = xlsx_problem(data, config.XLSX_MAX_UNCOMPRESSED, config.XLSX_MAX_RATIO)
    if problem:
        return problem
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data))
    except Exception:
        return "엑셀 파일을 열 수 없습니다. 엑셀에서 .xlsx로 다시 저장해 올려 주세요."
    if wb.vba_archive is not None:
        return "매크로가 들어 있는 파일은 받지 않습니다."
    storage.get().put(_key(form_key), data)
    cfg = load(form_key)
    if not cfg.get("sheet"):
        cfg["sheet"] = wb.active.title
    with db.transaction() as conn:
        _save(conn, form_key, cfg, True, actor)
        audit.record(conn, actor, "FORM_UPDATE", "excel_form", form_key, {"template": filename, "sheets": wb.sheetnames})
    return ""


def remove_template(form_key: str, actor: dict | None) -> None:
    storage.get().delete(_key(form_key))
    with db.transaction() as conn:
        _save(conn, form_key, load(form_key), False, actor)
        audit.record(conn, actor, "FORM_UPDATE", "excel_form", form_key, {"template": None})


def reset(form_key: str, actor: dict | None) -> None:
    storage.get().delete(_key(form_key))
    with db.transaction() as conn:
        conn.execute("DELETE FROM excel_forms WHERE form_key = ?", (form_key,))
        audit.record(conn, actor, "FORM_UPDATE", "excel_form", form_key, {"reset": True})


def template_bytes(form_key: str) -> bytes | None:
    return storage.get().get(_key(form_key)) if load(form_key).get("has_template") else None


def template_headers(form_key: str) -> tuple[list[str], list[str]]:
    """(시트 목록, 설정한 머리글 행의 값들) — 화면에서 회사 양식의 열 이름을 보여 주려고."""
    data = template_bytes(form_key)
    if data is None:
        return [], []
    cfg = load(form_key)
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True)
    ws = wb[cfg["sheet"]] if cfg.get("sheet") in wb.sheetnames else wb.active
    row = int(cfg.get("header_row") or 1)
    start = column_index_from_string(cfg.get("start_col") or "A")
    values = [c for c in next(ws.iter_rows(min_row=row, max_row=row, min_col=start, values_only=True), ())]
    while values and values[-1] in (None, ""):
        values.pop()
    return wb.sheetnames, ["" if v is None else str(v) for v in values]


# ── 내려받기 ─────────────────────────────────────────────────
def _fill_placeholders(ws, ctx: dict) -> None:
    values = {"{{제목}}": ctx.get("title", ""), "{{오늘}}": date.today().isoformat(),
              "{{작성자}}": ctx.get("user", ""), "{{기간}}": ctx.get("period", ""), "{{건수}}": f"{ctx.get('rows', 0):,}"}
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and "{{" in cell.value:
                text = cell.value
                for k, v in values.items():
                    text = text.replace(k, str(neutralize_formula(v)) if v else "")
                cell.value = text


def export(form_key: str, df: pd.DataFrame, ctx: dict | None = None) -> bytes:
    """df(기본 머리글)를 설정한 양식으로 엑셀 파일 내용으로 만든다. ctx: title·user·period."""
    from core import names
    cfg = load(form_key)
    ctx = {**(ctx or {}), "title": cfg.get("title") or label(form_key), "rows": len(df)}
    cols = [c for c in cfg.get("columns") or [] if c["source"] in df.columns] or \
           [{"source": c, "header": names.relabel(str(c)), "format": ""} for c in df.columns]   # 회사가 바꾼 거래 이름
    template = template_bytes(form_key)
    if template is not None:
        wb = openpyxl.load_workbook(io.BytesIO(template))
        ws = wb[cfg["sheet"]] if cfg.get("sheet") in wb.sheetnames else wb.active
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = (cfg.get("sheet") or label(form_key))[:31].replace("/", "·")
    header_row = int(cfg.get("header_row") or 1)
    write_header = cfg.get("write_header", True)
    start_row = int(cfg.get("start_row") or header_row + 1)
    first_col = column_index_from_string(cfg.get("start_col") or "A")

    if write_header:
        for i, c in enumerate(cols):
            cell = ws.cell(row=header_row, column=first_col + i, value=neutralize_formula(c["header"]))
            if template is None:
                cell.font = openpyxl.styles.Font(bold=True)
                cell.fill = openpyxl.styles.PatternFill("solid", fgColor="E8EDF9")
    # 양식의 첫 데이터 행 서식을 이어 쓴다
    styles = [copy(ws.cell(row=start_row, column=first_col + i)._style) for i in range(len(cols))]
    formats = [c.get("format", "") for c in cols]
    data = df[[c["source"] for c in cols]]
    for r, values in enumerate(data.itertuples(index=False), start=start_row):
        for i, v in enumerate(values):
            if v is None or (not isinstance(v, str) and pd.isna(v)):
                v = None
            elif hasattr(v, "item"):
                v = v.item()                           # numpy 숫자 → 파이썬 숫자
            cell = ws.cell(row=r, column=first_col + i, value=neutralize_formula(v))
            if template is not None and r > start_row:
                cell._style = copy(styles[i])
            if formats[i]:
                cell.number_format = formats[i]
    if template is None:                               # 새 파일이면 열 너비를 대충 맞춘다
        for i, c in enumerate(cols):
            width = max([len(str(c["header"]))] + [len(str(v)) for v in data.iloc[:200, i].tolist()]) + 2
            ws.column_dimensions[get_column_letter(first_col + i)].width = min(max(width * 1.2, 8), 50)
        ws.freeze_panes = ws.cell(row=header_row + 1, column=first_col)
    _fill_placeholders(ws, ctx)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def sample_df(form_key: str) -> pd.DataFrame:
    """양식 미리보기용 예시 두 줄."""
    cols = EXPORT_FORMS[form_key][1]
    return pd.DataFrame([[f"예시{i}-{c}" for c in cols] for i in (1, 2)], columns=cols)


# ── 올리기 ───────────────────────────────────────────────────
def _norm(text: str) -> str:
    return re.sub(r"[\s_()\[\]/·.-]", "", str(text)).lower()


def guess_source(header: str, sources: list[str]) -> str:
    """회사 머리글에 맞는 이 시스템 항목을 짐작한다 (같은 이름 → 동의어 순)."""
    n = _norm(header)
    for src in sources:
        if _norm(src) == n:
            return src
    for src in sources:
        if any(_norm(x) == n for x in SYNONYMS.get(src, [])):
            return src
    return ""


def read_import(form_key: str, data: bytes, filename: str, max_rows: int) -> tuple[pd.DataFrame, str]:
    """회사 엑셀/CSV를 읽어 표준 머리글로 바꾼다. (df, 문제)"""
    title, fields = IMPORT_FORMS[form_key]
    cfg = load(form_key)
    header = int(cfg.get("header_row") or 1) - 1
    name = filename.lower()
    try:
        if name.endswith(".csv"):
            raw = None
            for enc in ("utf-8-sig", "cp949"):           # 한글 엑셀 'CSV (쉼표로 분리)' 저장은 cp949
                try:
                    raw = pd.read_csv(io.BytesIO(data), header=header, nrows=max_rows + 1, encoding=enc)
                    break
                except UnicodeDecodeError:
                    continue
            if raw is None:
                raise ValueError("encoding")
        else:
            xl = pd.ExcelFile(io.BytesIO(data))
            sheet = cfg.get("sheet") if cfg.get("sheet") in xl.sheet_names else xl.sheet_names[0]
            raw = xl.parse(sheet, header=header, nrows=max_rows + 1)
    except Exception:
        return pd.DataFrame(), "파일을 읽을 수 없습니다. 엑셀(.xlsx) 또는 CSV(UTF-8·한글 엑셀 기본)로 저장해 주세요."
    lookup: dict[str, str] = {}
    for field, (std, defaults) in fields.items():
        for n in [std, *defaults, *cfg.get("aliases", {}).get(field, [])]:
            lookup.setdefault(_norm(n), std)
    raw = raw.rename(columns=lambda c: lookup.get(_norm(c), str(c).strip()))
    raw = raw.loc[:, ~raw.columns.duplicated()]
    return raw, ""
