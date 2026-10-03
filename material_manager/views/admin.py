"""시스템관리자 화면: 사용자·데이터 범위 · 플랜트·창고 · 배치 · 감사로그."""

import json
from datetime import date, timedelta

from flask import Blueprint, abort, flash, redirect, request, url_for

import config
from core import audit, auth, doctor, jobs, org, repository as repo, sso, version
from views.helpers import (Table, a_date, actor, as_id, f_str, form_response, log_export, page_arg, pager, render_page,
                           role_required, xlsx_response)

bp = Blueprint("admin", __name__, url_prefix="/admin")


# ── 사용자 ───────────────────────────────────────────────────
@bp.get("/users")
@role_required("ADMIN")
def users():
    df = auth.users_df()
    plants, whs = org.plants_df(), org.warehouses_df()
    plant_name = {int(r.id): r.code for r in plants.itertuples()}
    wh_name = {int(r.id): r.code for r in whs.itertuples()}
    records = df.to_dict("records")
    for u in records:
        p, w = org.user_scope(u["id"])
        u["scope_plants"], u["scope_whs"] = p, w
        u["scope_ver"] = version.of_scope(u["all_warehouses"], p, w)
        u["scope_label"] = ("전체" if u["role"] == "ADMIN" or u["all_warehouses"] else
                            ", ".join([plant_name.get(i, "?") + "(플랜트)" for i in sorted(p)]
                                      + [wh_name.get(i, "?") for i in sorted(w)]) or "없음")
    view = df.copy()
    view["role"] = view["role"].map(auth.role_label)
    view["active"] = view["active"].map({1: "사용", 0: "중지"})
    view["must_change_pw"] = view["must_change_pw"].map({1: "변경 필요", 0: ""})
    view["locked_until"] = view["locked_until"].fillna("")
    view["all_warehouses"] = [u["scope_label"] for u in records]
    view["auth_source"] = view["auth_source"].map({"sso": "SSO", "local": "비밀번호"}).fillna("비밀번호")
    view = view.drop(columns=["failed_count"]).rename(columns={"auth_source": "로그인",
        "id": "ID", "username": "아이디", "name": "이름", "role": "역할", "active": "상태", "all_warehouses": "데이터 범위",
        "email": "알림 메일", "messenger_id": "메신저 아이디",
        "must_change_pw": "비밀번호", "locked_until": "잠김 해제", "last_login_at": "최근 로그인",
        "created_at": "등록일시"})
    return render_page("admin_users.html", "users", grid=Table(view, {"ID": "{}"},
                       tones=["muted" if a == 0 else None for a in df["active"]]),
                       users=records, plants=plants.to_dict("records"), warehouses=whs.to_dict("records"),
                       sso_on=sso.enabled(), sso_outage=sso.outage_until() if sso.outage_active() else "",
                       outage_max=config.SSO_OUTAGE_MAX_HOURS)


@bp.post("/sso-outage")
@role_required("ADMIN")
def sso_outage():
    """사내 로그인 서버(IdP) 장애 때 비밀번호 계정 로그인을 정해진 시간만 허용한다."""
    try:
        hours = float(f_str("hours") or 0)
    except ValueError:
        abort(400, "시간을 숫자로 입력하세요.")
    until = sso.set_outage(hours, actor())
    flash(f"SSO 장애 모드를 {until}까지 켰습니다. 비밀번호 계정으로 로그인할 수 있습니다." if until
          else "SSO 장애 모드를 껐습니다.", "warning" if until else "success")
    return redirect(url_for("admin.users"))


@bp.post("/users")
@role_required("ADMIN")
def create_user():
    result = auth.create_user(f_str("username"), f_str("name"), f_str("role"),
                              request.form.get("password", ""), actor())
    flash(result.message + (" 첫 로그인 때 비밀번호를 바꾸게 됩니다." if result.ok else ""),
          "success" if result.ok else "error")
    return redirect(url_for("admin.users"))


@bp.post("/users/<int:user_id>")
@role_required("ADMIN")
def update_user(user_id: int):
    result = auth.update_user(user_id, f_str("name"), f_str("role"), request.form.get("active") == "1", actor(),
                              expected=f_str("_ver") or None)
    flash(result.message, "success" if result.ok else "error")
    return redirect(url_for("admin.users"))


@bp.post("/users/<int:user_id>/email")
@role_required("ADMIN")
def user_email(user_id: int):
    result = auth.set_email(user_id, f_str("email"), actor(), messenger_id=f_str("messenger_id"))
    flash(result.message, "success" if result.ok else "error")
    return redirect(url_for("admin.users"))


@bp.post("/users/<int:user_id>/scope")
@role_required("ADMIN")
def user_scope(user_id: int):
    result = org.set_user_scope(user_id, request.form.get("all") == "1",
                                repo.ids_in(request.form.getlist("plant")),
                                repo.ids_in(request.form.getlist("warehouse")), actor(), expected=f_str("_ver") or None)
    flash(result.message, "success" if result.ok else "error")
    return redirect(url_for("admin.users"))


@bp.post("/users/<int:user_id>/password")
@role_required("ADMIN")
def reset_password(user_id: int):
    result = auth.reset_password(user_id, request.form.get("password", ""), actor())
    flash(result.message, "success" if result.ok else "error")
    return redirect(url_for("admin.users"))


# ── 플랜트 · 창고 ───────────────────────────────────────────
@bp.get("/org")
@role_required("ADMIN")
def org_page():
    return render_page("admin_org.html", "org", plants=org.plants_df().to_dict("records"),
                       warehouses=org.warehouses_df().to_dict("records"))


@bp.post("/org/plant")
@role_required("ADMIN")
def plant_create():
    result = org.create_plant(f_str("code"), f_str("name"), f_str("sap_plant"), actor())
    flash(result.message, "success" if result.ok else "error")
    return redirect(url_for("admin.org_page"))


@bp.post("/org/plant/<int:plant_id>")
@role_required("ADMIN")
def plant_update(plant_id: int):
    result = org.update_plant(plant_id, f_str("name"), f_str("sap_plant"), request.form.get("active") == "1", actor(),
                              expected=f_str("_ver") or None)
    flash(result.message, "success" if result.ok else "error")
    return redirect(url_for("admin.org_page"))


@bp.post("/org/warehouse")
@role_required("ADMIN")
def warehouse_create():
    raw = f_str("plant_id")
    result = org.create_warehouse(as_id(raw) or 0, f_str("code"), f_str("name"),
                                  f_str("sap_sloc"), actor())
    flash(result.message, "success" if result.ok else "error")
    return redirect(url_for("admin.org_page"))


@bp.post("/org/warehouse/<int:wh_id>")
@role_required("ADMIN")
def warehouse_update(wh_id: int):
    result = org.update_warehouse(wh_id, f_str("name"), f_str("sap_sloc"), request.form.get("active") == "1", actor(),
                                  expected=f_str("_ver") or None)
    flash(result.message, "success" if result.ok else "error")
    return redirect(url_for("admin.org_page"))


# ── 배치 ─────────────────────────────────────────────────────
@bp.get("/jobs")
@role_required("ADMIN")
def jobs_page(checks=None):
    status = jobs.status_df()
    runs = jobs.runs_df().rename(columns={"id": "ID", "name": "작업", "holder": "실행 서버", "started_at": "시작",
                                          "finished_at": "종료", "status": "결과", "message": "메시지"})
    from core import notify
    notes = notify.recent_df()
    notes_view = notes.assign(status=notes["status"].map(notify.STATUS),
                              channel=notes["channel"].fillna("email").map(lambda c: notify.CHANNELS.get(c or "email", c))
                              ).rename(columns={
        "id": "ID", "created_at": "만든 시각", "channel": "보낼 길", "to_addr": "받는 사람", "subject": "제목", "status": "상태",
        "tries": "시도", "last_error": "오류", "sent_at": "보낸 시각"})
    return render_page("admin_jobs.html", "jobs", status=status.to_dict("records"), worker=jobs.worker_id(),
                       lease=config.JOB_LEASE_SECONDS, checks=checks, check_label=doctor.LABEL,
                       runs=Table(runs, {"ID": "{}"}, tones=["danger" if s == "ERROR" else None for s in runs["결과"]]),
                       notify_mode=config.NOTIFY_MODE, smtp_host=config.SMTP_HOST,
                       channels_on={"메일": bool(config.SMTP_HOST), "잔디": notify.jandi_on(), "네이버웍스": notify.naverworks_on()},
                       notes=Table(notes_view, {"ID": "{}", "시도": "{}"},
                                   tones=["danger" if s == "FAILED" else None for s in notes["status"]]))


@bp.post("/notify-test")
@role_required("ADMIN")
def notify_test():
    from core import notify
    n = notify.test_message(actor())
    flash(f"시험 알림 {n}건을 쌓았습니다. 1분 안에 배치가 보냅니다(지금 보내려면 '결재 알림 보내기'를 지금 실행)."
          if n else "쌓은 알림이 없습니다 — 내 메일·메신저 아이디가 없거나 알림이 꺼져 있습니다(MM_NOTIFY_MODE).", "info")
    return redirect(url_for("admin.jobs_page"))


@bp.post("/doctor")
@role_required("ADMIN")
def doctor_run():
    """운영 점검: 실제 DB·저장소·ERP·SSO·백업·배치 연결 확인 (거래는 보내지 않음)."""
    checks = doctor.run()
    audit.log(actor(), "DOCTOR", "system", "", {c.name: c.status for c in checks})
    return jobs_page(checks)


@bp.post("/jobs/<name>")
@role_required("ADMIN")
def job_run(name: str):
    if name not in jobs.JOBS:
        flash("없는 작업입니다.", "error")
        return redirect(url_for("admin.jobs_page"))
    audit.log(actor(), "JOB_RUN", "job", name)
    ran, msg = jobs.run(name, force=True)
    flash(f"{jobs.JOBS[name].label}: {msg}", "success" if ran and not msg.startswith("오류") else "warning")
    return redirect(url_for("admin.jobs_page"))


# ── 감사로그 ─────────────────────────────────────────────────
@bp.get("/audit")
@role_required("ADMIN")
def audit_log():
    start = a_date("start", date.today() - timedelta(days=7))
    end = a_date("end", date.today())
    user = request.args.get("user", "")
    action = request.args.get("action", "")
    keyword = request.args.get("q", "").strip()
    if request.args.get("export") == "xlsx":
        full = audit.audit_df(start.isoformat(), end.isoformat(), user, action, keyword, limit=config.EXPORT_MAX_ROWS)
        log_export("audit_log", len(full), start=start.isoformat(), end=end.isoformat())
        full["action"] = full["action"].map(lambda a: audit.ACTIONS.get(a, a))
        full = full.rename(columns={"id": "ID", "at": "일시", "user_name": "사용자", "action": "행위", "entity": "대상",
                                    "entity_id": "대상ID", "detail": "내용", "ip": "IP"})
        return form_response("audit", full, f"감사로그_{start:%Y%m%d}_{end:%Y%m%d}.xlsx", period=f"{start} ~ {end}")
    df, total = audit.audit_page(start.isoformat(), end.isoformat(), user, action, keyword,
                                 page=page_arg(), size=config.PAGE_SIZE)
    view = df.copy()
    view["action"] = view["action"].map(lambda a: audit.ACTIONS.get(a, a))
    view["detail"] = view["detail"].map(_short)
    view = view.rename(columns={"id": "ID", "at": "일시", "user_name": "사용자", "action": "행위",
                                "entity": "대상", "entity_id": "대상ID", "detail": "내용", "ip": "IP"})
    return render_page("admin_audit.html", "audit", start=start, end=end, user=user, action=action,
                       keyword=keyword, users=audit.user_names(), actions=audit.ACTIONS,
                       count=total, pager=pager(total, page_arg()), grid=Table(view, {"ID": "{}"}))


def _short(detail: str) -> str:
    """JSON 내용을 표에서 읽기 좋게 줄인다."""
    if not detail:
        return ""
    try:
        data = json.loads(detail)
    except ValueError:
        return detail[:120]
    text = ", ".join(f"{k}: {v}" for k, v in data.items() if v not in (None, "", []))
    return text if len(text) <= 160 else text[:157] + "..."


# ── 엑셀 양식 ────────────────────────────────────────────────
@bp.get("/forms")
@role_required("ADMIN")
def forms_list():
    from core import excel_forms
    rows = []
    for key, (title, _cols) in excel_forms.EXPORT_FORMS.items():
        cfg = excel_forms.load(key)
        rows.append({"key": key, "title": title, "kind": "내려받기", "custom": bool(cfg.get("columns")),
                     "template": cfg.get("has_template"), "updated": cfg.get("updated_at", ""), "by": cfg.get("updated_by", "")})
    for key, (title, _) in excel_forms.IMPORT_FORMS.items():
        cfg = excel_forms.load(key)
        rows.append({"key": key, "title": title, "kind": "올리기", "custom": bool(cfg.get("aliases")),
                     "template": False, "updated": cfg.get("updated_at", ""), "by": cfg.get("updated_by", "")})
    return render_page("admin_forms.html", "forms", forms=rows, placeholders=excel_forms.PLACEHOLDERS)


@bp.route("/forms/<key>", methods=["GET", "POST"])
@role_required("ADMIN")
def form_edit(key: str):
    from core import excel_forms
    if key not in excel_forms.EXPORT_FORMS and key not in excel_forms.IMPORT_FORMS:
        abort(404)
    if key in excel_forms.IMPORT_FORMS:
        title, fields = excel_forms.IMPORT_FORMS[key]
        if request.method == "POST":
            aliases = {f: request.form.get(f"alias_{f}", "").replace("\n", ",").split(",") for f in fields}
            problem = excel_forms.save_import(key, aliases, f_str("sheet"), _int("header_row", 1), actor())
            flash(problem or "양식을 저장했습니다.", "error" if problem else "success")
            if not problem:
                return redirect(url_for("admin.form_edit", key=key))
        cfg = excel_forms.load(key)
        return render_page("admin_form_import.html", "forms", key=key, form_title=title, fields=fields, cfg=cfg)

    title, sources = excel_forms.EXPORT_FORMS[key]
    if request.method == "POST":
        cols = [{"source": s, "header": h, "format": fmt} for s, h, fmt in
                zip(request.form.getlist("source"), request.form.getlist("header"), request.form.getlist("format"))]
        problem = excel_forms.save_export(key, cols, f_str("sheet"), _int("header_row", 1), _int("start_row", 2),
                                          f_str("start_col") or "A", request.form.get("write_header") == "1",
                                          f_str("title"), actor())
        flash(problem or "양식을 저장했습니다. '미리보기'로 결과를 확인하세요.", "error" if problem else "success")
        if not problem:
            return redirect(url_for("admin.form_edit", key=key))
    cfg = excel_forms.load(key)
    sheets, tmpl_headers = excel_forms.template_headers(key)
    if cfg.get("columns"):
        rows = cfg["columns"]
    elif tmpl_headers:                                 # 회사 양식의 머리글에 맞춰 항목을 짐작해 채운다
        rows = [{"source": excel_forms.guess_source(h, sources), "header": h, "format": ""} for h in tmpl_headers]
    else:
        rows = [{"source": s, "header": s, "format": ""} for s in sources]
    rows = rows + [{"source": "", "header": "", "format": ""}] * 3
    return render_page("admin_form_export.html", "forms", key=key, form_title=title, sources=sources, cfg=cfg,
                       rows=rows, sheets=sheets, tmpl_headers=tmpl_headers, formats=excel_forms.NUMBER_FORMATS,
                       placeholders=excel_forms.PLACEHOLDERS)


def _int(name: str, default: int) -> int:
    value = as_id(f_str(name))
    return default if value is None else value


@bp.post("/forms/<key>/template")
@role_required("ADMIN")
def form_template(key: str):
    from core import excel_forms
    file = request.files.get("file")
    if not file or not file.filename:
        flash("양식 파일을 고르세요.", "error")
    else:
        problem = excel_forms.upload_template(key, file.read(excel_forms.MAX_TEMPLATE_BYTES + 1), file.filename, actor())
        flash(problem or "회사 양식 파일을 올렸습니다. 열 연결과 시작 행을 확인한 뒤 저장하세요.",
              "error" if problem else "success")
    return redirect(url_for("admin.form_edit", key=key))


@bp.post("/forms/<key>/template/delete")
@role_required("ADMIN")
def form_template_delete(key: str):
    from core import excel_forms
    excel_forms.remove_template(key, actor())
    flash("양식 파일을 뺐습니다. 설정한 열 순서·머리글로 새 파일을 만듭니다.", "success")
    return redirect(url_for("admin.form_edit", key=key))


@bp.post("/forms/<key>/reset")
@role_required("ADMIN")
def form_reset(key: str):
    from core import excel_forms
    excel_forms.reset(key, actor())
    flash("기본 양식으로 되돌렸습니다.", "success")
    return redirect(url_for("admin.form_edit", key=key))


@bp.get("/forms/<key>/sample.xlsx")
@role_required("ADMIN")
def form_sample(key: str):
    """미리보기: 내려받기 양식은 예시 두 줄을 채운 결과, 올리기 양식은 회사 머리글로 된 빈 파일."""
    from core import excel_forms
    if key in excel_forms.EXPORT_FORMS:
        return form_response(key, excel_forms.sample_df(key), f"양식미리보기_{key}.xlsx", period="예시 기간")
    if key not in excel_forms.IMPORT_FORMS:
        abort(404)
    cfg = excel_forms.load(key)
    fields = excel_forms.IMPORT_FORMS[key][1]
    headers = [(cfg.get("aliases", {}).get(f) or [std])[0] for f, (std, _) in fields.items()]
    import io
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = (cfg.get("sheet") or "Sheet1")[:31]
    for i, h in enumerate(headers, start=1):
        ws.cell(row=int(cfg.get("header_row") or 1), column=i, value=h)
    buf = io.BytesIO()
    wb.save(buf)
    return xlsx_response(buf.getvalue(), f"업로드양식_{key}.xlsx")


@bp.get("/forms/<key>/template.xlsx")
@role_required("ADMIN")
def form_template_download(key: str):
    from core import excel_forms
    data = excel_forms.template_bytes(key)
    if data is None:
        abort(404)
    log_export("form_template", 0, form=key)
    return xlsx_response(data, f"회사양식_{key}.xlsx")
